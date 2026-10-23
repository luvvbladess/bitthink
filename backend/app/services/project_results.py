"""Account-owned structured results with optimistic revisions and real evidence."""
import io
import json
import math
import re
from contextvars import ContextVar
from app.services import project_memory as memory
turn_results: ContextVar[list | None] = ContextVar('project_turn_results',default=None)


def validate(spec):
    if not isinstance(spec, dict): raise ValueError('Нужен объект результата.')
    kind = spec.get('type')
    if kind not in {'table', 'report', 'slides'}: raise ValueError('Поддерживаются таблица, отчёт и слайды.')
    title = str(spec.get('title') or '').strip()
    if len(title)>255:raise ValueError('Название может содержать до 255 символов.')
    if not title: raise ValueError('Укажите название результата.')
    result = {'type': kind, 'title': title}
    if kind == 'report':
        result['text'] = str(spec.get('text') or '')
        if len(result['text'])>40000:raise ValueError('Отчёт может содержать до 40 000 символов.')
    elif kind == 'slides':
        slides = spec.get('slides')
        if not isinstance(slides, list) or not 1 <= len(slides) <= 30: raise ValueError('Нужно от 1 до 30 слайдов.')
        if any(isinstance(s,dict) and (len(str(s.get('title') or ''))>200 or len(str(s.get('text') or ''))>1800 or str(s.get('text') or '').count('\n')>16) for s in slides):raise ValueError('Разделите длинный текст на несколько слайдов: до 1800 символов и 17 строк на слайд.')
        result['slides'] = [{'id': str(s.get('id') or f'slide-{index+1}')[:64], 'title': str(s.get('title') or '')[:200], 'text': str(s.get('text') or '')[:1800]} for index,s in enumerate(slides) if isinstance(s, dict)]
        if len(result['slides']) != len(slides) or len({s['id'] for s in result['slides']}) != len(slides): raise ValueError('Нужны уникальные слайды.')
    else:
        columns, rows = spec.get('columns'), spec.get('rows')
        if not isinstance(columns, list) or not 1 <= len(columns) <= 12: raise ValueError('Нужно от 1 до 12 столбцов.')
        if not isinstance(rows, list) or len(rows) > 300: raise ValueError('Доступно до 300 строк.')
        result['columns'] = [{'id': str(c.get('id') or '')[:64], 'label': str(c.get('label') or '')[:100]} for c in columns if isinstance(c, dict)]
        ids = [c['id'] for c in result['columns']]
        if len(ids) != len(columns) or not all(ids) or len(set(ids)) != len(ids): raise ValueError('Нужны уникальные столбцы.')
        result['rows'] = []
        for index,row in enumerate(rows):
            if not isinstance(row, dict) or not isinstance(row.get('cells'), dict): raise ValueError('Нужны строки с cells.')
            if any(len(str(value))>1500 for value in row['cells'].values()):raise ValueError('Ячейка может содержать до 1500 символов.')
            result['rows'].append({'id': str(row.get('id') or f'row-{index+1}')[:64], 'cells': {key: str(row['cells'].get(key) if row['cells'].get(key) is not None else '')[:1500] for key in ids},
                                   'state': row.get('state') if row.get('state') in {'pending', 'checked', 'missing', 'error'} else 'pending'})
        if len({r['id'] for r in result['rows']}) != len(rows): raise ValueError('Нужны уникальные строки.')
        if spec.get('workflow') == 'procurement': result['workflow'] = 'procurement'
    if len(json.dumps(result, ensure_ascii=False)) > 59000: raise ValueError('Результат слишком большой. Разделите его на части.')
    return result


def save(uid, pid, spec, eid=None, revision=None, *, model=False):
    value = validate(spec)
    # An edited or model-authored table is never automatically "verified".
    # Keep the workflow state as explicit bookkeeping, not a fact-check badge.
    if model and value['type'] == 'table':
        from app.services.run_protocol import current_run
        record = current_run.get()
        from app.services.run_protocol import public_url
        previous_rows={}
        if eid:
            previous=json.loads(memory.detail(uid,pid,eid)['body'])
            previous_rows={row['id']:row for row in previous.get('rows',[])}
        for row in value['rows']:
            if previous_rows.get(row['id'])==row:continue
            urls = [public_url(cell) for cell in row['cells'].values() if cell.startswith(('http://', 'https://'))]
            if row['state'] == 'checked' and not (record and any(s['url'] in urls and s['status'] == 'read' for s in record.sources)):
                row['state'] = 'pending'
    if eid and memory.detail(uid, pid, eid)['kind'] != 'result': raise ValueError('Результат не найден.')
    saved=memory.save(uid, pid, value['title'], json.dumps(value, ensure_ascii=False), 'result', entry_id=eid, expected_revision=revision,deduplicate=bool(model and not eid))
    if model and turn_results.get() is not None:turn_results.get().append(saved['id'])
    return saved


def read(uid, pid, eid, revision=None):
    item = memory.detail(uid, pid, eid, revision)
    if item['kind'] != 'result': raise ValueError('Результат не найден.')
    return {**item, 'spec': json.loads(item['body'])}


def changed_spec(spec, changes):
    import copy
    spec = copy.deepcopy(spec)
    if not isinstance(changes,list) or not 1<=len(changes)<=100:raise ValueError('Нужно 1–100 изменений.')
    for change in changes:
        if not isinstance(change,dict):raise ValueError('Изменение должно быть объектом.')
        if spec['type'] == 'report' and 'start' in change:
            start, end, old = change.get('start'), change.get('end'), change.get('old')
            if type(start) is not int or type(end) is not int or not 0 <= start < end <= len(spec['text']) or spec['text'][start:end] != old:
                raise ValueError('Выделенный фрагмент изменился. Выделите его заново.')
            spec['text'] = spec['text'][:start] + str(change.get('value') or '') + spec['text'][end:]
        elif change.get('field')=='title' and not change.get('slide_id'):
            spec['title']=str(change.get('value') or '')
        elif spec['type']=='table':
            row=next((row for row in spec['rows'] if row['id']==change.get('row_id')),None)
            if not row or change.get('column_id') not in {c['id'] for c in spec['columns']}:raise ValueError('Строка или столбец не найдены.')
            row['cells'][change['column_id']]=str(change.get('value') if change.get('value') is not None else '')
            row['state']='pending'
        elif spec['type']=='slides':
            slide=next((s for s in spec['slides'] if s['id']==change.get('slide_id')),None)
            if not slide or change.get('field') not in {'title','text'}:raise ValueError('Слайд или поле не найдены.')
            slide[change['field']]=str(change.get('value') or '')
        elif spec['type']=='report':
            old=str(change.get('old') or '')
            if not old or spec['text'].count(old)!=1:raise ValueError('Нужен единственный точный фрагмент отчёта для замены.')
            spec['text']=spec['text'].replace(old,str(change.get('value') or ''),1)
        else:raise ValueError('Недопустимая правка.')
    return validate(spec)


def patch(uid,pid,eid,revision,changes,*,model=False):
    if not isinstance(revision,int) or revision<1:raise ValueError('Нужна expected_revision из project_result_read.')
    spec=read(uid,pid,eid)['spec']
    return save(uid,pid,changed_spec(spec, changes),eid,revision,model=model)



def tool_read(uid,pid,eid,offset=0):
    item=read(uid,pid,eid);spec=item['spec'];offset=max(0,int(offset))
    key='rows' if spec['type']=='table' else 'slides' if spec['type']=='slides' else 'text'
    original=spec[key];count=4 if isinstance(original,list) else 6000
    spec[key]=original[offset:offset+count]
    shortened=[]
    if key=='rows':
        for row in spec['rows']:
            for column,value in row['cells'].items():
                if len(value)>200:row['cells'][column]=value[:200];shortened.append({'row_id':row['id'],'column_id':column})
    return {'id':eid,'revision':item['revision'],'spec':spec,'total':len(original),'next_offset':offset+count if offset+count<len(original) else None,
            'shortened_cells':shortened,'instruction':'Для точечной правки используй project_result_patch; не пересохраняй обрезанный spec как полный результат.'}


def export(uid, pid, eid, format):
    spec = read(uid, pid, eid)['spec']; stream = io.BytesIO()
    if format == 'xlsx' and spec['type'] == 'table':
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        workbook = Workbook(); sheet = workbook.active; sheet.title = 'Результат'
        sheet.append([c['label'] for c in spec['columns']])
        for row in spec['rows']: sheet.append([row['cells'][c['id']] for c in spec['columns']])
        for row in sheet:
            for cell in row:
                cell.data_type = 's'; cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in sheet[1]: cell.fill=PatternFill('solid',fgColor='16343D'); cell.font=Font(color='FFFFFF',bold=True)
        for n in range(1,len(spec['columns'])+1): sheet.column_dimensions[get_column_letter(n)].width=28
        sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
        workbook.save(stream)
        mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    elif format == 'docx' and spec['type'] == 'report':
        from docx import Document
        doc=Document(); doc.add_heading(spec['title'],0)
        for paragraph in spec['text'].split('\n'): doc.add_paragraph(paragraph)
        doc.save(stream); mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    elif format == 'pptx' and spec['type'] == 'slides':
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        deck=Presentation();deck.slide_width=Inches(13.333);deck.slide_height=Inches(7.5)
        for item in spec['slides']:
            slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string('0C1012')
            for top,height,text,size,color in [(0.7,1.5,item['title'],30,'F2F4F5'),(2.5,4.2,item['text'],19,'C5CBCE')]:
                while size>11 and sum(max(1,math.ceil(len(line)/max(1,int(11.7*72/(size*.62))))) for line in text.split('\n'))*size*1.5>height*72-14:size-=1
                frame=slide.shapes.add_textbox(Inches(.8),Inches(top),Inches(11.7),Inches(height)).text_frame;frame.word_wrap=True
                frame.text=text
                for para in frame.paragraphs:para.font.size=Pt(size);para.font.name='Arial';para.font.color.rgb=RGBColor.from_string(color)
        deck.save(stream);mime='application/vnd.openxmlformats-officedocument.presentationml.presentation'
    else: raise ValueError('Этот формат не подходит результату.')
    return stream.getvalue(),mime


def purchase_spec(title, positions):
    if not isinstance(positions, list) or not 1 <= len(positions) <= 20: raise ValueError('Пилот поддерживает 1–20 позиций за план.')
    if any(not isinstance(p,dict) or not str(p.get('part') or '').strip() for p in positions):raise ValueError('У каждой позиции должен быть компонент или артикул.')
    cols=[('part','Компонент / артикул'),('quantity','Нужно'),('shop','Магазин'),('offer','Предложение'),('price','Цена и упаковка'),('available','Наличие'),('delivery','Доставка'),('source','Источник'),('checked','Проверено'),('note','Условия / пробелы')]
    return {'type':'table','title':title,'workflow':'procurement','columns':[{'id':key,'label':label} for key,label in cols],
            'rows':[{'cells':{'part':str(p.get('part') or ''),'quantity':str(p.get('quantity') or ''),'note':str(p.get('requirements') or '')},'state':'pending'} for p in positions if isinstance(p,dict)]}


def purchase_offer(uid,pid,eid,revision,row_id,url,quotes):
    from app.services.run_protocol import current_run, public_url
    record=current_run.get();url=public_url(url)
    source=next((s for s in record.sources if s['url']==url and s['status']=='read'),None) if record else None
    if not source:raise ValueError('Сначала прочитай конкретную карточку. Поисковая выдача не подтверждает предложение.')
    spec=read(uid,pid,eid)['spec']
    if spec.get('workflow')!='procurement':raise ValueError('Это не план закупки.')
    row=next((r for r in spec['rows'] if r['id']==row_id),None)
    if not row:raise ValueError('Позиция не найдена.')
    text=record.page_texts.get(url,'');normalized=' '.join(text.split()).casefold()
    if not isinstance(quotes,dict):raise ValueError('Нужны цитаты из карточки.')
    updates={}
    for field in ['offer','price','available','delivery','note']:
        quote=str(quotes.get(field) or '').strip()
        if quote and (len(quote)>600 or ' '.join(quote.split()).casefold() not in normalized):raise ValueError('Цитата не найдена в прочитанной странице: '+field)
        if quote:updates[field]=quote
    if not updates.get('price') or not updates.get('available'):raise ValueError('Для подтверждённого предложения нужны реальные цитаты о цене и наличии.')
    identifiers=re.findall(r'\b[a-z][a-z0-9-]*\d[a-z0-9-]*\b',row['cells'].get('part','').casefold())
    if any(identifier not in normalized for identifier in identifiers):raise ValueError('Точный артикул позиции не найден в карточке. Проверь товар или явно согласуй аналог.')
    from urllib.parse import urlsplit
    if updates.get('note'):updates['note']=row['cells'].get('note','')+'\n'+updates['note']
    row['cells'].update({**updates,'shop':urlsplit(url).hostname,'source':url,'checked':source['read_at']})
    row['state']='missing' if re.search(r'нет\s+в\s+налич|недоступ|законч',updates['available'],re.I) else 'checked'
    return save(uid,pid,spec,eid,revision,model=True)


def capture_answer(uid,cid,prompt,answer):
    """Store ordinary final tables/reports without another model request."""
    if not cid or not re.search(r'сравн|таблиц|отч[её]т|доклад|план\s+закуп',prompt,re.I):return None
    pid=memory.attached(uid,cid)['project_id']
    if not pid:return None
    lines=answer.splitlines();columns=None;rows=[]
    for index,line in enumerate(lines[:-1]):
        if '|' in line and re.match(r'^\s*\|?\s*:?-{3}',lines[index+1]):
            headers=[c.strip().strip('* ') for c in line.strip().strip('|').split('|')]
            if not 1<=len(headers)<=12:continue
            columns=[{'id':f'c{i}','label':h[:100]} for i,h in enumerate(headers)]
            for next_line in lines[index+2:]:
                if '|' not in next_line:break
                cells=[c.strip() for c in next_line.strip().strip('|').split('|')]
                if len(cells)!=len(columns):break
                rows.append({'cells':{c['id']:value for c,value in zip(columns,cells)},'state':'pending'})
                if len(rows)==300:break
            if rows:break
    brief=re.sub(r'https?://\S+','',prompt.splitlines()[0]).strip()[:100]
    if columns and rows:spec={'type':'table','title':'Сравнение'+(' · '+brief if brief else ''),'columns':columns,'rows':rows}
    elif re.search(r'отч[её]т|доклад',prompt,re.I) and 200<=len(answer)<=40000:spec={'type':'report','title':'Отчёт · '+prompt.splitlines()[0][:100],'text':answer}
    else:return None
    return save(uid,pid,spec,model=False)


def answer_references(uid, cid, entry_ids):
    """Only reference actual owned results from this turn and the bound project."""
    if not cid or not entry_ids:return []
    pid = memory.attached(uid,cid)['project_id']
    if not pid:return []
    refs=[]
    for eid in dict.fromkeys(entry_ids):
        try:item=read(uid,pid,eid)
        except ValueError:continue
        refs.append({'project_id':pid,'entry_id':eid,'title':item['title'],
            'type':item['spec']['type'],'revision':item['revision']})
        if len(refs)>=10:break
    return refs
