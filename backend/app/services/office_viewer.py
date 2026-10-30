"""Read-only spreadsheet viewing without imposing a printed page layout."""
import hashlib
import io
import zipfile
from pathlib import Path
from urllib.parse import quote

from app.config import get_settings


def spreadsheet_viewer(data: bytes, filename: str, conversation_id: str, mobile: bool = False):
    from app.api.editor import _ds_token, sign_file_link
    from app.services.document_preview import MAX_BYTES

    settings = get_settings()
    suffix = Path(filename).suffix.lower()
    if suffix not in {'.xlsx', '.xls', '.ods'} or not settings.ONLYOFFICE_JWT_SECRET:
        return None
    if len(data) > MAX_BYTES:
        raise ValueError('Для предпросмотра поддерживаются файлы до 30 МБ. Скачайте этот файл.')
    digest = hashlib.sha256(conversation_id.encode() + b'\0' + filename.encode() + b'\0' + data).hexdigest()
    url = (f'{settings.BACKEND_INTERNAL_URL.rstrip("/")}/conversations/'
           f'{quote(conversation_id, safe="")}/files/viewer-source?t={sign_file_link(conversation_id, filename)}'
           f'&v={hashlib.sha256(data).hexdigest()}')
    config = {
        'documentType': 'cell',
        'type': 'mobile' if mobile else 'desktop',
        'document': {
            'fileType': suffix[1:], 'key': f'preview-v2-{digest}', 'title': filename, 'url': url,
            'permissions': {'edit': False, 'comment': False, 'review': False, 'download': False, 'print': False},
        },
        'editorConfig': {
            'mode': 'view', 'lang': 'ru', 'user': {'id': 'preview', 'name': 'Просмотр'},
            'customization': {'compactHeader': True, 'compactToolbar': True, 'uiTheme': 'theme-light',
                              'help': False, 'feedback': False, 'macrosMode': 'disable'},
        },
    }
    return {'server': settings.ONLYOFFICE_PUBLIC_URL.rstrip('/'), 'config': {**config, 'token': _ds_token(config)}}


def spreadsheet_source(data: bytes, filename: str) -> bytes:
    """Populate optional chart caches in a viewing copy; never resave the workbook."""
    if not filename.lower().endswith('.xlsx'):
        return data
    from lxml import etree as xml
    from openpyxl import load_workbook
    from openpyxl.utils.cell import range_to_tuple

    ns = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
    prefix = '{' + ns['c'] + '}'
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as source:
            entries = source.infolist()
            if len(entries) > 4096 or sum(e.file_size for e in entries) > 120 * 1024 * 1024:
                raise ValueError('Таблица слишком велика для предпросмотра. Скачайте файл.')
            charts = [e for e in entries if e.filename.startswith('xl/charts/chart') and e.filename.endswith('.xml')]
            if not charts or len(charts) > 64:
                return data
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            changed = {}
            try:
                for entry in charts:
                    root = xml.fromstring(source.read(entry), parser=xml.XMLParser(resolve_entities=False, no_network=True))
                    # Some generators leave both axes on the left for line/area charts.
                    # Excel infers the category axis; the viewer needs its position explicit.
                    for plot in root.findall('.//c:plotArea', ns):
                        if plot.find('c:lineChart', ns) is not None or plot.find('c:areaChart', ns) is not None:
                            for axis in plot.findall('c:catAx/c:axPos', ns):
                                if axis.get('val') == 'l':
                                    axis.set('val', 'b')
                    for ref in root.findall('.//c:numRef', ns) + root.findall('.//c:strRef', ns):
                        formula = ref.find(prefix + 'f')
                        if formula is None or not formula.text or '[' in formula.text:
                            continue
                        try:
                            sheet, (left, top, right, bottom) = range_to_tuple(formula.text)
                        except ValueError:
                            continue
                        if sheet not in workbook or (right-left+1)*(bottom-top+1) > 10_000:
                            continue
                        values = [cell.value for row in workbook[sheet].iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right) for cell in row]
                        if not values or any(v is None for v in values):
                            continue  # Preserve existing caches for uncalculated formulas/sparse series.
                        strings = ref.getparent().tag == prefix + 'cat' and any(isinstance(v, str) for v in values)
                        if strings:
                            ref.tag = prefix + 'strRef'
                            for cached in list(ref):
                                if cached.tag in {prefix+'numCache', prefix+'strCache'}:
                                    ref.remove(cached)
                        cache_tag = prefix + ('strCache' if ref.tag == prefix+'strRef' else 'numCache')
                        if ref.find(cache_tag) is not None:
                            continue
                        if cache_tag == prefix+'numCache' and not all(isinstance(v, (int, float)) for v in values):
                            continue
                        cache = xml.SubElement(ref, cache_tag)
                        xml.SubElement(cache, prefix+'ptCount', val=str(len(values)))
                        for index, value in enumerate(values):
                            point = xml.SubElement(cache, prefix+'pt', idx=str(index))
                            xml.SubElement(point, prefix+'v').text = str(value)
                    changed[entry.filename] = xml.tostring(root)
            finally:
                workbook.close()
            output = io.BytesIO()
            with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as target:
                for entry in entries:
                    target.writestr(entry, changed.get(entry.filename, source.read(entry)))
            return output.getvalue()
    except (zipfile.BadZipFile, KeyError, OSError, xml.XMLSyntaxError) as exc:
        raise ValueError('Не удалось прочитать таблицу. Попробуйте скачать файл.') from exc
