"""Правка одного абзаца .docx как рецензирование: старый текст в w:del, новый в w:ins.

Нужна для телефона: мобильный редактор OnlyOffice в бесплатной версии только
показывает документ. Правка видна на компьютере как «было / стало», её можно
принять или отклонить, оформление абзаца (стиль, нумерация, шрифт) остаётся.
"""

from __future__ import annotations

import copy
import io
from datetime import datetime, timezone

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

W_P = qn("w:p")
W_R = qn("w:r")
W_T = qn("w:t")
_REVISIONS = {qn("w:ins"), qn("w:del"), qn("w:moveFrom"), qn("w:moveTo")}
# Stay in place: markers that are not text.
_MARKERS = {
    qn("w:bookmarkStart"), qn("w:bookmarkEnd"), qn("w:proofErr"), qn("w:permStart"), qn("w:permEnd"),
    qn("w:commentRangeStart"), qn("w:commentRangeEnd"),
}
# Runs with these inside are objects, fields or notes – not plain text.
_COMPLEX_RUN = {
    qn("w:drawing"), qn("w:object"), qn("w:pict"), qn("w:fldChar"), qn("w:instrText"),
    qn("w:footnoteReference"), qn("w:endnoteReference"), qn("w:sym"), qn("w:ruby"),
}


class EditRejected(Exception):
    """Пункт нельзя поправить с телефона; текст – для человека."""


class EditConflict(Exception):
    """Пункт успели изменить: список на экране устарел."""


def _text(node) -> str:
    parts = []
    for child in node:
        tag = child.tag
        if tag in (qn("w:del"), qn("w:moveFrom"), qn("w:pPr"), qn("w:rPr")):
            continue
        if tag == W_T:
            parts.append(child.text or "")
        elif tag == qn("w:tab"):
            parts.append("\t")
        elif tag in (qn("w:br"), qn("w:cr")):
            parts.append("\n")
        else:
            parts.append(_text(child))
    return "".join(parts)


def _style(paragraph) -> str:
    style = paragraph.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
    return style.get(qn("w:val"), "") if style is not None else ""


def _in_table(paragraph) -> bool:
    parent = paragraph.getparent()
    while parent is not None:
        if parent.tag == qn("w:tc"):
            return True
        parent = parent.getparent()
    return False


def list_paragraphs(data: bytes) -> list[dict]:
    """Непустые абзацы документа по порядку, с индексом для правки."""
    document = Document(io.BytesIO(data))
    items = []
    for index, paragraph in enumerate(document.element.body.iter(W_P)):
        text = _text(paragraph)
        if not text.strip():
            continue
        style = _style(paragraph).lower()
        items.append({
            "index": index,
            "text": text,
            "heading": "heading" in style or "заголов" in style or style.startswith("title"),
            "in_table": _in_table(paragraph),
        })
    return items


def _next_revision_id(document) -> int:
    ids = [0]
    for part in document.part.package.iter_parts():
        element = getattr(part, "element", None)
        if element is None:
            element = getattr(part, "_element", None)
        if element is None:
            continue
        for node in element.iter():
            value = node.get(qn("w:id"))
            if value and value.lstrip("-").isdigit():
                ids.append(int(value))
    return max(ids) + 1


def _revision(tag: str, rev_id: int, author: str, date: str):
    element = OxmlElement(tag)
    element.set(qn("w:id"), str(rev_id))
    element.set(qn("w:author"), author)
    element.set(qn("w:date"), date)
    return element


def _new_run(text: str, rpr):
    run = OxmlElement("w:r")
    if rpr is not None:
        run.append(copy.deepcopy(rpr))
    for line_no, line in enumerate(text.split("\n")):
        if line_no:
            run.append(OxmlElement("w:br"))
        for piece_no, piece in enumerate(line.split("\t")):
            if piece_no:
                run.append(OxmlElement("w:tab"))
            if piece:
                node = OxmlElement("w:t")
                node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                node.text = piece
                run.append(node)
    return run


def replace_paragraph(data: bytes, index: int, base_text: str, new_text: str, author: str) -> bytes:
    """Новый .docx: абзац index заменён на new_text как правка рецензирования.

    base_text – текст, который человек видел: если абзац с тех пор поменяли,
    правка не ложится поверх чужой (EditConflict).
    """
    document = Document(io.BytesIO(data))
    paragraphs = list(document.element.body.iter(W_P))
    if index < 0 or index >= len(paragraphs):
        raise EditConflict("Пункт не найден – обновите документ")
    paragraph = paragraphs[index]
    if _text(paragraph) != base_text:
        raise EditConflict("Этот пункт уже изменили – обновите документ")
    if new_text == base_text:
        return data

    content = [child for child in paragraph if child.tag != qn("w:pPr")]
    if any(child.tag in _REVISIONS for child in content):
        raise EditRejected("В пункте есть непринятые правки – примите или отклоните их на компьютере")
    if any(child.tag not in _MARKERS and child.tag != W_R for child in content):
        raise EditRejected("В пункте есть ссылки или поля – его лучше поправить с компьютера")
    text_runs = []
    for run in (child for child in content if child.tag == W_R):
        if any(node.tag in _COMPLEX_RUN for node in run):
            raise EditRejected("В пункте есть рисунок, поле или сноска – его лучше поправить с компьютера")
        if run.find(qn("w:commentReference")) is None:
            text_runs.append(run)

    date = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    rev_id = _next_revision_id(document)
    first_rpr = next((run.find(qn("w:rPr")) for run in text_runs if run.find(qn("w:rPr")) is not None), None)

    inserted = _revision("w:ins", rev_id + 1, author, date)
    inserted.append(_new_run(new_text, first_rpr))
    if text_runs:
        deleted = _revision("w:del", rev_id, author, date)
        text_runs[0].addprevious(deleted)
        for run in text_runs:
            for node in run.findall(W_T):
                node.tag = qn("w:delText")
            deleted.append(run)
        deleted.addnext(inserted)
    else:
        paragraph.append(inserted)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


if __name__ == "__main__":
    # Самопроверка: правка ложится как рецензирование и текст читается обратно.
    source = Document()
    source.add_paragraph("ДОГОВОР")
    source.add_paragraph("4.2. Оплата в течение 30 дней.").runs[0].bold = True
    raw = io.BytesIO()
    source.save(raw)
    items = list_paragraphs(raw.getvalue())
    target = next(item for item in items if item["text"].startswith("4.2"))
    edited = replace_paragraph(raw.getvalue(), target["index"], target["text"], "4.2. Оплата в течение 10 рабочих дней.", "Тест")
    after = next(item for item in list_paragraphs(edited) if item["index"] == target["index"])
    assert after["text"] == "4.2. Оплата в течение 10 рабочих дней.", after
    xml = Document(io.BytesIO(edited)).element.xml
    assert "w:del " in xml and "w:ins " in xml and "w:delText" in xml
    print("ok")
