"""
Дешёвый режим генерации больших документов (docgen).
Строит план документа, пишет разделы параллельно дешёвой моделью
с эскалацией по сложности, собирает единый .docx.
"""
import asyncio
import io
import json
import logging
import re
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Set, Tuple

from config import DEEPSEEK_API_KEY
from openai_client import get_chat_response

logger = logging.getLogger(__name__)

# Not unlimited by design: a planner that returns a runaway outline (tens of
# thousands of sections) would otherwise start a run lasting days with no way
# back. 5000 sections is ~7500 pages, which covers the 5-7 thousand page
# requirement with margin.
MAX_SECTIONS = 5000
# Batches (not a global semaphore) so previous_tail continuity and progress
# reporting stay simple; the trade-off is each batch waits for its slowest
# member. Acceptable at this concurrency — do not "fix" this into a semaphore
# without re-checking that trade-off still holds.
MAX_PARALLEL_SECTIONS = 12
CHUNK_TARGET_CHARS = 3000
TOP_K_CHUNKS = 6
MAX_CHUNK_CHARS_PER_SECTION = 12000
# First heading + start of each source file: enough to tell «ПЛК» from «насос»
# when filenames are forma1/данные2. No model call — already in RAM.
DIGEST_CHARS = 700

# Raising the per-file extraction caps (document_parser.MAX_*_EXTENDED) removes
# the only thing that used to keep the total corpus small by accident. Measured
# with the new caps: 50 files x 900 pages = 118 MB of text -> 57 500 chunks,
# 21.8s, 420MB RAM. That's fine; 200 such files would be ~1.7GB, which is not.
# This is the deliberate ceiling that replaces that accident.
MAX_SOURCE_CHARS_TOTAL = 150_000_000

# Measured on 149 files across 5 archives: unbounded previews + catalog reached
# ~300 000 chars (~120 000 tokens) in the planner call, which is supposed to be
# the cheap one, and grows linearly with file count. These two caps bound it.
PLANNER_PREVIEW_BUDGET = 60_000   # chars of preview bodies (headers extra)
# Previews are per archive, not per file, so these two give a real ceiling
# instead of a slower linear growth: at most PLANNER_PREVIEW_GROUPS previews no
# matter how many files those archives hold, archives beyond it listed by name
# only. Measured whole-block ceiling including headers and name lists: ~70k
# chars, flat from ~1000 files up to 5000 (it was 346k at 1000 before).
PLANNER_PREVIEW_GROUPS = 40
_NAMES_SHOWN_PER_GROUP = 8        # file names listed per archive before "… и ещё N"
PLANNER_CATALOG_CHUNKS = 400      # chunk titles shown, was an unbounded 2000

PLANNER_MODEL = "gpt-5.6-terra"
DEFAULT_WRITER_MODEL = "gpt-5.6-luna"
ESCALATED_WRITER_MODEL = "gpt-5.6-terra"

# Одним ответом план на тысячи разделов физически не помещается: элемент плана
# в JSON — это ~85 токенов (замерено), а потолок вывода планировщика 128 000
# токенов, то есть ~1500 разделов в идеале и заметно меньше на практике.
# Поэтому выше этого порога план строится в два уровня: сначала главы, затем
# каждая глава параллельно расписывается на разделы. Без этого запрос на 5000
# страниц упирался в оборванный JSON и молча откатывался на один раздел.
SINGLE_CALL_SECTION_LIMIT = 150
SECTIONS_PER_CHAPTER = 25
# Модель редко возвращает ровно столько глав, сколько попросили, а обычно
# меньше. Если держать 25 разделов на главу как жёсткий потолок, тридцать глав
# вместо ста тридцати дают 750 разделов вместо 3350 — заказ не выполнен. При
# нехватке глав разбор одной главы растягивается до этого числа: 100 разделов
# в JSON — это ~8500 токенов, что свободно помещается в ответ.
MAX_SECTIONS_PER_EXPANSION = 100
MAX_CHAPTERS = MAX_SECTIONS // SECTIONS_PER_CHAPTER  # 200 глав x 25 = MAX_SECTIONS
# Отдельный потолок на число файлов: 1000 документов — это 1000 «глав» плана,
# а не 1000×25 разделов одного тома. MAX_CHAPTERS резал бы такой заказ до 200.
MAX_DOCUMENTS = MAX_SECTIONS
# Сколько разделов просить на документ, когда объём не назван. ИТТ — это общие
# сведения, объект автоматизации, назначение, технические требования,
# комплектность, документация, приёмка, гарантии: восемь и есть типовой состав.
DEFAULT_SECTIONS_PER_DOCUMENT = 8
CHAPTER_CATALOG_CHUNKS = 80  # заголовков исходников в вызове на разбор одной главы

# Столько символов раздела приходится на страницу .docx: этим же числом
# считается «примерно N стр.» в подтверждении.
CHARS_PER_PAGE = 2000

# Модель пишет длиннее, чем предполагает план. Замерено живым прогоном на
# сервере: заказ на 6 страниц дал 11, то есть ~1.8 длины раздела. Оценка
# стоимости обещает «не больше», поэтому закладывает этот перелёт с запасом —
# иначе обещание про бюджет нарушается ровно на успешном прогоне.
SECTION_OVERSHOOT = 2

# Transient failures (429, provider timeout, dropped connection) are certain
# at thousands of calls. Bounded, no jitter/backoff growth: this is a bulk
# job, not a latency-sensitive path.
SECTION_ATTEMPTS = 3
SECTION_RETRY_DELAYS = (2, 8)  # seconds before attempts 2 and 3

_TOKEN_RE = re.compile(r"[a-zA-Zа-яА-ЯёЁ0-9]{3,}")
# «5000 страниц», «5 000 стр.», «5-7 тысяч страниц», «300 разделов». Единица
# измерения обязательна: голое число в запросе объёмом не является.
_SIZE_REQUEST_RE = re.compile(
    r"(\d[\d\s]{0,9}?)\s*(?:[-–—]\s*(\d[\d\s]{0,9}?)\s*)?"
    r"(тыс\.?|тысяч[иа]?|к)?\s*(страниц\w*|стр\.|раздел\w*)",
    re.IGNORECASE,
)
# «1000 документов», «10 документов». Не «документацию»: это один том, не N файлов.
_DOC_COUNT_RE = re.compile(
    r"(\d[\d\s]{0,9}?)\s*(?:[-–—]\s*(\d[\d\s]{0,9}?)\s*)?"
    r"(тыс\.?|тысяч[иа]?|к)?\s*документ(?:ов|а|ы)?\b",
    re.IGNORECASE,
)
# Заглушки, которыми клиенты моделей отвечают вместо содержимого, когда
# ответа не было: openai_client.py:573,668 и deepseek_client.py:265,345
# возвращают их как обычный текст, не поднимая исключения.
_NOT_A_SECTION = frozenset({"Нет ответа от модели"})

_PAGE_MARKER = re.compile(r"\n---\s*Страница\s+\d+\s*---\n")
_SECTION_FAILED_PREFIX = "[Не удалось сгенерировать раздел"
# Число в маркере обязательно: без него подстрока может встретиться в самом
# тексте документа и оболгать его как усечённый.
_PAGES_TRUNCATED_RE = re.compile(r"Прочитал первые \d+")

# Confirmation chip labels (Task 12). Phase two matches these exact strings in
# the clarify reply text to detect the user's choice — a literal retyped
# elsewhere instead of importing these constants would silently break that.
_CONFIRM_START = "Да, начинай"
_CONFIRM_CANCEL = "Нет, отменить"
# Третий вариант появляется только когда есть что заменять. «Да, начинай»
# означает буквально «начинай»: значения из примеров остаются как есть.
# Метки [УКАЗАТЬ: …] — отдельное решение: на 5000 страниц их двадцать пять
# штук, и человек, нажавший главную кнопку просто чтобы посмотреть результат,
# такого не заказывал.
_CONFIRM_START_PLACEHOLDERS = "Начинай, поставь метки"
_TEMPLATE_CONFIRMED = "Да"
_TEMPLATE_NONE = "Шаблона нет, всё это база знаний"
_TEMPLATE_WRONG = "Нет, шаблон другой файл"

# Task 15: requisites that repeat across the example set (decimal designations,
# organisation names, ФИО, dates, long document numbers) get flagged before
# generation so the user can say what they become in the new document instead
# of find-replacing a finished 5000-page .docx afterwards.
MAX_REPLACEMENT_CANDIDATES = 25

_DECIMAL_DESIGNATION_RE = re.compile(r"[А-ЯЁ]{4}\.\d{6}\.\d{3}")
_ORG_NAME_RE = re.compile(r"(?:[А-ЯЁ]{2,6}\s+)?«[^»\n]{2,80}»")
_FIO_RE = re.compile(r"[А-ЯЁ][а-яё]+\s+[А-ЯЁ]\.\s?[А-ЯЁ]\.")
_REQUISITE_DATE_RE = re.compile(r"\b\d{2}\.\d{2}\.\d{4}\b")
_DOC_NUMBER_RE = re.compile(r"\b\d{6,}\b")
# Standards, not requisites — never proposed, even though a bare digit run
# inside "ГОСТ 2.105-95" could otherwise coincidentally match one of the
# patterns above.
_STANDARD_REF_PRECEDED_RE = re.compile(r"(?:ГОСТ|ОСТ|ТУ)\s*$")

# Order matters: earlier patterns claim their span first, so a decimal
# designation's own digits are not also proposed separately as a "document
# number" by the last, broadest pattern.
_CANDIDATE_PATTERNS: List[Tuple[re.Pattern, str]] = [
    (_DECIMAL_DESIGNATION_RE, "Децимальный номер"),
    (_ORG_NAME_RE, "Организация"),
    (_FIO_RE, "Разработал"),
    (_REQUISITE_DATE_RE, "Дата"),
    (_DOC_NUMBER_RE, "Номер документа"),
]

_REPLACEMENT_LINE_RE = re.compile(r"^(?P<left>.*?)(?:::=|→|->)(?P<right>.*)$")
_KEEP_AS_IS = ("как есть", "=")

# Words that appear in almost every GOST form and would otherwise glue
# unrelated files together when matching by digest.
_DIGEST_STOP = frozenset({
    "требования", "исходные", "технические", "документ", "раздел", "общие",
    "положения", "форма", "шаблон", "образец", "лист", "утверждаю",
    "согласовано", "таблица", "приложение", "гост", "заказчик", "разработчик",
    "наименование", "обозначение", "содержание", "введение", "данные",
    "файл", "архив", "страница", "the", "and", "for", "sheet", "page",
    "для", "или", "при", "как", "что", "этот", "этой", "этого",
})


@dataclass
class Chunk:
    id: int
    doc_name: str
    title: str
    text: str
    tokens: frozenset = field(default_factory=frozenset)
    title_tokens: frozenset = field(default_factory=frozenset)


@dataclass
class Section:
    id: int
    title: str
    brief: str
    complexity: str  # "simple" | "complex"
    chapter: str = ""  # заголовок главы при двухуровневом плане, иначе пусто
    document: str = ""  # название отдельного файла, если заказано несколько документов


def _tokenize(text: str) -> set:
    return {t.lower() for t in _TOKEN_RE.findall(text or "")}


def _requested_sections(text: str) -> Optional[int]:
    """Сколько разделов заказал пользователь, если он назвал объём.

    Планировщик сам по себе строит план «на глазок» — обычно это десятки
    разделов, то есть несколько десятков страниц. Человек, который просит
    документ на 5000 страниц, получал именно это и никак не узнавал, почему.
    Объём из запроса («5000 страниц», «5-7 тысяч страниц», «300 разделов»)
    становится целью для планировщика; из диапазона берётся верхняя граница.
    Возвращает None, если объём не назван — тогда поведение прежнее.
    """
    best = 0
    for m in _SIZE_REQUEST_RE.finditer(text or ""):
        low, high, scale, unit = m.group(1), m.group(2), m.group(3), m.group(4)
        multiplier = 1000 if scale else 1
        value = 0
        for raw in (low, high):
            if not raw:
                continue
            value = max(value, int(re.sub(r"\D", "", raw) or 0) * multiplier)
        if unit.lower().startswith("раздел"):
            sections = value
        else:
            sections = -(-value * CHARS_PER_PAGE // CHUNK_TARGET_CHARS)
        best = max(best, sections)
    if best <= 0:
        return None
    return min(best, MAX_SECTIONS)


def _requested_documents(text: str) -> Optional[int]:
    """Сколько отдельных файлов заказал пользователь.

    «5000 страниц» — объём одного документа. «1000 документов» — тысяча файлов.
    Без этого планировщик либо строил один том, либо упирался в MAX_CHAPTERS=200.
    """
    best = 0
    for m in _DOC_COUNT_RE.finditer(text or ""):
        low, high, scale = m.group(1), m.group(2), m.group(3)
        multiplier = 1000 if scale else 1
        value = 0
        for raw in (low, high):
            if not raw:
                continue
            value = max(value, int(re.sub(r"\D", "", raw) or 0) * multiplier)
        best = max(best, value)
    if best <= 0:
        return None
    return min(best, MAX_DOCUMENTS)


def _unique_doc_names(chunks: List[Chunk]) -> List[str]:
    seen: List[str] = []
    for chunk in chunks:
        if chunk.doc_name not in seen:
            seen.append(chunk.doc_name)
    return seen


def _file_stem(name: str) -> str:
    """Имя файла без пути и расширения: «архив.zip/ИТТ_ПЛК.docx» → «ИТТ ПЛК»."""
    base = (name or "").replace("\\", "/").rsplit("/", 1)[-1]
    base = re.sub(r"\.[^.]+$", "", base)
    return re.sub(r"[_\-]+", " ", base).strip()


def _name_score(query: str, candidate: str) -> int:
    """Насколько имя шаблона/исходника относится к названию итогового документа.

    Точное совпадение стемов, вхождение одной строки в другую, пересечение
    токенов. Ноль — нет связи: тогда этот файл не должен попасть в чужой документ.
    """
    q = _file_stem(query).lower()
    c = _file_stem(candidate).lower()
    if not q or not c:
        return 0
    if q == c:
        return 1000
    if q in c or c in q:
        return 100 + min(len(q), len(c))
    qt = _tokenize(q) | _tokenize(query.replace("/", " ").replace("\\", " "))
    ct = _tokenize(c) | _tokenize(candidate.replace("/", " ").replace("\\", " ").replace(".", " "))
    if not qt or not ct:
        return 0
    overlap = len(qt & ct)
    if not overlap:
        return 0
    return overlap * 10 + (3 if overlap == len(qt) else 0)


def _file_digests(chunks: List[Chunk]) -> Dict[str, str]:
    """Краткое содержание каждого исходника: заголовок и начало файла.

    Считается по уже прочитанным чанкам, без вызова модели — на тысяче файлов
    это дешевле любой суммаризации и достаточно, чтобы отличить ПЛК от насоса,
    когда файлы названы forma1.docx / данные2.xlsx.
    """
    parts: Dict[str, List[str]] = {}
    used: Dict[str, int] = {}
    for chunk in chunks:
        n = used.get(chunk.doc_name, 0)
        if n >= DIGEST_CHARS:
            continue
        bits = parts.setdefault(chunk.doc_name, [])
        if not bits and chunk.title:
            bits.append(chunk.title.strip())
            n += len(chunk.title)
        take = (chunk.text or "")[: max(0, DIGEST_CHARS - n)]
        if take:
            bits.append(take)
            n += len(take)
        used[chunk.doc_name] = n
    return {name: " ".join(bits)[:DIGEST_CHARS] for name, bits in parts.items()}


def _common_digest_tokens(digests: Dict[str, str]) -> Set[str]:
    """Токены, которые есть у большинства файлов, плюс стоп-слова ГОСТ.

    Без этого «требования» и «раздел» склеивали бы все шаблоны между собой.
    """
    common = set(_DIGEST_STOP)
    if len(digests) < 3:
        return common
    df: Dict[str, int] = {}
    for text in digests.values():
        for token in set(_tokenize(text)):
            df[token] = df.get(token, 0) + 1
    threshold = max(3, (len(digests) + 1) // 2)
    common.update(token for token, count in df.items() if count >= threshold)
    return common


def _digest_snippet(text: str, limit: int = 400) -> str:
    return re.sub(r"\s+", " ", text or "").strip()[:limit]


def _match_score(
    query: str,
    candidate_name: str,
    candidate_digest: str = "",
    query_digest: str = "",
    common: Optional[Set[str]] = None,
) -> int:
    """Имя плюс краткое содержание. Имя побеждает при равенстве; содержание
    спасает, когда файлы названы forma1 / данные2, а тема видна только из текста.
    """
    score = _name_score(query, candidate_name)
    common = common if common is not None else set(_DIGEST_STOP)
    query_tokens = (_tokenize(query) | _tokenize(query_digest)) - common
    cand_tokens = (
        _tokenize(_file_stem(candidate_name)) | _tokenize(candidate_digest)
    ) - common
    overlap = len(query_tokens & cand_tokens)
    if overlap:
        score += overlap * 8
    body_q = _tokenize(query_digest) - common
    body_c = _tokenize(candidate_digest) - common
    body_overlap = len(body_q & body_c)
    if body_overlap:
        score += body_overlap * 5
    return score


def _title_texts_for_matching(
    outline: List[Section],
    digests: Dict[str, str],
    template_names: Optional[Set[str]],
) -> Dict[str, str]:
    """Текст, с которым сравниваем шаблон: название документа, брифы разделов
    и выжимка той базы знаний, которая уже совпала по имени."""
    buckets: Dict[str, List[str]] = {}
    for section in outline:
        if not section.document:
            continue
        bucket = buckets.setdefault(section.document, [])
        if section.title:
            bucket.append(section.title)
        if section.brief:
            bucket.append(section.brief)
    knowledge = [name for name in digests if name not in (template_names or set())]
    result: Dict[str, str] = {title: " ".join(parts) for title, parts in buckets.items()}
    for title, text in list(result.items()):
        extras = [
            digests[name] for name in knowledge
            if digests.get(name) and _name_score(title, name) > 0
        ]
        if extras:
            result[title] = text + " " + " ".join(extras)
    return result


def _heuristic_template_names(chunks: List[Chunk]) -> Set[str]:
    """Файлы, которые по пути/имени выглядят как шаблоны оформления.

    Нужно до вызова планировщика: чтобы из архива «шаблоны.zip» + «знания.zip»
    сразу понять, какие .docx — формы, а какие — база, и сопоставить их
    по именам без LLM, который тысячу названий всё равно обрежет.
    """
    hit: Set[str] = set()
    for name in _unique_doc_names(chunks):
        if _TEMPLATE_NAME_RE.search(name) or _TEMPLATE_NAME_RE.search(_folder_of(name)):
            hit.add(name)
    return hit


def _seed_document_titles(
    requested: int,
    knowledge_names: List[str],
    template_names: Set[str],
) -> Optional[List[Tuple[str, str]]]:
    """Если в архивах уже лежит по файлу на каждый итоговый документ — не
    просить планировщик выдумать тысячу названий: он потеряет большую часть.

    База знаний — обычная единица работы (каждый xlsx/docx → свой выходной
    файл). Столько же шаблонов — запасной вариант: папка ГОСТовских форм и
    одна общая база. Возвращает (название, исходный файл) — файл нужен, чтобы
    подтянуть его краткое содержание в бриф и в сопоставление шаблона.
    """
    if requested < 2:
        return None
    knowledge = [n for n in knowledge_names if n not in (template_names or set())]
    templates = [n for n in sorted(template_names or ()) if n.lower().endswith((".docx", ".doc"))]
    if len(knowledge) == requested:
        return [(_file_stem(n) or n, n) for n in knowledge]
    if len(templates) == requested:
        return [(_file_stem(n) or n, n) for n in templates]
    return None


def _assign_templates(
    titles: List[str],
    template_names: Set[str],
    digests: Optional[Dict[str, str]] = None,
    title_texts: Optional[Dict[str, str]] = None,
) -> Dict[str, str]:
    """Название итогового документа → файл-шаблон. Уникально, пока шаблонов хватает.

    Сначала имя файла, затем краткое содержание: «ИТТ на ПЛК» забирает
    «ИТТ_ПЛК.docx», а если шаблоны названы forma1/forma2 — тот, в чьём начале
    речь про ПЛК, а не про насос.
    """
    templates = [n for n in sorted(template_names or ()) if n.lower().endswith((".docx", ".doc"))]
    unique_titles: List[str] = []
    seen: Set[str] = set()
    for title in titles:
        key = title or ""
        if key not in seen:
            seen.add(key)
            unique_titles.append(key)
    if not templates or not unique_titles:
        return {}
    if len(templates) == 1:
        return {title: templates[0] for title in unique_titles}

    digests = digests or {}
    title_texts = title_texts or {}
    common = _common_digest_tokens(digests) if digests else set(_DIGEST_STOP)
    pairs: List[Tuple[int, str, str]] = []
    for title in unique_titles:
        extra = title_texts.get(title, "")
        for tmpl in templates:
            score = _match_score(title, tmpl, digests.get(tmpl, ""), extra, common)
            if score:
                pairs.append((score, title, tmpl))
    pairs.sort(key=lambda item: (item[0], item[1], item[2]), reverse=True)
    assigned: Dict[str, str] = {}
    used_tmpl: Set[str] = set()
    for _, title, tmpl in pairs:
        if title in assigned or tmpl in used_tmpl:
            continue
        assigned[title] = tmpl
        used_tmpl.add(tmpl)
    leftover_titles = [t for t in unique_titles if t not in assigned]
    leftover_tmpls = [m for m in templates if m not in used_tmpl]
    for title, tmpl in zip(leftover_titles, leftover_tmpls):
        assigned[title] = tmpl
        used_tmpl.add(tmpl)
    fallback = templates[0]
    for title in unique_titles:
        assigned.setdefault(title, fallback)
    return assigned


def _scope_chunks_for_section(
    section: Section,
    chunks: List[Chunk],
    template_names: Optional[Set[str]],
    assigned_templates: Optional[Dict[str, str]] = None,
    digests: Optional[Dict[str, str]] = None,
) -> Tuple[List[Chunk], Set[str]]:
    """Фрагменты и шаблон, которые относятся к этому разделу, а не ко всему заказу.

    Имя файла плюс краткое содержание: «данные1.xlsx» про ПЛК попадёт в документ
    про ПЛК, даже если в названии нет «ПЛК». Нет пересечения — оставляем всю
    базу (лучше общий контекст, чем пустой файл), но чужие шаблоны отсекаем.
    """
    templates = set(template_names or ())
    knowledge_names = {c.doc_name for c in chunks if c.doc_name not in templates}
    chosen = None
    if section.document and assigned_templates:
        chosen = assigned_templates.get(section.document)
    templates_for_section: Set[str] = {chosen} if chosen else templates

    related_knowledge = knowledge_names
    if section.document and knowledge_names:
        digests = digests or {}
        common = _common_digest_tokens(digests) if digests else set(_DIGEST_STOP)
        query = " ".join(
            part for part in (section.document, section.title, section.brief) if part
        )
        extra = (digests.get(chosen, "") if chosen else "") + " " + query
        scored = [
            (
                name,
                _match_score(query, name, digests.get(name, ""), extra, common),
            )
            for name in knowledge_names
        ]
        positive = [(name, score) for name, score in scored if score > 0]
        if positive:
            best = max(score for _, score in positive)
            # Shared boilerplate («уникальный маркер», «требования») даёт слабый
            # ненулевой счёт сразу нескольким файлам. Держим только тех, кто
            # рядом с лучшим совпадением, иначе чужая база просачивается в документ.
            # min(..., best) обязателен: при слабом лучшем совпадении (счёт 13
            # против порога 16) фильтр отсекал ВСЮ базу знаний, и раздел уходил
            # писаться по одному шаблону, без единого факта. Причём полное
            # отсутствие совпадения сохраняло базу целиком — то есть намёк на
            # тему делал результат хуже, чем его отсутствие.
            floor = min(max(best * 0.5, 16), best)
            related_knowledge = {name for name, score in positive if score >= floor}

    allow = related_knowledge | templates_for_section
    scoped = [c for c in chunks if c.doc_name in allow]
    return (scoped or chunks), templates_for_section


def _split_into_chunks(doc_name: str, text: str, start_id: int) -> List[Chunk]:
    pieces = [p.strip() for p in _PAGE_MARKER.split(text) if p.strip()]
    chunks: List[Chunk] = []
    cid = start_id
    for piece in pieces:
        for i in range(0, len(piece), CHUNK_TARGET_CHARS):
            slice_text = piece[i:i + CHUNK_TARGET_CHARS].strip()
            if not slice_text:
                continue
            title = slice_text.splitlines()[0][:120]
            chunks.append(Chunk(
                id=cid, doc_name=doc_name, title=title, text=slice_text,
                tokens=frozenset(_tokenize(slice_text)),
                title_tokens=frozenset(_tokenize(title)),
            ))
            cid += 1
    return chunks


def _extract_source_chunks(user_id: int) -> Tuple[List[Chunk], List[str], List[str]]:
    """Чанки всех документов беседы, имена тех, что прочитаны не целиком, и
    имена тех, что не поместились в общий потолок MAX_SOURCE_CHARS_TOTAL.

    Усечение ищется в целом тексте документа, до нарезки: маркер длиной в
    десятки символов легко попадает на границу чанков, и тогда ни в одном
    чанке целиком не находится.

    Про сами усечения: повышенные лимиты извлечения включаются, только если
    режим «Документы» уже выбран в момент загрузки файла. Если человек сначала
    приложил файлы и только потом переключил режим, большой исходник уже
    усечён, и без предупреждения он выглядит как полный документ.

    Про общий потолок: документ, который вместе с уже накопленными не
    поместится в MAX_SOURCE_CHARS_TOTAL, пропускается — но следующие, меньшие
    по размеру, всё ещё проверяются и могут поместиться. Имя пропущенного
    документа обязательно возвращается вызывающему: документ, молча
    исчезнувший из генерации на 5000 страниц, — ровно тот отказ, от которого
    этот режим и защищали все предыдущие лимиты.
    """
    from conversations import conversation_manager
    from document_parser import TEXT_TRUNCATED_NOTICE

    docs = conversation_manager.get_documents(int(user_id))
    chunks: List[Chunk] = []
    truncated: List[str] = []
    skipped: List[str] = []
    total_chars = 0
    for doc in docs:
        name = str(doc.get("filename") or "документ")
        body = str(doc.get("content") or "")
        if not body:
            continue
        if total_chars + len(body) > MAX_SOURCE_CHARS_TOTAL:
            skipped.append(name)
            continue
        total_chars += len(body)
        if TEXT_TRUNCATED_NOTICE.strip() in body or _PAGES_TRUNCATED_RE.search(body):
            truncated.append(name)
        chunks.extend(_split_into_chunks(name, body, len(chunks)))
    return chunks, truncated, skipped


def _select_relevant_chunks(
    section_title: str, section_brief: str, chunks: List[Chunk], top_k: int = TOP_K_CHUNKS
) -> List[Chunk]:
    if not chunks:
        return []
    query_tokens = _tokenize(f"{section_title} {section_brief}")
    if not query_tokens:
        return chunks[:top_k]
    scored = []
    for chunk in chunks:
        overlap = len(query_tokens & chunk.tokens)
        # title_tokens is precomputed by _split_into_chunks; a Chunk built by hand
        # elsewhere still gets the title bonus instead of silently losing it.
        title_tokens = chunk.title_tokens or _tokenize(chunk.title)
        title_overlap = len(query_tokens & title_tokens)
        score = overlap + title_overlap * 3
        if score > 0:
            scored.append((score, chunk))
    if not scored:
        return chunks[:top_k]
    scored.sort(key=lambda pair: pair[0], reverse=True)
    return [chunk for _, chunk in scored[:top_k]]


def _join_labeled(chunks: List[Chunk], budget: int) -> str:
    parts = []
    remaining = budget
    for chunk in chunks:
        if remaining <= 0:
            break
        piece = chunk.text[:remaining]
        parts.append(f"[{chunk.doc_name}] {piece}")
        remaining -= len(piece)
    return "\n\n".join(parts)


def _build_context_block(chunks: List[Chunk], budget: int, template_names: Optional[Set[str]] = None) -> str:
    if not chunks:
        return ""
    if not template_names:
        return _join_labeled(chunks, budget)
    template_chunks = [c for c in chunks if c.doc_name in template_names]
    knowledge_chunks = [c for c in chunks if c.doc_name not in template_names]
    if not knowledge_chunks:
        return f"Формат по шаблону:\n{_join_labeled(template_chunks, budget)}"
    if not template_chunks:
        return f"Факты из базы знаний:\n{_join_labeled(knowledge_chunks, budget)}"
    half = budget // 2
    template_block = _join_labeled(template_chunks, half)
    knowledge_block = _join_labeled(knowledge_chunks, budget - len(template_block))
    return f"Формат по шаблону:\n{template_block}\n\nФакты из базы знаний:\n{knowledge_block}"


# Отрицательный lookbehind обязателен: без него «форма» находится внутри
# «ин-форма-ции», и архив «База информации для наполнения.zip» целиком
# уезжает в шаблоны. База знаний тогда пуста, а разделы пишутся без фактов —
# проверено на реальных именах архивов заказчика. Множественное число
# («Формы», «Образцы») тоже должно опознаваться, поэтому окончания
# перечислены, а не отброшены.
_TEMPLATE_NAME_RE = re.compile(
    r"(?<![а-яёa-z])(шаблон|образ(?:ец|цы|цов|цам)|форм[аыу]?|бланк|пример|template|form|sample|shablon)",
    re.IGNORECASE,
)


def _classify_documents_hint(chunks: List[Chunk]) -> str:
    """Дешёвая подсказка планировщику по именам файлов —
    что похоже на шаблон оформления, а что на базу знаний. Планировщик
    вправе проигнорировать её, если пользователь в запросе явно называет
    другой файл шаблоном — это только эвристика на случай, когда
    пользователь не сказал прямо."""
    seen: List[str] = []
    for c in chunks:
        if c.doc_name not in seen:
            seen.append(c.doc_name)
    template_like = [n for n in seen if _TEMPLATE_NAME_RE.search(n)]
    if not template_like:
        return ""
    knowledge_like = [n for n in seen if n not in template_like]
    return (
        f"По именам файлов похоже на шаблон оформления: {', '.join(template_like)}. "
        f"Остальное похоже на базу знаний: {', '.join(knowledge_like) or '(нет других файлов)'}. "
        "Это только подсказка по имени файла — если пользователь в запросе прямо называет "
        "другой файл шаблоном, доверяй запросу, а не этой подсказке."
    )


def _folder_of(doc_name: str) -> str:
    """The containing folder — everything before the LAST '/', or the bare
    name when there is none. extract_zip_archive names nested archive members
    f"{archive}/{dir}/{inner}", so a file two folders deep groups with its own
    folder, not with the archive as a whole — that's what lets a planner tell
    "arhiv.zip/shablony" (templates) apart from "arhiv.zip/baza" (knowledge)."""
    return doc_name.rsplit("/", 1)[0] if "/" in doc_name else doc_name


def _group_by_folder(names: List[str]) -> Dict[str, List[str]]:
    groups: Dict[str, List[str]] = {}
    for name in names:
        groups.setdefault(_folder_of(name), []).append(name)
    return groups


def _document_previews(chunks: List[Chunk], budget_per_doc: int = 1500) -> str:
    """Одно превью на папку (архив или папку внутри него), а не на каждый
    файл, плюс список имён внутри.

    Превью на файл не ограничивало ничего: пол в 300 символов побеждал деление
    бюджета, и блок снова рос линейно — замерено 66 878 символов уже на 150
    файлах и 346 028 на 1000, то есть хуже исходной проблемы. Папок же в
    беседе единицы, поэтому бюджет на папку даёт настоящий потолок:
    PLANNER_PREVIEW_GROUPS × budget_per_doc, независимо от числа файлов.

    Файлы внутри одной папки обычно однотипны (формы, тома одного
    справочника), так что для распознавания шаблона хватает начала первого из
    них; остальные планировщик видит по именам — они дешёвые, и именно их он
    возвращает.
    """
    order: List[str] = []
    bodies: Dict[str, str] = {}
    for c in chunks:
        if c.doc_name not in bodies:
            order.append(c.doc_name)
            bodies[c.doc_name] = ""
        if len(bodies[c.doc_name]) < budget_per_doc:
            bodies[c.doc_name] += c.text

    groups = _group_by_folder(order)
    previewed = list(groups.items())[:PLANNER_PREVIEW_GROUPS]
    per_group_budget = max(300, min(budget_per_doc, PLANNER_PREVIEW_BUDGET // max(1, len(previewed))))

    def _listed(names: List[str]) -> str:
        shown = ", ".join(names[:_NAMES_SHOWN_PER_GROUP])
        if len(names) > _NAMES_SHOWN_PER_GROUP:
            shown += f", … и ещё {len(names) - _NAMES_SHOWN_PER_GROUP}"
        return shown

    parts: List[str] = []
    for folder, members in previewed:
        head = members[0]
        if len(members) > 1:
            parts.append(
                f"=== Группа {folder}: {len(members)} файлов ({_listed(members)}) ===\n"
                f"--- начало файла {head} ---\n{bodies[head][:per_group_budget]}"
            )
        else:
            parts.append(f"--- {head} (начало файла) ---\n{bodies[head][:per_group_budget]}")

    hidden = list(groups)[PLANNER_PREVIEW_GROUPS:]
    if hidden:
        parts.append(f"=== Ещё источники без превью ({len(hidden)}): {_listed(hidden)} ===")
    return "\n\n".join(parts)


def _parse_outline_response(response_text: str) -> Tuple[List[Section], List[str]]:
    cleaned = (response_text or "").strip()
    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```(?:json)?\n", "", cleaned)
        cleaned = re.sub(r"\n```$", "", cleaned)
    try:
        raw = json.loads(cleaned)
    except json.JSONDecodeError:
        return [], []
    if not isinstance(raw, dict):
        return [], []
    raw_sections = raw.get("sections")
    if not isinstance(raw_sections, list):
        return [], []
    sections: List[Section] = []
    for i, item in enumerate(raw_sections):
        if not isinstance(item, dict) or not item.get("title"):
            continue
        complexity = item.get("complexity") if item.get("complexity") in ("simple", "complex") else "simple"
        sections.append(Section(
            id=i,
            title=str(item["title"])[:200],
            brief=str(item.get("brief", ""))[:500],
            complexity=complexity,
            document=str(item.get("document", ""))[:200].strip(),
        ))
    raw_templates = raw.get("template_documents")
    # A planner that answers with the old singular shape (a bare string) must
    # not silently produce "no template" — wrap it instead of dropping it.
    if isinstance(raw_templates, str):
        raw_templates = [raw_templates]
    if not isinstance(raw_templates, list):
        raw_templates = []
    template_names = [str(n).strip() for n in raw_templates if str(n).strip()]
    return sections[:MAX_SECTIONS], template_names


def _expand_template_names(names: List[str], chunks: List[Chunk]) -> Set[str]:
    """Resolves planner-named templates to the set of document names they
    select. An exact match selects that document; a name matching an archive
    prefix (some document starts with f"{name}/") selects every member of
    that archive; unknown names are dropped. A name that is both a real file
    and a prefix of others selects the union of both — nothing named by the
    planner is silently lost."""
    all_docs: List[str] = []
    doc_set: Set[str] = set()
    for c in chunks:
        if c.doc_name not in doc_set:
            doc_set.add(c.doc_name)
            all_docs.append(c.doc_name)

    resolved: Set[str] = set()
    for raw in names:
        # Планировщик вполне может назвать архив с хвостовым слэшем
        # ("shablony.zip/"); без нормализации такое имя не совпадёт ни с чем,
        # и шаблоны молча окажутся пустыми.
        name = raw.strip().rstrip("/")
        if not name:
            continue
        if name in doc_set:
            resolved.add(name)
        prefix = f"{name}/"
        resolved.update(d for d in all_docs if d.startswith(prefix))
    return resolved


async def _plan_outline(
    user_text: str,
    chunks: List[Chunk],
    user_id: int,
    extra_instruction: str = "",
    want_count: Optional[int] = None,
    as_chapters: bool = False,
    as_documents: bool = False,
) -> Tuple[List[Section], Set[str], bool]:
    """Разделы, набор файлов-шаблонов (может быть пустым) и флаг «план
    построить не удалось».

    Флаг нужен, чтобы молчаливый откат на один раздел «Документ» был виден
    пользователю: без него сбой планировщика выглядит как обычный короткий
    документ, и человек не понимает, почему структура не та, что он просил.
    """
    if chunks:
        catalog = "\n".join(f"{c.id}: {c.title}" for c in chunks[:PLANNER_CATALOG_CHUNKS])
    else:
        catalog = "(исходники не найдены — опирайся только на запрос пользователя)"

    classification_hint = _classify_documents_hint(chunks)
    document_previews = _document_previews(chunks)

    if as_documents and as_chapters:
        size_instruction = (
            f"Это перечень отдельных документов: верни ровно {want_count} документов. "
            "Каждый элемент — отдельный файл, не глава одного тома. "
            "В title — название документа (оно же станет именем файла), в brief — что в нём должно быть. "
            "document заполни тем же названием, что title.\n"
        )
    elif as_documents:
        size_instruction = (
            f"Пользователю нужно ровно {want_count} отдельных документов — столько и верни. "
            "У каждого раздела заполни document названием файла, к которому он относится; "
            "разделы одного документа идут подряд.\n"
        )
    elif as_chapters:
        size_instruction = (
            f"Это оглавление ОЧЕНЬ большого документа: верни ровно {want_count} глав верхнего "
            f"уровня, каждую из которых потом распишут на {SECTIONS_PER_CHAPTER} разделов. "
            "В title — название главы, в brief — что в неё входит, чтобы разделы не пересекались "
            "между главами. Если нужны несколько отдельных документов — у каждой главы заполни "
            "document названием файла, к которому она относится: разделы главы унаследуют его.\n"
        )
    elif want_count:
        size_instruction = f"Пользователю нужен объём примерно в {want_count} раздел(ов) — столько и верни.\n"
    else:
        size_instruction = ""

    prompt = (
        "Построй план большого документа по запросу пользователя.\n"
        "ГЛАВНОЕ ПРАВИЛО: запрос пользователя — это техническое задание, а не пожелание. "
        "Если он назвал количество документов, их перечень или структуру — выполняй буквально, "
        "не добавляя своего и не сокращая.\n"
        + size_instruction
        + "Верни СТРОГО JSON-объект без markdown-обёрток и без пояснений, в формате:\n"
        '{"template_documents": ["точное имя файла или имя архива из списка ниже", ...] (можно пустой список), '
        '"sections": [{"title": "Название раздела", "brief": "Что должно быть в разделе, 1-3 предложения", '
        '"complexity": "simple" | "complex", "document": "название отдельного файла"}, ...]}.\n'
        "document — заполняй ТОЛЬКО если пользователь просит НЕСКОЛЬКО отдельных документов "
        "(например «сделай 10 ИТТ на разные узлы» или «1000 документов»). Тогда у каждого раздела "
        "стоит название того документа, к которому он относится, и разделы одного документа идут подряд. "
        "Если нужен один документ — оставь document пустым у всех разделов.\n"
        "Если шаблонов оформления несколько, каждому итоговому файлу соответствует ОДИН шаблон "
        "с похожим именем; факты бери только из исходников этой же темы, не смешивай базы разных узлов.\n"
        "complexity=complex — для расчётов, таблиц с цифрами, юридических формулировок, "
        "технических требований с точными значениями. Остальное — simple.\n\n"
        f"Запрос пользователя:\n{user_text}\n\n"
        + (f"Уточнения пользователя: {extra_instruction}\n\n" if extra_instruction else "")
        + f"Заголовки доступных фрагментов исходников (id: заголовок):\n{catalog}\n\n"
        "template_documents — точные имена документов ниже (можно несколько), ТОЛЬКО если запрос\n"
        "пользователя или подсказка ниже указывают, что это шаблон оформления (пример структуры для\n"
        "итогового документа). Вместо перечисления всех файлов можно указать имя целого архива, имя\n"
        "папки внутри архива (например, \"архив.zip/папка\") или имя одного файла — тогда шаблоном\n"
        "станут все файлы этого архива, все файлы этой папки или один этот файл.\n"
        "Слова пользователя в запросе всегда важнее подсказки по имени файла.\n"
        "Если ничто не указывает на шаблон, верни template_documents: []. Если template_documents\n"
        "задан, построй список разделов, максимально повторяя структуру (заголовки, их порядок) этих\n"
        "документов, а не придумывай новую; иначе строй план свободно по сути запроса и остальных\n"
        "материалов.\n\n"
        f"{classification_hint}\n\n"
        f"Начало каждого документа (для распознавания шаблона):\n{document_previews}"
    )
    messages = [
        {"role": "system", "content": "Ты планировщик документов. Отвечаешь только валидным JSON-объектом."},
        {"role": "user", "content": prompt},
    ]
    try:
        response_text, _, _, _ = await get_chat_response(
            messages, model=PLANNER_MODEL, user_id=user_id, use_tools=False, use_skills=False,
        )
        sections, template_names_raw = _parse_outline_response(response_text)
        if sections:
            return sections, _expand_template_names(template_names_raw, chunks), False
        logger.error("docgen outline planning returned no sections: %s", (response_text or "")[:300])
    except Exception as e:
        logger.error(f"docgen outline planning failed: {e}", exc_info=True)
    return [Section(id=0, title="Документ", brief=user_text[:500], complexity="complex")], set(), True


async def _expand_chapter(
    chapter: Section, user_text: str, catalog: str, per_chapter: int, user_id: int
) -> List[Section]:
    """Одна глава -> её разделы. Отдельный дешёвый вызов на главу: только так
    план на тысячи разделов вообще помещается в ответы модели.

    Вход намеренно маленький (запрос, глава и короткий каталог заголовков, без
    превью документов) — таких вызовов сотни, и полный контекст планировщика
    в каждом стоил бы дороже, чем сама генерация текста.

    Провал вызова не теряет главу: она возвращается одним разделом, то есть
    документ становится короче, но не дырявым.
    """
    prompt = (
        f"Распиши одну главу большого документа на разделы.\n"
        f"Верни СТРОГО JSON-объект без markdown-обёрток: "
        '{"sections": [{"title": "...", "brief": "Что должно быть в разделе, 1-3 предложения", '
        '"complexity": "simple" | "complex"}, ...]}.\n'
        f"Нужно примерно {per_chapter} раздел(ов), только по этой главе, без пересечений с другими.\n"
        "complexity=complex — для расчётов, таблиц с цифрами, юридических формулировок, "
        "технических требований с точными значениями. Остальное — simple.\n\n"
        f"Запрос пользователя:\n{user_text}\n\n"
        f"Глава: {chapter.title}\n"
        f"Что входит в главу: {chapter.brief}\n\n"
        f"Заголовки доступных фрагментов исходников:\n{catalog}"
    )
    messages = [
        {"role": "system", "content": "Ты планировщик документов. Отвечаешь только валидным JSON-объектом."},
        {"role": "user", "content": prompt},
    ]
    try:
        response_text, _, _, _ = await asyncio.wait_for(
            get_chat_response(
                messages, model=DEFAULT_WRITER_MODEL, user_id=user_id, use_tools=False, use_skills=False,
            ),
            timeout=120,
        )
        sections, _ = _parse_outline_response(response_text)
    except Exception as e:
        logger.warning("docgen chapter '%s' expansion failed: %s", chapter.title, e)
        sections = []
    if not sections:
        logger.warning("docgen chapter '%s' produced no sections, kept as one", chapter.title)
        return [Section(id=0, title=chapter.title, brief=chapter.brief,
                        complexity=chapter.complexity, chapter=chapter.title,
                        document=chapter.document)]
    inherited = sections[:per_chapter]
    for section in inherited:
        section.chapter = chapter.title
        # The chapter-level plan is where "10 separate documents of 5000 pages"
        # is decided; expansion JSON has no `document` field, so without this
        # stamp a two-level run silently collapses back into one file.
        section.document = chapter.document
    return inherited


async def _expand_chapters(
    chapters: List[Section],
    target_sections: int,
    user_text: str,
    chunks: List[Chunk],
    user_id: int,
    status_msg: Any,
    template_names: Set[str],
    chapter_cap: int,
) -> Tuple[List[Section], Set[str], bool]:
    """Расписывает список глав/документов на разделы. chapter_cap — сколько
    глав реально разбирать: для одного тома это MAX_CHAPTERS, для N файлов —
    MAX_DOCUMENTS, иначе заказ на 1000 документов обрезался бы до 200."""
    if not chapters:
        return chapters, template_names, True
    chapters = chapters[:chapter_cap]
    per_chapter = min(MAX_SECTIONS_PER_EXPANSION, max(1, -(-target_sections // len(chapters))))
    catalog = "\n".join(f"{c.id}: {c.title}" for c in chunks[:CHAPTER_CATALOG_CHUNKS])
    sections: List[Section] = []
    for batch_start in range(0, len(chapters), MAX_PARALLEL_SECTIONS):
        batch = chapters[batch_start:batch_start + MAX_PARALLEL_SECTIONS]
        results = await asyncio.gather(*[
            _expand_chapter(chapter, user_text, catalog, per_chapter, user_id) for chapter in batch
        ])
        for chapter_sections in results:
            sections.extend(chapter_sections)
        planned = min(batch_start + MAX_PARALLEL_SECTIONS, len(chapters))
        await _update_status(
            status_msg,
            f"📄 План: {_progress_bar(planned, len(chapters))} глава {planned} из {len(chapters)}, "
            f"разделов уже {len(sections)}",
        )
    for i, section in enumerate(sections):
        section.id = i
    logger.info(
        "docgen two-level plan: %d chapters -> %d sections (target %d)",
        len(chapters), len(sections), target_sections,
    )
    return sections[:MAX_SECTIONS], template_names, False


DOCUMENT_TOP_UP_ATTEMPTS = 2


async def _top_up_documents(
    chapters: List[Section],
    want: int,
    user_text: str,
    chunks: List[Chunk],
    user_id: int,
    extra_instruction: str,
) -> List[Section]:
    """Дозапрашивает названия документов, если список пришёл короче заказа.

    Замерено на реальном заказе «все 10 документов»: один и тот же запрос на
    одних и тех же архивах давал то 10 названий, то 2. Просьба «верни ровно 10»
    остаётся просьбой, поэтому недостача добирается отдельными вызовами — уже
    зная, что есть, модели остаётся придумать только остаток.
    """
    for _ in range(DOCUMENT_TOP_UP_ATTEMPTS):
        missing = want - len(chapters)
        if missing <= 0:
            return chapters
        have = "; ".join(c.title for c in chapters)
        logger.warning(
            "docgen: планировщик вернул %d документов из %d, дозапрашиваю %d",
            len(chapters), want, missing,
        )
        extra, _, failed = await _plan_outline(
            user_text, chunks, user_id,
            extra_instruction=(
                f"{extra_instruction}\nУже запланированы документы: {have}. "
                f"Верни ТОЛЬКО ещё {missing} других документов на другие узлы, "
                "не повторяя перечисленные."
            ),
            want_count=missing, as_chapters=True, as_documents=True,
        )
        if failed or not extra:
            break
        known = {c.title.strip().lower() for c in chapters}
        for section in extra:
            title = section.title.strip()
            if not title or title.lower() in known:
                continue
            known.add(title.lower())
            section.document = section.document or title
            chapters.append(section)
            if len(chapters) >= want:
                break
    if len(chapters) < want:
        logger.error(
            "docgen: заказано %d документов, в плане осталось %d", want, len(chapters)
        )
    return chapters


async def _plan_as_documents(
    user_text: str,
    chunks: List[Chunk],
    user_id: int,
    extra_instruction: str,
    status_msg: Any,
    documents_wanted: int,
    target_sections: Optional[int],
) -> Tuple[List[Section], Set[str], bool]:
    """План на N отдельных файлов: имена из архивов, если они уже лежат
    по файлу на документ, иначе — один вызов планировщика на перечень."""
    heuristic = _heuristic_template_names(chunks)
    knowledge_names = [n for n in _unique_doc_names(chunks) if n not in heuristic]
    seed = _seed_document_titles(documents_wanted, knowledge_names, heuristic)
    template_names: Set[str] = set(heuristic)
    digests = _file_digests(chunks)

    if seed:
        chapters = []
        for i, (title, source) in enumerate(seed):
            snippet = _digest_snippet(digests.get(source, ""))
            brief = (
                f"Документ «{title}». Кратко по исходнику: {snippet}"
                if snippet else
                f"Полный документ «{title}» по относящимся исходникам этой темы."
            )
            chapters.append(Section(
                id=i,
                title=title,
                brief=brief,
                complexity="complex",
                document=title,
            ))
    else:
        # Всегда сначала перечень документов, а не плоский список разделов.
        # На заказе «все 10 документов» плоский список отдавал число файлов на
        # усмотрение модели: она вернула разделы всего двух документов из
        # десяти. Перечислить десять названий — задача, которую модель
        # выполняет надёжно, а дальше каждый документ раскрывается отдельно.
        want = min(documents_wanted, MAX_DOCUMENTS)
        chapters, template_names, failed = await _plan_outline(
            user_text, chunks, user_id, extra_instruction,
            want_count=want, as_chapters=True, as_documents=True,
        )
        if failed or not chapters:
            return chapters, template_names, True
        chapters = await _top_up_documents(
            chapters, want, user_text, chunks, user_id, extra_instruction,
        )
        chapters = chapters[:want]
        if heuristic:
            template_names = template_names | heuristic
        for chapter in chapters:
            if not chapter.document:
                chapter.document = chapter.title

    # Без указанного объёма документ получал ОДИН раздел: в готовых ИТТ было
    # по три заголовка на файл. Документ такого рода — это несколько разделов
    # (общие сведения, назначение, требования, комплектность, приёмка), и
    # столько и запрашивается по умолчанию.
    if target_sections:
        per_doc = min(
            MAX_SECTIONS_PER_EXPANSION,
            max(1, -(-min(target_sections, MAX_SECTIONS) // max(1, len(chapters)))),
        )
    else:
        per_doc = DEFAULT_SECTIONS_PER_DOCUMENT
    if per_doc <= 1:
        for i, chapter in enumerate(chapters):
            chapter.id = i
        return chapters, template_names, False
    return await _expand_chapters(
        chapters, len(chapters) * per_doc, user_text, chunks, user_id,
        status_msg, template_names, chapter_cap=MAX_DOCUMENTS,
    )


async def _plan_document(
    user_text: str,
    chunks: List[Chunk],
    user_id: int,
    extra_instruction: str = "",
    status_msg: Any = None,
) -> Tuple[List[Section], Set[str], bool]:
    """План документа: один вызов на небольшой документ, два уровня на большой.

    Порог — SINGLE_CALL_SECTION_LIMIT: выше него один ответ планировщика
    физически не вмещает план (см. комментарий к константе), поэтому сначала
    строятся главы, а потом они параллельно расписываются на разделы.

    «N документов» — отдельная ветка: это N файлов, а не N разделов одного тома.
    """
    combined = f"{user_text}\n{extra_instruction}"
    documents_wanted = _requested_documents(combined)
    target = _requested_sections(combined)

    if documents_wanted and documents_wanted > 1:
        return await _plan_as_documents(
            user_text, chunks, user_id, extra_instruction, status_msg,
            documents_wanted, target,
        )

    if not target or target <= SINGLE_CALL_SECTION_LIMIT:
        return await _plan_outline(user_text, chunks, user_id, extra_instruction, want_count=target)

    chapter_count = min(MAX_CHAPTERS, -(-target // SECTIONS_PER_CHAPTER))
    chapters, template_names, planning_failed = await _plan_outline(
        user_text, chunks, user_id, extra_instruction,
        want_count=chapter_count, as_chapters=True,
    )
    if planning_failed or not chapters:
        return chapters, template_names, True

    # Планировщик может вернуть глав больше, чем просили. Без этого среза
    # каждая лишняя глава стоила бы отдельного вызова на разбор, а разделы
    # сверх лимита всё равно отрезаются в самом конце — деньги за них уже
    # были бы потрачены.
    return await _expand_chapters(
        chapters, target, user_text, chunks, user_id, status_msg, template_names,
        chapter_cap=chapter_count,
    )


def _find_replacement_candidates(chunks: List[Chunk]) -> List[Tuple[str, str]]:
    """Requisites that repeat across the example set (Task 15): decimal
    designations, organisation names in guillemets, ФИО, dates and long
    document/contract numbers. Kept only if seen in at least two different
    documents (chunk.doc_name) — a requisite repeats across the example set, a
    one-off number in a single table does not, and that frequency rule is what
    keeps genuine technical data off the list. ГОСТ/ОСТ/ТУ references are
    excluded outright: they are standards, not project requisites.

    Pure regex + frequency over text already in memory — no model call here,
    so it is safe to run in the same thread as the corpus load. Returns
    (value, generic_label) pairs, ranked by document count, capped at
    MAX_REPLACEMENT_CANDIDATES.
    """
    occurrences: Dict[str, Set[str]] = {}
    generic_labels: Dict[str, str] = {}
    for chunk in chunks:
        text = chunk.text
        claimed: List[Tuple[int, int]] = []
        for pattern, generic_label in _CANDIDATE_PATTERNS:
            for m in pattern.finditer(text):
                start, end = m.span()
                if any(start < c_end and end > c_start for c_start, c_end in claimed):
                    continue
                if _STANDARD_REF_PRECEDED_RE.search(text[max(0, start - 15):start]):
                    continue
                value = m.group(0).strip()
                if not value:
                    continue
                claimed.append((start, end))
                occurrences.setdefault(value, set()).add(chunk.doc_name)
                generic_labels.setdefault(value, generic_label)

    repeated = [(value, docs) for value, docs in occurrences.items() if len(docs) >= 2]
    repeated.sort(key=lambda pair: len(pair[1]), reverse=True)
    return [(value, generic_labels[value]) for value, _ in repeated[:MAX_REPLACEMENT_CANDIDATES]]


def _parse_candidate_labels(response_text: str, known_values: Set[str]) -> List[Tuple[str, str]]:
    cleaned = (response_text or "").strip()
    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```(?:json)?\n", "", cleaned)
        cleaned = re.sub(r"\n```$", "", cleaned)
    try:
        raw = json.loads(cleaned)
    except json.JSONDecodeError:
        return []
    if not isinstance(raw, dict):
        return []
    items = raw.get("items")
    if not isinstance(items, list):
        return []
    result: List[Tuple[str, str]] = []
    seen: Set[str] = set()
    for item in items:
        if not isinstance(item, dict):
            continue
        value = str(item.get("value") or "").strip()
        label = str(item.get("label") or "").strip()[:60]
        if not value or value not in known_values or not label or value in seen:
            continue
        seen.add(value)
        result.append((value, label))
    return result


async def _label_replacement_candidates(
    candidates: List[Tuple[str, str]], user_id: int
) -> List[Tuple[str, str]]:
    """One cheap labelling call on DEFAULT_WRITER_MODEL (not the planner): it
    sees only the candidate values — never the corpus — and decides which are
    genuinely project-specific requisites (as opposed to real technical data
    that happens to repeat), giving each a short Russian label. A candidate it
    does not return is dropped.

    Guarded exactly like _parse_outline_response: no candidates skips the
    call entirely, and any failure or unparseable answer falls back to the
    regex candidates with their generic labels — never an exception.
    """
    if not candidates:
        return []
    values = [value for value, _ in candidates]
    prompt = (
        "Вот значения, повторяющиеся в примерах документов. Выбери из них те, что являются "
        "специфичными для проекта реквизитами (заказчик, разработчик, децимальный номер, "
        "название организации, ФИО, дата и т.п.) и поменяются в новом документе — а не "
        "техническими данными вроде характеристик изделия. Верни СТРОГО JSON-объект без "
        'markdown-обёрток: {"items": [{"value": "точное значение из списка", '
        '"label": "короткая русская подпись"}, ...]}.\n\n'
        f"Значения:\n{chr(10).join(values)}"
    )
    messages = [
        {"role": "system", "content": "Ты помощник по разметке реквизитов. Отвечаешь только валидным JSON-объектом."},
        {"role": "user", "content": prompt},
    ]
    try:
        response_text, _, _, _ = await get_chat_response(
            messages, model=DEFAULT_WRITER_MODEL, user_id=user_id, use_tools=False, use_skills=False,
        )
        labelled = _parse_candidate_labels(response_text, set(values))
        if labelled:
            return labelled
        logger.error("docgen candidate labelling returned nothing usable: %s", (response_text or "")[:300])
    except Exception as e:
        logger.error(f"docgen candidate labelling failed: {e}", exc_info=True)
    return candidates


def _replacement_table(text: str) -> Tuple[Dict[str, str], bool]:
    """Parses a filled «было → стало» reply (Task 15).

    Returns (mapping, is_table). mapping is {было: стало}; an empty справа
    value maps to "" (placeholder later), and "как есть" / "=" drops that
    pair so the old value is kept. is_table is True when the message is
    shaped like the confirmation table — even if every row was "как есть"
    and the mapping is therefore empty. The phase-two gate needs that
    distinction: _parse_replacements returning {} used to mean both "keep
    everything" and "this is ordinary prose", so a filled table of «как есть»
    re-asked confirmation forever instead of starting.

    Accepts →, -> and ::= as the separator — the last is what
    app/api/documents.py's /edit-docx norm-control endpoint already uses.
    A line may carry an optional "Метка: " prefix before the было value.
    """
    mapping: Dict[str, str] = {}
    recognised: List[str] = []
    for line in (text or "").splitlines():
        match = _REPLACEMENT_LINE_RE.match(line.strip())
        if not match:
            continue
        left = match.group("left").strip()
        right = match.group("right").strip()
        if ":" in left:
            _, _, left = left.partition(":")
            left = left.strip()
        if not left:
            continue
        # Левая часть обязана иметь форму реквизита — того же вида, по которому
        # значение и попало в предложенный список. Иначе обычная фраза со
        # стрелкой («я думаю -> надо переделать») разбирается как непустой
        # список замен, а непустой список означает «запускай»: одно случайное
        # сообщение стоило бы тысяч вызовов модели.
        if not any(pattern.fullmatch(left) for pattern, _ in _CANDIDATE_PATTERNS):
            continue
        recognised.append(left)
        if right.lower() in _KEEP_AS_IS:
            continue
        mapping[left] = right
    # Одинокая строка из одних цифр («100000 -> 120000») подходит под форму
    # номера документа, но в переписке про сметы и сроки встречается сама по
    # себе. Согласием она быть не может: рядом должна стоять хотя бы ещё одна
    # замена или значение с узнаваемой формой — фамилия, организация, дата,
    # децимальный номер. Заполненный список из таблицы всегда такой.
    if len(recognised) == 1:
        lone = recognised[0]
        recognisable = any(
            pattern.fullmatch(lone) for pattern, _ in _CANDIDATE_PATTERNS if pattern is not _DOC_NUMBER_RE
        )
        if not recognisable:
            return {}, False
    return mapping, bool(recognised)


def _parse_replacements(text: str) -> Dict[str, str]:
    mapping, _ = _replacement_table(text)
    return mapping


def _replacements_offered_in(messages: List[Dict[str, Any]]) -> Dict[str, str]:
    """The last assistant message's «было → стало» table, with empty
    right-hand sides. Used when the user starts via the «Да, начинай» chip
    instead of filling the list: empty cells become [УКАЗАТЬ: …] placeholders,
    which is what Task 15 promised for a blank answer. Without this, chip-start
    left foreign customer names in the finished 5000 pages — the worst outcome
    the table exists to prevent.

    Only the last assistant message is considered, so an older confirmation
    in the same chat cannot leak its table into a later run that had no
    candidates of its own.
    """
    for message in reversed(messages or []):
        if not isinstance(message, dict) or message.get("role") != "assistant":
            continue
        mapping, is_table = _replacement_table(str(message.get("content") or ""))
        return mapping if is_table else {}
    return {}


def _label_for_value(value: str) -> str:
    """A human-readable description for an old value that has no answer, for
    the [УКАЗАТЬ: <label>] placeholder — reusing the same shape-based
    categories as _find_replacement_candidates. Never the old value itself:
    that would defeat the whole point of the sweep."""
    for pattern, label in _CANDIDATE_PATTERNS:
        if pattern.fullmatch(value):
            return label
    return "Значение"


def _apply_replacements(text: str, mapping: Optional[Dict[str, str]]) -> str:
    """Deterministic sweep (Task 15): plain str.replace per pair, longest
    было-value first, so a decimal designation is not half-replaced by a
    shorter overlapping candidate before its own turn comes up. An empty
    стало value becomes a visible [УКАЗАТЬ: <label>] placeholder — the model
    cannot be trusted to copy values exactly across thousands of calls, this
    is what actually guarantees the old value does not survive."""
    if not mapping or not text:
        return text
    for was in sorted(mapping, key=len, reverse=True):
        become = mapping[was]
        replacement = become if become else f"[УКАЗАТЬ: {_label_for_value(was)}]"
        text = text.replace(was, replacement)
    return text


def _section_context(
    section: Section, chunks: List[Chunk], template_names: Optional[Set[str]],
    replacements: Optional[Dict[str, str]] = None,
    assigned_templates: Optional[Dict[str, str]] = None,
    digests: Optional[Dict[str, str]] = None,
) -> str:
    scoped, templates_for_section = _scope_chunks_for_section(
        section, chunks, template_names, assigned_templates, digests,
    )
    relevant = _select_relevant_chunks(section.title, section.brief, scoped)
    block = _build_context_block(relevant, MAX_CHUNK_CHARS_PER_SECTION, templates_for_section)
    return _apply_replacements(block, replacements) if replacements else block


async def _write_section(
    section: Section, chunks: List[Chunk], previous_tail: str, user_id: int,
    template_names: Optional[Set[str]] = None,
    replacements: Optional[Dict[str, str]] = None,
    user_request: str = "",
    assigned_templates: Optional[Dict[str, str]] = None,
    digests: Optional[Dict[str, str]] = None,
) -> str:
    # Selecting chunks + building the context block tokenizes/scans the whole
    # corpus per section; at 12 concurrent writers that's real CPU time on the
    # (shared) event loop, so it runs in a thread instead.
    context_block = await asyncio.to_thread(
        _section_context, section, chunks, template_names, replacements,
        assigned_templates, digests,
    )

    prompt_parts = []
    # Писатель раздела видел только задачу раздела, а не то, что просил
    # пользователь. На тысяче вызовов это значит, что его требования —
    # к стилю, составу, терминам — не доходят ни до одного из них.
    if user_request:
        prompt_parts.append(
            "Требования пользователя ко всей работе (выполнять буквально, они важнее "
            f"общих соображений):\n{user_request[:2000]}"
        )
    if section.document:
        prompt_parts.append(f"Документ: {section.document}")
        tmpl = (assigned_templates or {}).get(section.document)
        if tmpl:
            prompt_parts.append(
                f"Оформление этого файла берётся из шаблона «{tmpl}». "
                "Структуру повторяй по нему; факты — только из исходников этой же темы, "
                "не из соседних документов заказа."
            )
            snippet = _digest_snippet((digests or {}).get(tmpl, ""), 280)
            if snippet:
                prompt_parts.append(f"Кратко о шаблоне этого файла: {snippet}")
    prompt_parts += [f"Раздел документа: {section.title}", f"Задача раздела: {section.brief}"]
    if context_block:
        prompt_parts.append(f"Релевантные фрагменты исходников:\n{context_block}")
    if previous_tail:
        prompt_parts.append(
            "Конец предыдущего раздела (для согласованности терминов и стиля, не повторяй "
            f"его содержание):\n{previous_tail}"
        )
    prompt_parts.append(
        "Напиши текст ТОЛЬКО этого раздела в Markdown, без заголовка раздела (его добавят "
        "отдельно), без вступлений вида «в этом разделе» и без итоговых выводов в конце. "
        "Не выдумывай цифры и факты, которых нет в исходниках или в задаче раздела."
    )
    prompt = "\n\n".join(prompt_parts)

    messages = [
        {"role": "system", "content": "Ты технический писатель. Пишешь один раздел документа, по существу, без воды."},
        {"role": "user", "content": prompt},
    ]

    # asyncio.CancelledError is a BaseException, not an Exception, so it is
    # never caught below — it passes straight through the retry loop and out
    # of this function. That is required: it's how the user's stop button
    # (generation_hub.hub.cancel) reaches a running job. Do not widen this
    # except clause.
    last_error: Optional[Exception] = None
    for attempt in range(1, SECTION_ATTEMPTS + 1):
        try:
            if section.complexity == "complex":
                if DEEPSEEK_API_KEY:
                    from deepseek_client import get_deepseek_response
                    text, _, _, _ = await asyncio.wait_for(
                        get_deepseek_response(
                            messages, model="deepseek-v4-pro", user_id=user_id, use_tools=False,
                        ),
                        timeout=120,
                    )
                else:
                    text, _, _, _ = await asyncio.wait_for(
                        get_chat_response(
                            messages, model=ESCALATED_WRITER_MODEL, user_id=user_id, use_tools=False, use_skills=False,
                        ),
                        timeout=120,
                    )
            else:
                text, _, _, _ = await asyncio.wait_for(
                    get_chat_response(
                        messages, model=DEFAULT_WRITER_MODEL, user_id=user_id, use_tools=False, use_skills=False,
                    ),
                    timeout=120,
                )
            # Клиенты моделей не бросают исключение на ошибке API — они
            # возвращают её обычным текстом. Без этой проверки такой ответ
            # становится телом раздела: повтор не срабатывает, раздел
            # считается удачным, и прогон на тысячи разделов рапортует
            # «Готово», собрав документ из сообщений об ошибке.
            stripped = text.strip()
            if not stripped or stripped.startswith("❌") or stripped in _NOT_A_SECTION:
                raise RuntimeError(stripped[:200] or "пустой ответ модели")
            return stripped
        except Exception as e:
            last_error = e
            if attempt < SECTION_ATTEMPTS:
                logger.warning(
                    "docgen section '%s' attempt %d/%d failed, retrying: %s",
                    section.title, attempt, SECTION_ATTEMPTS, e,
                )
                # Clamped, not indexed directly: raising SECTION_ATTEMPTS without
                # extending the delays tuple would otherwise IndexError here — in
                # the one code path whose whole job is surviving failures.
                await asyncio.sleep(SECTION_RETRY_DELAYS[min(attempt, len(SECTION_RETRY_DELAYS)) - 1])
            else:
                logger.error(
                    f"docgen section '{section.title}' failed after {SECTION_ATTEMPTS} attempts: {e}",
                    exc_info=True,
                )
    return f"{_SECTION_FAILED_PREFIX}: {str(last_error)[:200]}]"


def _progress_bar(done: int, total: int, cells: int = 10) -> str:
    filled = min(cells, max(0, round(done / total * cells))) if total else 0
    return "▰" * filled + "▱" * (cells - filled)


def _progress_line(done: int, total: int, chars_written: int) -> str:
    """Строка хода работы: доля, разделы и главное — страницы.

    Разделами прогресс мерить бесполезно: человек заказывал страницы и ждёт
    часами, поэтому «готово 1620 стр. из ~5025» — единственное число, по
    которому видно, туда ли всё идёт. Готовые страницы считаются по реально
    написанным символам, а не по плану: если модель пишет короче ожидаемого,
    это должно быть видно сразу, а не в конце.
    """
    pages_done = chars_written // CHARS_PER_PAGE
    pages_total = max(1, total * CHUNK_TARGET_CHARS // CHARS_PER_PAGE)
    percent = round(done / total * 100) if total else 0
    return (
        f"📄 {_progress_bar(done, total)} {percent}% · "
        f"раздел {done} из {total} · готово ~{pages_done} стр. из ~{pages_total}"
    )


async def _update_status(status_msg: Any, text: str) -> None:
    if not status_msg:
        return
    try:
        await status_msg.edit_text(text, parse_mode="Markdown")
    except Exception as e:
        logger.debug(f"Failed to edit docgen status message: {e}")


async def get_docgen_response(
    messages: List[Dict[str, Any]],
    user_text: str,
    user_id: int,
    status_msg: Any,
) -> Tuple[str, List[Dict[str, Any]], str, List[Dict[str, str]]]:
    """Two-phase entry (Task 12): a first turn plans and asks for confirmation
    before any writer call runs; only a clarify reply on a later turn actually
    generates. At up to 5000 sections a run is thousands of model calls and
    hours of wall time, so starting one on the wrong knowledge base or the
    wrong template file by mistake is expensive — this makes that mistake
    cheap to catch before it burns anything."""
    from app.billing.quota import billing_pool
    from clarify import is_clarify_reply

    token = billing_pool.set("computer")
    try:
        # Task 15: a second accepted shape besides the clarify chip reply — a
        # filled «было → стало» list is an ordinary message, not a chip
        # reply, so without this check phase two would never be reached and
        # the mode would re-ask forever. is_clarify_reply(user_text) still
        # covers the chip path (including "skipped") unchanged.
        if not is_clarify_reply(user_text) and not _replacement_table(user_text)[1]:
            return await _confirm_before_generating(user_text, user_id, status_msg)
        return await _run_docgen_after_confirmation(messages, user_text, user_id, status_msg)
    finally:
        billing_pool.reset(token)


def _recover_original_request(messages: List[Dict[str, Any]], user_text: str) -> str:
    """The last user message that isn't itself a confirmation — i.e. the
    request from before the confirmation question was asked. Falls back to
    user_text so a conversation that somehow starts with a confirmation
    doesn't crash.

    Task 15: a filled «было → стало» reply is a confirmation too, but it is
    not a clarify chip reply (is_clarify_reply is False for it) — without
    also skipping it here, this would return the replacement list itself as
    the "original request" instead of walking further back to find it.
    """
    from clarify import is_clarify_reply

    for message in reversed(messages or []):
        if not isinstance(message, dict) or message.get("role") != "user":
            continue
        content = str(message.get("content") or "")
        if is_clarify_reply(content) or _replacement_table(content)[1]:
            continue
        return content
    return user_text


def _format_grouped_names(names: List[str], max_members: int = 5) -> str:
    """Names grouped by containing folder for the confirmation text, which has
    no length limit — unlike the clarify chip prompts. Each group's member
    list is truncated after max_members with "… и ещё N" so a 200-file base
    stays readable instead of listing every file."""
    if not names:
        return ""
    parts = []
    for folder, members in _group_by_folder(names).items():
        if len(members) == 1 and members[0] == folder:
            parts.append(folder)
            continue
        prefix = f"{folder}/"
        shown = [m[len(prefix):] if m.startswith(prefix) else m for m in members[:max_members]]
        suffix = f", … и ещё {len(members) - max_members}" if len(members) > max_members else ""
        parts.append(f"{folder} ({len(members)} файл(ов): {', '.join(shown)}{suffix})")
    return "; ".join(parts)


def _template_blanks(user_id: int, template_names: Set[str]) -> Dict[str, bytes]:
    """Пустые заготовки по каждому .docx-шаблону: стили, поля, колонтитулы
    исходника без его текста. Негодные файлы пропускаются.

    Ключ — имя файла-шаблона, чтобы сборка могла выдать каждому итоговому
    документу свою форму, а не первый попавшийся .docx на весь заказ.
    """
    from document_parser import load_source_docx
    from docx_generator import blank_copy_of_template

    blanks: Dict[str, bytes] = {}
    for name in sorted(template_names):
        # .doc тоже годится: в хранилище оформления он лежит уже
        # сконвертированным в .docx (document_parser.store_source_docx).
        if not name.lower().endswith((".docx", ".doc")):
            continue
        data = load_source_docx(user_id, name)
        if not data:
            continue
        blank = blank_copy_of_template(data)
        if blank and _template_renders_cleanly(blank):
            logger.info("docgen: оформление взято из шаблона %s", name)
            blanks[name] = blank
        elif blank:
            logger.warning("docgen: шаблон %s не годится как основа, оформление стандартное", name)
    return blanks


def _template_base_bytes(user_id: int, template_names: Set[str]) -> Optional[bytes]:
    """Одна заготовка — для пробы и для заказа с единственным шаблоном."""
    blanks = _template_blanks(user_id, template_names)
    if not blanks:
        return None
    return next(iter(blanks.values()))


def _template_renders_cleanly(blank_template: bytes) -> bool:
    """Проверяет заготовку одним пробным документом до того, как по ней будут
    собраны все файлы.

    convert_markdown_to_docx на сбое htmldocx не падает, а кладёт в документ
    СЫРОЙ markdown («## Заголовок», «- пункт») и возвращает валидный .docx.
    Заказчик получил бы десять таких файлов, а сводка сообщила бы, что
    оформление взято из шаблона. Один пробный прогон на весь заказ — это
    ничто рядом с тысячами вызовов модели, а отличает он катастрофу от
    просто стандартного оформления.
    """
    from docx_generator import convert_markdown_to_docx, CONVERSION_FAILED_MARKER

    probe = "## Проверка заголовка\n\nАбзац.\n\n- пункт списка\n- второй пункт\n\n1. первый\n2. второй"
    try:
        result = convert_markdown_to_docx(probe, base_template_bytes=blank_template)
        # Именно текстом документа, а не поиском по байтам: .docx — это zip,
        # внутри него текст сжат, и подстрока в архиве не находится никогда.
        # Из-за этого предохранитель отвечал «шаблон годен» всегда, и заказчик
        # получил документы с сырым markdown вместо текста.
        from docx import Document

        body = "\n".join(p.text for p in Document(io.BytesIO(result)).paragraphs)
    except Exception as e:
        logger.warning("docgen: пробная сборка по шаблону не удалась: %s", e)
        return False
    if CONVERSION_FAILED_MARKER in body:
        return False
    return "Проверка заголовка" in body and "пункт списка" in body


def _document_filename(title: str, used: Set[str]) -> str:
    """Имя файла из названия документа: без разделителей пути и без совпадений."""
    clean = re.sub(r'[\\/:*?"<>|]+', " ", title or "").strip()
    # Планировщик нередко называет документ вместе с расширением («ИТТ_ПЛК.docx»),
    # и файл выходил «ИТТ_ПЛК.docx.docx».
    clean = re.sub(r"\.(docx?|pdf|txt|rtf|odt)$", "", clean, flags=re.IGNORECASE).strip()
    clean = re.sub(r"\s+", " ", clean)[:120] or "Документ"
    candidate = f"{clean}.docx"
    index = 2
    while candidate.lower() in used:
        candidate = f"{clean} ({index}).docx"
        index += 1
    used.add(candidate.lower())
    return candidate


def _group_by_document(
    outline: List[Section], section_texts: List[str]
) -> List[Tuple[str, str]]:
    """(название документа, markdown) для каждого отдельного файла.

    Пока планировщик не проставил document, всё собирается в один файл — так
    режим вёл себя всегда. Заголовок главы печатается один раз при смене:
    иначе двухуровневый план стал бы плоской простынёй.
    """
    # Ключ — название документа, порядок — по первому появлению. Собирать
    # подряд идущие разделы было бы короче, но стоит планировщику перемешать
    # разделы двух документов — и вместо десяти файлов выходит двадцать,
    # половина с дописанным «(2)» в имени.
    parts_by_document: Dict[str, List[str]] = {}
    chapter_by_document: Dict[str, str] = {}
    for section, text in zip(outline, section_texts):
        parts = parts_by_document.setdefault(section.document, [])
        if section.chapter and chapter_by_document.get(section.document) != section.chapter:
            chapter_by_document[section.document] = section.chapter
            parts.append(f"# {section.chapter}")
        parts.append(f"## {section.title}\n\n{text}")
    return [(title, "\n\n".join(parts)) for title, parts in parts_by_document.items()]


def _zip_documents(files: List[Dict[str, Any]], archive_name: str = "Документы.zip") -> Dict[str, Any]:
    """Несколько .docx одним архивом — десять отдельных вложений в переписке
    выглядят как десять сообщений, а скачивать их нужно вместе."""
    import zipfile

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for item in files:
            archive.writestr(item["filename"], item["bytes"])
    return {"filename": archive_name, "bytes": buffer.getvalue()}


def _estimate_cost_usd(outline: List[Section], has_sources: bool) -> float:
    """Верхняя оценка счёта за прогон, в долларах.

    Считается по потолкам, а не по среднему: вход раздела берётся как полный
    MAX_CHUNK_CHARS_PER_SECTION плюс служебная часть промпта, выход — как
    целый раздел. Реальный счёт выходит меньше, и это правильная сторона
    ошибки для числа, по которому человек решает, запускать ли прогон на
    тысячи вызовов. Повторы после сбоев (SECTION_ATTEMPTS) сюда не входят:
    они редки и не меняют порядок величины.
    """
    from app.billing.costs import model_cost_usd
    from model_context import estimate_tokens

    context_chars = MAX_CHUNK_CHARS_PER_SECTION if has_sources else 0
    input_tokens = estimate_tokens("x" * (context_chars + 1200))
    output_tokens = estimate_tokens("x" * CHUNK_TARGET_CHARS * SECTION_OVERSHOOT)
    escalated_model = "deepseek-v4-pro" if DEEPSEEK_API_KEY else ESCALATED_WRITER_MODEL

    total = 0.0
    for section in outline:
        model = escalated_model if section.complexity == "complex" else DEFAULT_WRITER_MODEL
        total += model_cost_usd(model, input_tokens, output_tokens)
    return total


async def _confirm_before_generating(
    user_text: str, user_id: int, status_msg: Any
) -> Tuple[str, List[Dict[str, Any]], str, List[Dict[str, str]]]:
    """Phase one: plan, then ask — never write a single section. One planner
    call is cheap next to the thousands of writer calls it would otherwise
    green-light blindly, and it buys precise numbers for the question instead
    of a blind "are you sure?"."""
    from clarify import normalize_questions, pack_search

    await _update_status(status_msg, "📄 Читаю исходники и строю план документа...")
    chunks, truncated_sources, skipped_sources = await asyncio.to_thread(_extract_source_chunks, user_id)
    outline, template_names, planning_failed = await _plan_document(
        user_text, chunks, user_id, status_msg=status_msg
    )
    candidates = await asyncio.to_thread(_find_replacement_candidates, chunks)
    candidates = await _label_replacement_candidates(candidates, user_id)

    total = len(outline)
    approx_pages = max(1, total * CHUNK_TARGET_CHARS // CHARS_PER_PAGE)
    approx_cost = _estimate_cost_usd(outline, has_sources=bool(chunks))

    doc_names: List[str] = []
    for chunk in chunks:
        if chunk.doc_name not in doc_names:
            doc_names.append(chunk.doc_name)
    knowledge_names = [n for n in doc_names if n not in template_names] if template_names else doc_names

    planned_documents = len({s.document for s in outline if s.document}) or 1
    requested_documents = _requested_documents(user_text)
    lines = [
        (f"Документов: {planned_documents}, разделов: {total}, примерно {approx_pages} стр."
         if planned_documents > 1
         else f"Разделов: {total}, примерно {approx_pages} стр."),
    ]
    # Недостача видна до запуска, а не после: иначе человек узнаёт о ней,
    # открыв архив на два файла вместо десяти.
    if requested_documents and planned_documents < requested_documents:
        lines.append(
            f"Заказано документов: {requested_documents}, в плане только {planned_documents}. "
            "Отмените и переформулируйте перечень, если нужны все."
        )
    lines += [
        f"Оценка стоимости API: не больше ${approx_cost:.2f} (по потолку контекста на раздел)."
        + (
            " Сложные разделы пишет DeepSeek." if DEEPSEEK_API_KEY
            else f" DEEPSEEK_API_KEY не задан, сложные разделы пойдут на {ESCALATED_WRITER_MODEL} — дороже в разы."
        ),
    ]
    if template_names:
        lines.append(
            f"Шаблон оформления ({len(template_names)} файл(ов)): "
            f"{_format_grouped_names(sorted(template_names))}."
        )
        planned_titles = list(dict.fromkeys(s.document for s in outline if s.document))
        if planned_titles and len(template_names) > 1:
            digests = _file_digests(chunks)
            title_texts = _title_texts_for_matching(outline, digests, template_names)
            assigned = _assign_templates(
                planned_titles, template_names, digests=digests, title_texts=title_texts,
            )
            matched = sum(1 for title in planned_titles if assigned.get(title))
            lines.append(
                f"Каждому документу — свой шаблон по имени и содержанию файла "
                f"({matched} из {len(planned_titles)} сопоставлено)."
            )
    else:
        lines.append("Шаблон оформления не определён.")
    lines.append(
        f"База знаний: {_format_grouped_names(knowledge_names)}." if knowledge_names
        else "База знаний: файлы не найдены."
    )
    if truncated_sources:
        lines.append(
            f"Прочитаны не целиком: {', '.join(truncated_sources)}. Загрузите их заново в режиме «Документы»."
        )
    if skipped_sources:
        lines.append(
            f"Не поместились в общий лимит базы знаний ({MAX_SOURCE_CHARS_TOTAL // 1_000_000} млн символов) "
            f"и не читались вовсе: {', '.join(skipped_sources)}."
        )
    if planning_failed:
        lines.append("План построить не удалось — запасной вариант: один раздел на весь документ.")
    if candidates:
        lines.append("")
        lines.append("Это повторяется в примерах — укажите, чем заменить в новых файлах:")
        lines.append("")
        for value, label in candidates:
            lines.append(f"{label}: {value} → ")
        lines.append("")
        lines.append("Ответьте этим же списком, дописав значения справа.")
        lines.append("Пустая строка — поставлю метку [УКАЗАТЬ: …]. Напишите «как есть», если менять не нужно.")
        lines.append(
            f"Кнопка «{_CONFIRM_START}» оставит эти значения как в примерах; "
            f"«{_CONFIRM_START_PLACEHOLDERS}» — подставит вместо них метки [УКАЗАТЬ: …]."
        )
    text = "\n".join(lines)

    start_options = [_CONFIRM_START]
    if candidates:
        start_options.append(_CONFIRM_START_PLACEHOLDERS)
    start_options.append(_CONFIRM_CANCEL)
    questions_raw = [
        {
            "prompt": f"Начинать генерацию? Разделов: {total}, примерно {approx_pages} страниц",
            "options": start_options,
        }
    ]
    if template_names:
        questions_raw.append({
            # Count-based, not name-based: a chip prompt over 160 chars is
            # silently dropped by clarify.normalize_questions, and naming
            # dozens of template files here would do exactly that. The detail
            # goes in `text` above instead, which has no limit.
            "prompt": f"Шаблоны определены верно? Файлов: {len(template_names)}",
            "options": [_TEMPLATE_CONFIRMED, _TEMPLATE_NONE, _TEMPLATE_WRONG],
        })
    questions = normalize_questions(questions_raw)
    return text, [], "", pack_search(questions)


async def _run_docgen_after_confirmation(
    messages: List[Dict[str, Any]], user_text: str, user_id: int, status_msg: Any
) -> Tuple[str, List[Dict[str, Any]], str, List[Dict[str, str]]]:
    """Phase two: honour the answers. Anything short of an explicit start exits
    before anything is loaded or planned; otherwise the original request is
    recovered and the reply itself is threaded into planning so an answer like
    "шаблон другой файл" actually changes the plan instead of being asked for
    and then ignored.

    The gate requires _CONFIRM_START rather than merely the absence of
    _CONFIRM_CANCEL, because the clarify card is a shared component with a
    "Пропустить" button bound to Enter: skipping every question sends
    "Уточнения пропущены…", which is a clarify reply containing neither label.
    Treating that as consent would launch thousands of model calls on one
    accidental keypress — exactly what this confirmation exists to prevent.

    Task 15 adds a second accepted shape: a filled «было → стало» list is
    itself treated as consent to start (the user would not bother filling in
    replacement values for a run they don't want), so a message that parses
    into at least one replacement pair also passes the gate. Everything else —
    including a chip reply that is neither start nor cancel — still refuses,
    keeping Task 12's fail-safe intact.
    """
    replacements, table_reply = _replacement_table(user_text)
    wants_placeholders = _CONFIRM_START_PLACEHOLDERS in user_text
    started = wants_placeholders or _CONFIRM_START in user_text
    if not started and not table_reply:
        if _CONFIRM_CANCEL in user_text:
            return "Отменил, ничего не генерировал.", [], "", []
        return (
            "Не начинаю: подтверждение не получено. Нажмите «"
            f"{_CONFIRM_START}», если документ нужно сгенерировать.",
            [], "", [],
        )
    # Только по отдельной кнопке: пустые ячейки таблицы превращаются в метки
    # [УКАЗАТЬ: …]. «Да, начинай» значит «начинай», а не «замени двадцать пять
    # реквизитов на метки» — заполненный список остаётся третьим, самым
    # точным способом сказать, чем именно их заменить.
    if wants_placeholders and not table_reply:
        replacements = _replacements_offered_in(messages)
    original_request = _recover_original_request(messages, user_text)
    return await _run_docgen(
        original_request, user_id, status_msg, extra_instruction=user_text, replacements=replacements
    )


async def _run_docgen(
    user_text: str,
    user_id: int,
    status_msg: Any,
    extra_instruction: str = "",
    replacements: Optional[Dict[str, str]] = None,
) -> Tuple[str, List[Dict[str, Any]], str, List[Dict[str, str]]]:
    """Режим Документы: план -> параллельная генерация разделов -> сборка в .docx."""
    await _update_status(status_msg, "📄 Читаю исходники и строю план документа...")
    # Синхронный запрос к БД плюс токенизация всего корпуса: на тысячах страниц
    # это секунды CPU, которые нельзя держать в event loop — он общий на всех.
    chunks, truncated_sources, skipped_sources = await asyncio.to_thread(_extract_source_chunks, user_id)
    no_sources = not chunks
    outline, template_names, planning_failed = await _plan_document(
        user_text, chunks, user_id, extra_instruction, status_msg=status_msg
    )

    total = len(outline)
    if planning_failed:
        plan_status = "📄 План построить не удалось — пишу документ одним разделом..."
    else:
        plan_status = f"📄 План готов: {total} раздел(ов). Пишу текст..."
        if no_sources:
            plan_status += " Исходники не найдены — пишу по одному промпту."
    await _update_status(status_msg, plan_status)

    section_texts: List[str] = [""] * total
    digests = _file_digests(chunks)
    title_texts = _title_texts_for_matching(outline, digests, template_names)
    assigned = _assign_templates(
        [section.document for section in outline if section.document],
        template_names,
        digests=digests,
        title_texts=title_texts,
    )
    tails: Dict[str, str] = {}
    last_global_tail = ""
    done = 0
    for batch_start in range(0, total, MAX_PARALLEL_SECTIONS):
        batch = outline[batch_start:batch_start + MAX_PARALLEL_SECTIONS]
        jobs = [
            _write_section(
                section, chunks,
                tails.get(section.document, "") if section.document else last_global_tail,
                user_id, template_names, replacements,
                user_request=user_text, assigned_templates=assigned, digests=digests,
            )
            for section in batch
        ]
        results = await asyncio.gather(*jobs)
        for offset, text in enumerate(results):
            section_texts[batch_start + offset] = text
            if text and not text.startswith(_SECTION_FAILED_PREFIX):
                last_global_tail = text[-500:]
                if batch[offset].document:
                    tails[batch[offset].document] = last_global_tail
        done += len(batch)
        chars_written = sum(
            len(text) for text in section_texts
            if text and not text.startswith(_SECTION_FAILED_PREFIX)
        )
        await _update_status(status_msg, _progress_line(done, total, chars_written))

    await _update_status(status_msg, "📄 Собираю итоговый .docx...")
    documents = _group_by_document(outline, section_texts)
    blanks = await asyncio.to_thread(_template_blanks, user_id, template_names)
    file_assigned = _assign_templates(
        [title for title, _ in documents],
        set(blanks) or template_names,
        digests=digests,
        title_texts=title_texts,
    )

    from docx_generator import convert_markdown_to_docx
    files: List[Dict[str, Any]] = []
    used_names: Set[str] = set()
    used_templates: Set[str] = set()
    full_markdown = ""
    try:
        for title, markdown_text in documents:
            # Deterministic final sweep (Task 15): the model cannot be trusted
            # to copy values exactly across thousands of calls, so this — not
            # the per-section context replacement above — is what actually
            # guarantees no old requisite survives. Runs on the assembled
            # markdown, before the .docx conversion, so it covers text the
            # writer produced on its own too.
            markdown_text = _apply_replacements(markdown_text, replacements)
            full_markdown += markdown_text
            tmpl_name = file_assigned.get(title or "")
            template_bytes = blanks.get(tmpl_name) if tmpl_name else None
            if template_bytes is None and len(blanks) == 1:
                template_bytes = next(iter(blanks.values()))
                tmpl_name = next(iter(blanks))
            if template_bytes and tmpl_name:
                used_templates.add(tmpl_name)
            docx_bytes = await asyncio.to_thread(
                convert_markdown_to_docx, markdown_text, template_bytes
            )
            files.append({
                "filename": _document_filename(title or "Документ", used_names),
                "bytes": docx_bytes,
            })
    except Exception as e:
        logger.error(f"docgen docx assembly failed: {e}", exc_info=True)
        return f"Не удалось собрать документ: {str(e)[:200]}", [], "", []

    approx_pages = max(1, len(full_markdown) // CHARS_PER_PAGE)
    if len(files) > 1:
        summary = (
            f"Готово. Документов: {len(files)}, разделов: {total}, примерно {approx_pages} стр. "
            "Все файлы — в архиве во вложении."
        )
    else:
        summary = (
            f"Готово. Документ из {total} раздел(ов), примерно {approx_pages} стр. — файл во вложении."
        )
    if used_templates:
        if len(used_templates) > 1:
            summary += f" Оформление: каждому документу свой шаблон ({len(used_templates)} шт.)."
        else:
            summary += " Оформление взято из файла-шаблона."
    else:
        summary += (
            " Оформление — стандартное: файла-шаблона в формате .docx среди исходников не нашлось."
        )
    if no_sources:
        summary += " Исходники не найдены — документ написан по одному промпту."
    if planning_failed:
        summary += " План разделов построить не удалось, поэтому документ написан одним разделом."
    failed_sections = sum(1 for text in section_texts if text.startswith(_SECTION_FAILED_PREFIX))
    if failed_sections:
        summary += f" Не удалось сгенерировать разделов: {failed_sections} из {total} — они помечены в тексте."
    if truncated_sources:
        summary += (
            f" Прочитаны не целиком: {', '.join(truncated_sources)}. "
            "Загрузите эти файлы заново, уже в режиме «Документы» — тогда они будут прочитаны полностью."
        )
    if skipped_sources:
        summary += (
            f" Не поместились в общий лимит базы знаний и не читались вовсе: {', '.join(skipped_sources)}."
        )
    if replacements:
        replaced_count = sum(1 for v in replacements.values() if v)
        placeholder_count = sum(1 for v in replacements.values() if not v)
        if replaced_count:
            summary += f" Заменено реквизитов: {replaced_count}."
        if placeholder_count:
            summary += f" Оставлены метки [УКАЗАТЬ: …] вместо не указанных значений: {placeholder_count}."
    # Десять отдельных вложений в переписке — десять карточек, которые качают
    # по одной. Когда документов больше одного, отдаётся архив.
    if len(files) > 1:
        files = [_zip_documents(files)]
    return summary, files, "", []
