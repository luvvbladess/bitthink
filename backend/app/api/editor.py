"""Документ беседы в OnlyOffice: открыть, сохранить версию, ИИ-правка фрагмента.

Все, кто открыл документ с одним key, правят его вместе вживую – это делает
сам сервер документов. Мы отдаём файл, принимаем сохранения и пишем текст
по выделенному фрагменту для плагина «ИИ-правка».
"""

from __future__ import annotations

import asyncio
import io
import logging
import time
import uuid
from pathlib import Path
from typing import Any, Optional
from urllib.parse import urlsplit

from fastapi import APIRouter, Depends, HTTPException, Request, status
from fastapi.responses import FileResponse
from jose import JWTError, jwt

from app.auth import get_current_user
from app.config import get_settings
from app.core.repository import repo

logger = logging.getLogger(__name__)
router = APIRouter()

PLUGIN_GUID = "asc.{7B1E5A2C-3D4F-4A6B-9C8D-1E2F3A4B5C6D}"
LINK_TTL = 12 * 3600
KEEP_VERSIONS = 10
MAX_CONTEXT_CHARS = 40_000
MAX_SELECTION_CHARS = 20_000
_SAVE_STATUSES = {2, 6}  # 2 – все закрыли документ, 6 – принудительное сохранение
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


# doc_id → ids of people with the document open for editing in OnlyOffice.
# Filled from callbacks (status 1 lists current editors, 2/4 mean everyone left).
# ponytail: in-process memory, fine for the single uvicorn worker; after a
# backend restart it is empty until the next callback. Move to the DB if the
# backend ever runs several workers.
_EDITING: dict[str, set[str]] = {}
_PHONE_LOCKS: dict[str, asyncio.Lock] = {}


def _enabled() -> bool:
    return bool(get_settings().ONLYOFFICE_JWT_SECRET)


def _require_enabled() -> None:
    if not _enabled():
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail="Редактор документов не настроен")


# ---------------------------------------------------------------- links

def sign_link(doc_id: str, purpose: str, user: int = 0, ttl: int = LINK_TTL) -> str:
    """Подписанная ссылка на одно действие с одним документом."""
    settings = get_settings()
    payload = {"doc": doc_id, "p": purpose, "u": int(user), "exp": int(time.time()) + ttl}
    return jwt.encode(payload, settings.SECRET_KEY, algorithm=settings.ALGORITHM)


def check_link(token: str, doc_id: str, purpose: str) -> dict:
    settings = get_settings()
    try:
        payload = jwt.decode(token or "", settings.SECRET_KEY, algorithms=[settings.ALGORITHM])
    except JWTError as exc:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Ссылка недействительна") from exc
    if payload.get("doc") != doc_id or payload.get("p") != purpose:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Ссылка недействительна")
    return payload


def _ds_token(payload: dict) -> str:
    return jwt.encode(payload, get_settings().ONLYOFFICE_JWT_SECRET, algorithm="HS256")


def read_callback(body: dict, authorization: str = "") -> dict:
    """Тело колбэка сервера документов, проверенное его JWT. Без подписи – отказ."""
    secret = get_settings().ONLYOFFICE_JWT_SECRET
    token = body.get("token")
    header = (authorization or "").removeprefix("Bearer ").strip()
    try:
        if token:
            return jwt.decode(token, secret, algorithms=["HS256"])
        if header:
            decoded = jwt.decode(header, secret, algorithms=["HS256"])
            return decoded.get("payload") or decoded
    except JWTError as exc:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Неверная подпись") from exc
    raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Нет подписи")


def internal_download_url(url: str) -> str:
    """Файл забираем только у своего сервера документов, какой бы хост ни был в ссылке.

    Сервер строит ссылку от адреса, по которому к нему пришёл браузер
    (https://сайт/onlyoffice/cache/...). Путь сохраняем, хост меняем на внутренний.
    """
    settings = get_settings()
    parts = urlsplit(url or "")
    path = parts.path or "/"
    prefix = urlsplit(settings.ONLYOFFICE_PUBLIC_URL).path.rstrip("/")
    if prefix and path.startswith(prefix + "/"):
        path = path[len(prefix):]
    query = f"?{parts.query}" if parts.query else ""
    return f"{settings.ONLYOFFICE_INTERNAL_URL.rstrip('/')}{path}{query}"


# ---------------------------------------------------------------- storage

def _doc_dir(doc_id: str) -> Path:
    # Не в UPLOAD_DIR: всё оттуда раздаётся статикой по /uploads.
    return get_settings().DATA_DIR / "editor" / doc_id


def version_path(doc_id: str, version: int) -> Path:
    return _doc_dir(doc_id) / f"v{int(version)}.docx"


def _prune_versions(doc_id: str, current: int) -> None:
    # v1 – исходник, его не трогаем: к нему всегда можно вернуться.
    for old in range(2, current - KEEP_VERSIONS + 1):
        try:
            version_path(doc_id, old).unlink(missing_ok=True)
        except OSError:
            pass


def _find_attachment(conv, filename: str) -> Optional[tuple[dict, Optional[int]]]:
    """Карточка файла в переписке и её автор. Последняя с таким именем."""
    for message in reversed(conv.messages):
        attachment = message.attachment or {}
        for item in [attachment, *(attachment.get("files") or [])]:
            if isinstance(item, dict) and str(item.get("name") or "") == filename:
                return item, message.author_user_id
    return None


def source_bytes(conv, filename: str) -> Optional[bytes]:
    """Исходный .docx из переписки: сгенерированный файл или оригинал загрузки."""
    found = _find_attachment(conv, filename)
    if not found:
        return None
    attachment, author = found
    url = str(attachment.get("url") or "")
    if url.startswith("/uploads/"):
        root = get_settings().UPLOAD_DIR.resolve()
        candidate = (root / url.removeprefix("/uploads/")).resolve()
        if candidate.is_relative_to(root) and candidate.is_file():
            return candidate.read_bytes()
    from document_parser import load_source_docx

    for owner in dict.fromkeys([author, conv.owner_id]):
        if owner:
            data = load_source_docx(int(owner), filename)
            if data:
                return data
    return None


def _open_record(conv_id: str, filename: str, bot_id: int, data_loader) -> tuple[str, int]:
    """(id, version) документа беседы. Первый раз – копия исходника как v1."""
    from sqlalchemy.exc import IntegrityError

    from app.db.engine import SyncSessionLocal
    from app.db.models import EditorDocument

    with SyncSessionLocal() as session:
        record = session.query(EditorDocument).filter_by(conversation_id=conv_id, filename=filename).first()
        if record and version_path(record.id, record.version).is_file():
            return record.id, record.version
        data = data_loader()
        if not data:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Исходный .docx не найден")
        if record is None:
            record = EditorDocument(id=uuid.uuid4().hex, conversation_id=conv_id, filename=filename, version=1, created_by=bot_id)
            session.add(record)
            try:
                session.commit()
            except IntegrityError:
                # Второй участник открыл тот же файл одновременно – берём его запись.
                session.rollback()
                record = session.query(EditorDocument).filter_by(conversation_id=conv_id, filename=filename).one()
        path = version_path(record.id, record.version)
        if not path.is_file():
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(data)
        return record.id, record.version


def _record(doc_id: str):
    from app.db.engine import SyncSessionLocal
    from app.db.models import EditorDocument

    with SyncSessionLocal() as session:
        record = session.get(EditorDocument, doc_id)
        if record is None:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Документ не найден")
        session.expunge(record)
        return record


def store_saved_version(doc_id: str, data: bytes, bump: bool, edited_by: Optional[int] = None) -> int:
    """Записать сохранённый файл. bump – все закрыли документ: следующее открытие
    получит новый key и свежий файл. Принудительное сохранение пишет поверх
    текущей версии, чтобы не разорвать живую сессию.

    Правка становится файлом беседы: карточка в чате, ИИ, Пилот и «Документы»
    дальше работают с ней. Исходник остаётся версией v1."""
    from datetime import datetime, timezone

    from app.db.engine import SyncSessionLocal
    from app.db.models import EditorDocument

    with SyncSessionLocal() as session:
        record = session.get(EditorDocument, doc_id)
        if record is None:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Документ не найден")
        version = record.version + 1 if bump else record.version
        path = version_path(doc_id, version)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(data)
        if bump:
            record.version = version
        if edited_by:
            record.edited_by = int(edited_by)
        record.updated_at = datetime.now(timezone.utc)
        session.commit()
        conv_id, filename, created_by = record.conversation_id, record.filename, record.created_by
    _prune_versions(doc_id, version)
    _write_back(conv_id, filename, created_by, data)
    _refresh_chat_text(conv_id, filename, created_by, data)
    return version


def _write_back(conv_id: str, filename: str, bot_id: int, data: bytes) -> None:
    """Положить правку туда, откуда файл беседы читают все остальные."""
    try:
        conv = repo._manager.conversation_view(int(bot_id), conv_id)
        found = _find_attachment(conv, filename) if conv else None
        if not found:
            return
        attachment, author = found
        url = str(attachment.get("url") or "")
        if url.startswith("/uploads/"):
            root = get_settings().UPLOAD_DIR.resolve()
            target = (root / url.removeprefix("/uploads/")).resolve()
            if target.is_relative_to(root) and target.is_file():
                target.write_bytes(data)
            return
        from document_parser import _source_store_path

        path = _source_store_path(int(author or conv.owner_id), filename)
        if path is not None:
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(data)
    except Exception:
        logger.exception("Editor: write-back failed for %s", filename)


def _docx_text(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(line for line in lines if line.strip())


def _refresh_chat_text(conv_id: str, filename: str, bot_id: int, data: bytes) -> None:
    """Чат должен видеть исправленный документ, а не версию до правок."""
    try:
        from conversations import conversation_manager

        conversation_manager.add_document(int(bot_id), filename, _docx_text(data), conv_id=conv_id)
    except Exception:
        logger.exception("Editor: chat text refresh failed for %s", filename)


# ---------------------------------------------------------------- files of a conversation

def _editor_records(conv_id: str) -> dict:
    from app.db.engine import SyncSessionLocal
    from app.db.models import EditorDocument

    with SyncSessionLocal() as session:
        rows = session.query(EditorDocument).filter_by(conversation_id=conv_id).all()
        for row in rows:
            session.expunge(row)
    return {row.filename: row for row in rows}


def _iso(value) -> str:
    return value.isoformat() if hasattr(value, "isoformat") else str(value or "")


def sign_file_link(conv_id: str, filename: str, ttl: int = 3600) -> str:
    settings = get_settings()
    payload = {"conv": conv_id, "name": filename, "p": "download", "exp": int(time.time()) + ttl}
    return jwt.encode(payload, settings.SECRET_KEY, algorithm=settings.ALGORITHM)


def check_file_link(token: str) -> tuple[str, str]:
    settings = get_settings()
    try:
        payload = jwt.decode(token or "", settings.SECRET_KEY, algorithms=[settings.ALGORITHM])
    except JWTError as exc:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Ссылка устарела, обновите список файлов") from exc
    if payload.get("p") != "download":
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Ссылка недействительна")
    return str(payload.get("conv") or ""), str(payload.get("name") or "")


def _has_stored_source(conv, filename: str, author: Optional[int]) -> bool:
    if not filename.lower().endswith((".docx", ".doc")):
        return False
    from document_parser import _source_store_path

    for owner in dict.fromkeys([author, conv.owner_id]):
        path = _source_store_path(int(owner), filename) if owner else None
        if path is not None and path.is_file():
            return True
    return False


def conversation_files(conv) -> list[dict]:
    """Все файлы беседы: кто и когда загрузил, кто и когда последний раз правил."""
    records = _editor_records(conv.id)
    found: dict[str, dict] = {}
    for message in conv.messages:
        attachment = message.attachment or {}
        for item in [attachment, *(attachment.get("files") or [])]:
            name = str((item or {}).get("name") or "").strip() if isinstance(item, dict) else ""
            if not name or item.get("status") == "error":
                continue
            author = message.author_user_id or (conv.owner_id if message.role == "user" else None)
            found[name] = {
                "name": name,
                "kind": str(item.get("type") or "document"),
                "size": int(item.get("size") or 0),
                "url": str(item.get("url") or ""),
                "uploaded_at": message.timestamp,
                "uploaded_by_id": author,
                "from_assistant": message.role == "assistant",
            }
    people = {row.edited_by for row in records.values() if row.edited_by}
    people |= {item["uploaded_by_id"] for item in found.values() if item["uploaded_by_id"]}
    cards = repo._manager.author_cards(list(people))
    files = []
    for name, item in found.items():
        record = records.get(name)
        edited = bool(record and (record.version > 1 or record.edited_by))
        has_bytes = bool(item["url"]) or bool(record) or _has_stored_source(conv, name, item["uploaded_by_id"])
        files.append({
            "name": name,
            "kind": item["kind"],
            "size": item["size"],
            "uploaded_at": item["uploaded_at"],
            "uploaded_by": "Bit-Think" if item["from_assistant"] else (cards.get(item["uploaded_by_id"]) or {}).get("name") or "Участник",
            "editable": name.lower().endswith((".docx", ".doc")) and has_bytes,
            "edited_at": _iso(record.updated_at) if edited else None,
            "edited_by": (cards.get(record.edited_by) or {}).get("name") if edited and record.edited_by else None,
            "version": record.version if record else 1,
            "download_url": (
                f"/conversations/{conv.id}/files/download?t={sign_file_link(conv.id, name)}" if has_bytes else None
            ),
        })
    files.sort(key=lambda f: f["edited_at"] or f["uploaded_at"], reverse=True)
    return files


def file_bytes(conv, filename: str) -> Optional[bytes]:
    """Текущая версия файла беседы: правка из редактора или то, что загрузили."""
    record = _editor_records(conv.id).get(filename)
    if record:
        path = version_path(record.id, record.version)
        if path.is_file():
            return path.read_bytes()
    return source_bytes(conv, filename)


# ---------------------------------------------------------------- access

async def _conversation_for(web_user_id: str, conv_id: str):
    bot_id = await repo.ensure_user(web_user_id)
    conv = await asyncio.to_thread(repo._manager.conversation_view, bot_id, conv_id)
    if conv is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Беседа не найдена")
    return bot_id, conv


def _allowed(bot_id: int, conv_id: str) -> bool:
    return repo._manager.conversation_view(int(bot_id), conv_id) is not None


# ---------------------------------------------------------------- routes

@router.get("/config")
async def editor_config():
    return {"enabled": _enabled()}


@router.post("/open")
async def open_document(data: dict, user_id: str = Depends(get_current_user)):
    _require_enabled()
    settings = get_settings()
    conv_id = str(data.get("conversation_id") or "")
    filename = str(data.get("filename") or "").strip()
    if not conv_id or not filename.lower().endswith((".docx", ".doc")):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Открыть можно .docx из этой беседы")
    bot_id, conv = await _conversation_for(user_id, conv_id)
    doc_id, version = await asyncio.to_thread(
        _open_record, conv.id, filename, bot_id, lambda: source_bytes(conv, filename)
    )
    names = await asyncio.to_thread(repo._manager.author_cards, [bot_id])
    backend = settings.BACKEND_INTERNAL_URL.rstrip("/")
    app_url = (settings.PUBLIC_APP_URL or str(data.get("origin") or "")).rstrip("/")
    if not app_url.startswith(("http://", "https://")):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Нет адреса приложения")
    title = filename if filename.lower().endswith(".docx") else f"{filename}x"
    config: dict[str, Any] = {
        "documentType": "word",
        "document": {
            "fileType": "docx",
            "key": f"{doc_id}-{version}",
            "title": title,
            "url": f"{backend}/editor/{doc_id}/file?t={sign_link(doc_id, 'file')}",
            "permissions": {"edit": True, "comment": True, "review": True, "download": True, "print": True},
        },
        "editorConfig": {
            "mode": "edit",
            "lang": "ru",
            "callbackUrl": f"{backend}/editor/{doc_id}/callback?t={sign_link(doc_id, 'callback', ttl=30 * 24 * 3600)}",
            "user": {"id": str(bot_id), "name": (names.get(bot_id) or {}).get("name") or "Участник"},
            "customization": {"forcesave": True, "comments": True, "compactHeader": True, "uiTheme": "theme-light"},
            "plugins": {"autostart": [PLUGIN_GUID], "pluginsData": [f"{app_url}/onlyoffice-plugin/config.json"]},
        },
    }
    if data.get("mobile"):
        # Free OnlyOffice only views on phones; edits there go through
        # /paragraphs as tracked changes. Viewers join the live session key.
        config["type"] = "mobile"
        config["editorConfig"]["mode"] = "view"
        config["editorConfig"].pop("plugins", None)
        config["document"]["permissions"] = {"edit": False, "comment": False, "review": False, "download": True, "print": True}
    config["token"] = _ds_token(config)
    return {
        "doc_id": doc_id,
        "server": settings.ONLYOFFICE_PUBLIC_URL.rstrip("/"),
        "config": config,
        "ai_token": sign_link(doc_id, "ai", user=bot_id),
    }


@router.get("/{doc_id}/file")
async def document_file(doc_id: str, t: str = ""):
    check_link(t, doc_id, "file")
    record = await asyncio.to_thread(_record, doc_id)
    path = version_path(doc_id, record.version)
    if not path.is_file():
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Файл не найден")
    return FileResponse(path, media_type=DOCX_MIME)


@router.post("/{doc_id}/callback")
async def document_callback(doc_id: str, request: Request, t: str = ""):
    check_link(t, doc_id, "callback")
    body = read_callback(await request.json(), request.headers.get("authorization", ""))
    state = int(body.get("status") or 0)
    if state == 1:
        _EDITING[doc_id] = {str(item) for item in (body.get("users") or [])}
    elif state in (2, 3, 4):
        _EDITING.pop(doc_id, None)
    if state in _SAVE_STATUSES and body.get("url"):
        import httpx

        try:
            async with httpx.AsyncClient(timeout=60) as client:
                response = await client.get(internal_download_url(str(body["url"])))
                response.raise_for_status()
        except httpx.HTTPError:
            logger.exception("Editor: saved file download failed for %s", doc_id)
            return {"error": 1}
        users = [str(item) for item in (body.get("users") or [])]
        edited_by = int(users[0]) if users and users[0].isdigit() else None
        await asyncio.to_thread(store_saved_version, doc_id, response.content, state == 2, edited_by)
        record = await asyncio.to_thread(_record, doc_id)
        # Участники общей беседы увидят новую версию в списке файлов сразу.
        await repo.notify_room(record.conversation_id)
    return {"error": 0}


def _editing_names(doc_id: str, me: int) -> list[str]:
    """Who has the document open for editing on a computer. Includes the phone
    user too: their own desktop session would still save over a phone edit."""
    ids = [int(item) for item in _EDITING.get(doc_id, set()) if item.isdigit()]
    if not ids:
        return []
    cards = repo._manager.author_cards(ids)
    return ["вы" if item == me else (cards.get(item) or {}).get("name") or "участник" for item in ids]


@router.get("/{doc_id}/paragraphs")
async def document_paragraphs(doc_id: str, user_id: str = Depends(get_current_user)):
    """Пункты текущей версии для правки с телефона."""
    from app.services.tracked_edit import list_paragraphs

    record = await asyncio.to_thread(_record, doc_id)
    bot_id, _ = await _conversation_for(user_id, record.conversation_id)
    path = version_path(doc_id, record.version)
    if not path.is_file():
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Файл не найден")
    paragraphs = await asyncio.to_thread(list_paragraphs, path.read_bytes())
    busy = await asyncio.to_thread(_editing_names, doc_id, bot_id)
    return {"version": record.version, "paragraphs": paragraphs, "busy": busy}


@router.post("/{doc_id}/paragraphs")
async def edit_paragraph(doc_id: str, data: dict, user_id: str = Depends(get_current_user)):
    """Заменить один пункт как правку рецензирования и сделать это новой версией файла."""
    from app.services.tracked_edit import EditConflict, EditRejected, replace_paragraph

    record = await asyncio.to_thread(_record, doc_id)
    bot_id, _ = await _conversation_for(user_id, record.conversation_id)
    busy = await asyncio.to_thread(_editing_names, doc_id, bot_id)
    if busy:
        # Their live session would save over this file when they close it.
        raise HTTPException(
            status_code=status.HTTP_423_LOCKED,
            detail=f"Документ сейчас правят на компьютере ({', '.join(busy)}). Попробуйте, когда его закроют.",
        )
    try:
        index = int(data.get("index"))
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Не указан пункт") from exc
    text = str(data.get("text") or "").strip()
    if not text:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Пустой текст пункта")
    cards = await asyncio.to_thread(repo._manager.author_cards, [bot_id])
    author = (cards.get(bot_id) or {}).get("name") or "Участник"
    # Two phones saving at once must not write over each other's version.
    async with _PHONE_LOCKS.setdefault(doc_id, asyncio.Lock()):
        record = await asyncio.to_thread(_record, doc_id)
        path = version_path(doc_id, record.version)
        try:
            edited = await asyncio.to_thread(
                replace_paragraph, path.read_bytes(), index, str(data.get("base_text") or ""), text, author,
            )
        except EditConflict as exc:
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(exc)) from exc
        except EditRejected as exc:
            raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail=str(exc)) from exc
        version = await asyncio.to_thread(store_saved_version, doc_id, edited, True, bot_id)
    await repo.notify_room(record.conversation_id)
    return {"version": version}


EDIT_SYSTEM_PROMPT = (
    "Ты редактор деловых и юридических документов на русском. Тебе дают выделенный фрагмент "
    "документа и просьбу, что с ним сделать. Верни ТОЛЬКО новый текст фрагмента, которым его "
    "заменят в документе: без пояснений, вступлений, кавычек вокруг, markdown и заголовков. "
    "Сохраняй нумерацию пунктов, термины и стиль документа. Не выдумывай реквизиты, суммы и даты: "
    "если значения нет в документе и просьбе, оставь как было. Если выделения нет – напиши текст "
    "для вставки по просьбе. Не используй длинное тире «—», только «–»."
)


def _clean_ai_text(text: str) -> str:
    text = (text or "").strip()
    if text.startswith("```"):
        text = text.split("\n", 1)[1] if "\n" in text else ""
        text = text.rsplit("```", 1)[0]
    text = text.strip()
    if len(text) >= 2 and text[0] in "«\"" and text[-1] in "»\"":
        text = text[1:-1].strip()
    return text.replace("—", "–")


@router.post("/{doc_id}/ai")
async def ai_edit(doc_id: str, data: dict, request: Request):
    """Текст для выделенного фрагмента. Доступ – по ai-ссылке из /open: плагин
    работает внутри редактора и не видит сессию приложения."""
    _require_enabled()
    token = (request.headers.get("authorization") or "").removeprefix("Bearer ").strip()
    bot_id = int(check_link(token, doc_id, "ai").get("u") or 0)
    record = await asyncio.to_thread(_record, doc_id)
    if not bot_id or not await asyncio.to_thread(_allowed, bot_id, record.conversation_id):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Нет доступа к документу")
    instruction = str(data.get("instruction") or "").strip()
    selection = str(data.get("selection") or "")[:MAX_SELECTION_CHARS]
    comment = str(data.get("comment") or "").strip()
    if not instruction and not comment:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Напишите, что сделать с фрагментом")

    from app.billing.plans import clamp_model
    from app.billing.quota import QuotaError, assert_can_use, usage_view
    from conversations import conversation_manager
    from openai_client import get_chat_response

    sub = await asyncio.to_thread(conversation_manager.get_subscription, bot_id)
    view = usage_view(sub)
    model = clamp_model(view["tier"], "gpt-6-sol")
    charged = view["tier"] == "free"
    try:
        await asyncio.to_thread(assert_can_use, bot_id, "chat", model)
        if charged:
            limit = int(view["free_daily"]["replies_limit"] or 0)
            if not await asyncio.to_thread(conversation_manager.consume_daily_counter, bot_id, "daily_gpt54", limit):
                raise QuotaError("Лимит на сегодня закончился.", "quota")
    except QuotaError as exc:
        raise HTTPException(status_code=status.HTTP_402_PAYMENT_REQUIRED, detail=str(exc)) from exc

    path = version_path(doc_id, record.version)
    context = ""
    if path.is_file():
        try:
            context = (await asyncio.to_thread(_docx_text, path.read_bytes()))[:MAX_CONTEXT_CHARS]
        except Exception:
            logger.debug("Editor: context extraction failed", exc_info=True)
    request_text = (
        f"Документ «{record.filename}» (последняя сохранённая версия, для контекста):\n{context}\n\n"
        f"Выделенный фрагмент:\n{selection or '(ничего не выделено)'}\n\n"
        + (f"Комментарий к фрагменту:\n{comment}\n\n" if comment else "")
        + f"Просьба:\n{instruction or 'Исправь фрагмент по комментарию.'}"
    )
    answer, _, _, _ = await get_chat_response(
        [{"role": "system", "content": EDIT_SYSTEM_PROMPT}, {"role": "user", "content": request_text}],
        model=model,
        user_id=bot_id,
        use_tools=False,
        reasoning_effort="low",
        use_skills=False,
    )
    if not answer or answer.startswith("❌"):
        if charged:
            await asyncio.to_thread(conversation_manager.refund_daily_counter, bot_id, "daily_gpt54")
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY, detail="Модель не ответила, попробуйте ещё раз")
    return {"text": _clean_ai_text(answer)}
