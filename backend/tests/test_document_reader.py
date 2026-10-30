import asyncio
import io
import json
import hashlib
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from fastapi import HTTPException
from PIL import Image

from app.services.document_reader import reading_content, image_bytes


def document_bytes():
    doc = Document(); doc.add_heading('Условия сотрудничества', 1)
    p = doc.add_paragraph('Оплата '); p.add_run('150 000 ₽').bold = True
    p.add_run(' в течение 10 дней.').italic = True
    doc.add_paragraph('Первый пункт', style='List Number')
    doc.add_paragraph('Второй пункт', style='List Number')
    doc.add_paragraph('Важное условие', style='List Bullet')
    table = doc.add_table(rows=2, cols=3)
    for index, value in enumerate(['Услуга', 'Срок', 'Стоимость']): table.rows[0].cells[index].text = value
    table.rows[1].cells[0].text = 'Подготовка документов'; table.rows[1].cells[1].text = '10 дней'; table.rows[1].cells[2].text = '150 000 ₽'
    picture = io.BytesIO(); Image.new('RGB', (400, 200), 'blue').save(picture, 'PNG'); picture.seek(0); doc.add_picture(picture)
    p = doc.add_paragraph(); link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), doc.part.relate_to('https://example.com/условия', RT.HYPERLINK, is_external=True))
    r=OxmlElement('w:r'); t=OxmlElement('w:t'); t.text='Полные условия';r.append(t);link.append(r);p._p.append(link)
    bad = OxmlElement('w:hyperlink'); bad.set(qn('r:id'), doc.part.relate_to('javascript:alert(1)', RT.HYPERLINK, is_external=True)); r=OxmlElement('w:r'); t=OxmlElement('w:t'); t.text='Небезопасная ссылка';r.append(t);bad.append(r);p._p.append(bad)
    b=io.BytesIO();doc.save(b);return b.getvalue()


def test_reading_preserves_content_lists_tables_images_and_safe_links():
    data=document_bytes(); content=reading_content(data, lambda rid: '/private-image/'+rid)
    blocks=content['blocks']; assert blocks[0]['heading']==1
    assert blocks[1]['content'][1] == {'text':'150 000 ₽','bold':True}
    assert [b['list']['label'] for b in blocks if 'list' in b] == ['1.','2.','•']
    table=next(b for b in blocks if b['type']=='table');assert len(table['rows'])==2 and len(table['rows'][0]['cells'])==3
    image=next(item for b in blocks if b['type']=='paragraph' for item in b['content'] if 'image' in item)
    body,mime=image_bytes(data,image['image'].split('/')[-1]);assert mime=='image/png' and body.startswith(b'\x89PNG')
    links=[i for b in blocks if b['type']=='paragraph' for i in b['content'] if 'href' in i]
    assert len(links)==1 and links[0]['href'].startswith('https://')
    assert 'javascript:' not in json.dumps(content)
    assert 'Небезопасная ссылка' in json.dumps(content,ensure_ascii=False)


def test_reading_keeps_merged_cells_and_does_not_repeat_their_text():
    doc=Document();table=doc.add_table(rows=1,cols=3);table.cell(0,0).merge(table.cell(0,1)).text='Объединённая ячейка'
    b=io.BytesIO();doc.save(b);content=reading_content(b.getvalue(),lambda rid:rid)
    cells=content['blocks'][0]['rows'][0]['cells'];assert len(cells)==2 and cells[0]['span']==2
    assert json.dumps(content,ensure_ascii=False).count('Объединённая ячейка')==1


def test_reading_corrupt_document_and_unrelated_image_are_rejected():
    with pytest.raises(ValueError): reading_content(b'not docx',lambda rid:rid)
    with pytest.raises(ValueError): image_bytes(document_bytes(),'rIdDoesNotExist')


def test_vertical_merge_does_not_duplicate_content_across_rows():
    doc=Document();table=doc.add_table(rows=3,cols=2)
    table.cell(0,0).merge(table.cell(2,0)).text='Общее условие'
    b=io.BytesIO();doc.save(b);content=reading_content(b.getvalue(),lambda rid:rid)
    rows=content['blocks'][0]['rows']
    assert rows[0]['cells'][0]['rowSpan']==3 and len(rows[1]['cells'])==1
    assert json.dumps(content,ensure_ascii=False).count('Общее условие')==1


def test_roman_numbering_and_explicit_restart_are_preserved():
    doc=Document();root=doc.part.numbering_part.element
    abstract=OxmlElement('w:abstractNum');abstract.set(qn('w:abstractNumId'),'99')
    level=OxmlElement('w:lvl');level.set(qn('w:ilvl'),'0')
    for tag,value in [('start','1'),('numFmt','upperRoman'),('lvlText','%1)')]:
        child=OxmlElement('w:'+tag);child.set(qn('w:val'),value);level.append(child)
    abstract.append(level);root.append(abstract)
    number=OxmlElement('w:num');number.set(qn('w:numId'),'99')
    child=OxmlElement('w:abstractNumId');child.set(qn('w:val'),'99');number.append(child)
    override=OxmlElement('w:lvlOverride');override.set(qn('w:ilvl'),'0');start=OxmlElement('w:startOverride');start.set(qn('w:val'),'4');override.append(start);number.append(override);root.append(number)
    p=doc.add_paragraph('Особое условие');num=OxmlElement('w:numPr');id=OxmlElement('w:numId');id.set(qn('w:val'),'99');num.append(id);p._p.get_or_add_pPr().append(num)
    b=io.BytesIO();doc.save(b)
    assert reading_content(b.getvalue(),lambda rid:rid)['blocks'][0]['list']['label']=='IV)'


def test_objects_without_readable_image_offer_original_instead_of_silently_disappearing():
    doc=Document();doc.add_paragraph('Документ с диаграммой').add_run()._r.append(OxmlElement('w:drawing'))
    b=io.BytesIO();doc.save(b)
    assert 'Диаграммы и плавающие объекты доступны в «Оригинале».' in reading_content(b.getvalue(),lambda rid:rid)['warnings']


def test_reading_endpoint_is_opt_in_and_checks_membership(monkeypatch):
    from app.api import conversations,editor
    data=document_bytes()
    async def allowed(*args):return 1,object()
    monkeypatch.setattr(editor,'_conversation_for',allowed)
    monkeypatch.setattr(editor,'file_bytes',lambda *args:data)
    response=asyncio.run(conversations.preview_conversation_file('mine','Contract.docx',user_id='owner',native=True,reading=True))
    parsed=json.loads(response.body);assert 'reader' in parsed and 'pages' not in parsed and 'download_url' in parsed
    assert response.headers['cache-control']=='private, no-store'
    url=next(i['image'] for b in parsed['reader']['blocks'] if b['type']=='paragraph' for i in b['content'] if 'image' in i)
    from urllib.parse import urlsplit,parse_qs
    query=parse_qs(urlsplit(url).query)
    assert editor.check_file_link(query['t'][0])==('mine','Contract.docx')
    async def denied(*args):raise HTTPException(status_code=404)
    monkeypatch.setattr(editor,'_conversation_for',denied)
    with pytest.raises(HTTPException):asyncio.run(conversations.preview_conversation_file('other','Contract.docx',user_id='owner',native=True,reading=True))


def test_images_cannot_cross_conversations_and_stale_versions_fail(monkeypatch):
    from app.api import conversations,editor
    data=document_bytes()
    monkeypatch.setattr(editor,'check_file_link',lambda token:('mine','Contract.docx'))
    with pytest.raises(HTTPException) as error:asyncio.run(conversations.document_reader_image('other','signed','rId1','v'))
    assert error.value.status_code==403
    monkeypatch.setattr(conversations.repo,'conversations',SimpleNamespace(owner=lambda cid:1,view=lambda uid,cid:object()))
    monkeypatch.setattr(editor,'file_bytes',lambda *args:data)
    with pytest.raises(HTTPException) as error:asyncio.run(conversations.document_reader_image('mine','signed','rId1','stale'))
    assert error.value.status_code==409
    content=reading_content(data,lambda rid:rid);rid=next(i['image'] for b in content['blocks'] if b['type']=='paragraph' for i in b['content'] if 'image' in i)
    response=asyncio.run(conversations.document_reader_image('mine','signed',rid,hashlib.sha256(data).hexdigest()))
    assert response.media_type=='image/png' and response.headers['cache-control']=='private, no-store'


def test_docx_original_uses_read_only_word_viewer(monkeypatch):
    from app.config import get_settings
    from app.services.office_viewer import office_viewer
    monkeypatch.setattr(get_settings(),'ONLYOFFICE_JWT_SECRET','test-word-viewer')
    c=office_viewer(document_bytes(),'Contract.docx','mine')['config']
    assert c['documentType']=='word' and c['document']['fileType']=='docx'
    assert c['editorConfig']['mode']=='view' and not c['document']['permissions']['edit']
