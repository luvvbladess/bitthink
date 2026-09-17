from app.config import get_settings  # noqa: F401 — adds bot_core to sys.path

import io

import docgen_router as dg
from docgen_router import (
    MAX_CHUNK_CHARS_PER_SECTION,
    MAX_SOURCE_CHARS_TOTAL,
    PLANNER_PREVIEW_BUDGET,
    PLANNER_PREVIEW_GROUPS,
    Chunk,
    Section,
    _build_context_block,
    _classify_documents_hint,
    _document_previews,
    _expand_template_names,
    _folder_of,
    _group_by_folder,
    _section_context,
    _select_relevant_chunks,
    _parse_outline_response,
)

# _extract_source_chunks(user_id) now reads via conversation_manager.get_documents
# (a raw DB read) instead of parsing chat messages — no cheap way to exercise that
# without a real DB, so per the task-9 brief it goes without a unit test here,
# same as _plan_outline/_write_section (covered by the manual smoke check instead).


def test_select_relevant_chunks_ranks_keyword_match_first():
    chunks = [
        Chunk(id=0, doc_name="d", title="Гидравлическая система", text="Давление в гидравлической системе буксира составляет 250 бар.", tokens=frozenset()),
        Chunk(id=1, doc_name="d", title="Штатное расписание", text="В экипаж входят капитан, механик и два матроса.", tokens=frozenset()),
        Chunk(id=2, doc_name="d", title="Электрооборудование", text="Генератор мощностью 500 кВт питает силовую установку.", tokens=frozenset()),
    ]

    selected = _select_relevant_chunks("Гидравлическая система", "Опиши давление в гидравлике буксира", chunks, top_k=2)

    assert selected[0].id == 0


def test_select_relevant_chunks_falls_back_to_first_chunks_when_no_match():
    chunks = [Chunk(id=i, doc_name="d", title=f"Часть {i}", text="нейтральный текст без общих слов", tokens=frozenset()) for i in range(3)]

    selected = _select_relevant_chunks("совершенно другая тема", "ничего общего", chunks, top_k=2)

    assert [c.id for c in selected] == [0, 1]


def test_build_context_block_respects_char_budget():
    chunks = [Chunk(id=i, doc_name="d", title="t", text="x" * 5000, tokens=frozenset()) for i in range(5)]

    block = _build_context_block(chunks, budget=8000)

    # Allow a small, bounded overhead per chunk for the "[doc_name] " label.
    assert len(block) <= 8000 + len(chunks) * 10


def test_build_context_block_labels_template_and_knowledge_when_both_present():
    chunks = [
        Chunk(id=0, doc_name="шаблон.docx", title="t", text="структура раздела", tokens=frozenset()),
        Chunk(id=1, doc_name="данные.xlsx", title="t", text="конкретные цифры", tokens=frozenset()),
    ]

    block = _build_context_block(chunks, budget=8000, template_names={"шаблон.docx"})

    assert "Формат по шаблону:" in block
    assert "Факты из базы знаний:" in block
    template_part, knowledge_part = block.split("Факты из базы знаний:")
    assert "[шаблон.docx] структура раздела" in template_part
    assert "[данные.xlsx] конкретные цифры" in knowledge_part


def test_build_context_block_falls_back_when_template_name_missing_from_chunks():
    chunks = [
        Chunk(id=0, doc_name="данные.xlsx", title="t", text="конкретные цифры", tokens=frozenset()),
        Chunk(id=1, doc_name="отчёт.docx", title="t", text="другой текст", tokens=frozenset()),
    ]

    block = _build_context_block(chunks, budget=8000, template_names={"несуществующий.docx"})

    assert block.startswith("Факты из базы знаний:")
    assert "Формат по шаблону:" not in block
    assert "[данные.xlsx] конкретные цифры" in block
    assert "[отчёт.docx] другой текст" in block


def test_build_context_block_labels_two_template_files_under_one_heading():
    chunks = [
        Chunk(id=0, doc_name="шаблон1.docx", title="t", text="структура раздела один", tokens=frozenset()),
        Chunk(id=1, doc_name="шаблон2.docx", title="t", text="структура раздела два", tokens=frozenset()),
        Chunk(id=2, doc_name="данные.xlsx", title="t", text="конкретные цифры", tokens=frozenset()),
    ]

    block = _build_context_block(chunks, budget=8000, template_names={"шаблон1.docx", "шаблон2.docx"})

    template_part, knowledge_part = block.split("Факты из базы знаний:")
    assert "[шаблон1.docx] структура раздела один" in template_part
    assert "[шаблон2.docx] структура раздела два" in template_part
    assert "[данные.xlsx] конкретные цифры" in knowledge_part


def test_parse_outline_response_extracts_sections_from_fenced_json():
    response = (
        '```json\n'
        '{"template_documents": [], "sections": '
        '[{"title": "Введение", "brief": "Общее описание", "complexity": "simple"},'
        ' {"title": "Расчёт нагрузки", "brief": "Числа", "complexity": "complex"}]}\n'
        '```'
    )

    sections, templates = _parse_outline_response(response)

    assert [s.title for s in sections] == ["Введение", "Расчёт нагрузки"]
    assert sections[0].complexity == "simple"
    assert sections[1].complexity == "complex"
    assert templates == []


def test_parse_outline_response_returns_empty_on_garbage():
    assert _parse_outline_response("не json вообще") == ([], [])


def test_parse_outline_response_defaults_unknown_complexity_to_simple():
    response = '{"sections": [{"title": "Раздел", "brief": "текст", "complexity": "нечто странное"}]}'

    sections, template = _parse_outline_response(response)

    assert sections[0].complexity == "simple"


def test_parse_outline_response_skips_items_without_title():
    response = '{"sections": [{"brief": "нет заголовка"}, {"title": "Есть заголовок", "brief": "ок"}]}'

    sections, template = _parse_outline_response(response)

    assert [s.title for s in sections] == ["Есть заголовок"]


def test_parse_outline_response_extracts_template_documents_array():
    response = '{"template_documents": ["a.docx", "b.docx"], "sections": [{"title": "Раздел", "brief": "текст"}]}'

    sections, templates = _parse_outline_response(response)

    assert templates == ["a.docx", "b.docx"]


def test_parse_outline_response_wraps_bare_template_document_string():
    response = '{"template_documents": "шаблон.docx", "sections": [{"title": "Раздел", "brief": "текст"}]}'

    sections, templates = _parse_outline_response(response)

    assert templates == ["шаблон.docx"]


def test_parse_outline_response_empty_list_when_key_missing():
    response = '{"sections": [{"title": "Раздел", "brief": "текст"}]}'

    sections, templates = _parse_outline_response(response)

    assert templates == []


def test_expand_template_names_exact_match_selects_that_document():
    chunks = [
        Chunk(id=0, doc_name="a.docx", title="t", text="x", tokens=frozenset()),
        Chunk(id=1, doc_name="b.docx", title="t", text="x", tokens=frozenset()),
    ]

    assert _expand_template_names(["a.docx"], chunks) == {"a.docx"}


def test_expand_template_names_archive_name_selects_all_its_members_only():
    chunks = [
        Chunk(id=0, doc_name="shablony.zip/a.docx", title="t", text="x", tokens=frozenset()),
        Chunk(id=1, doc_name="shablony.zip/b.docx", title="t", text="x", tokens=frozenset()),
        Chunk(id=2, doc_name="other.docx", title="t", text="x", tokens=frozenset()),
    ]

    resolved = _expand_template_names(["shablony.zip"], chunks)

    assert resolved == {"shablony.zip/a.docx", "shablony.zip/b.docx"}


def test_expand_template_names_drops_unknown_name():
    chunks = [Chunk(id=0, doc_name="a.docx", title="t", text="x", tokens=frozenset())]

    assert _expand_template_names(["призрак.docx"], chunks) == set()


def test_expand_template_names_name_both_real_file_and_archive_prefix_selects_union():
    # A name that is both an exact document and a prefix of other documents
    # (e.g. "papka" uploaded standalone, alongside "papka/inner.docx") selects
    # both — nothing the planner named is silently dropped.
    chunks = [
        Chunk(id=0, doc_name="papka", title="t", text="x", tokens=frozenset()),
        Chunk(id=1, doc_name="papka/inner.docx", title="t", text="x", tokens=frozenset()),
    ]

    resolved = _expand_template_names(["papka"], chunks)

    assert resolved == {"papka", "papka/inner.docx"}


def test_expand_template_names_folder_inside_archive_selects_only_that_folder():
    # arhiv.zip holds two folders — shablony (templates) and baza
    # (knowledge) — plus a sub-subfolder inside shablony. Naming just
    # "arhiv.zip/shablony" must select every file under it, however deep,
    # and nothing from "arhiv.zip/baza".
    chunks = [
        Chunk(id=0, doc_name="arhiv.zip/shablony/forma1.docx", title="t", text="x", tokens=frozenset()),
        Chunk(id=1, doc_name="arhiv.zip/shablony/formy/forma2.docx", title="t", text="x", tokens=frozenset()),
        Chunk(id=2, doc_name="arhiv.zip/baza/otchet1.docx", title="t", text="x", tokens=frozenset()),
    ]

    resolved = _expand_template_names(["arhiv.zip/shablony"], chunks)

    assert resolved == {"arhiv.zip/shablony/forma1.docx", "arhiv.zip/shablony/formy/forma2.docx"}


def test_folder_of_bare_name_has_no_slash():
    assert _folder_of("report.docx") == "report.docx"


def test_folder_of_flat_archive_member_groups_by_the_archive():
    assert _folder_of("arhiv.zip/forma1.docx") == "arhiv.zip"


def test_folder_of_nested_archive_member_groups_by_its_own_folder():
    assert _folder_of("arhiv.zip/shablony/formy/forma1.docx") == "arhiv.zip/shablony/formy"


def test_group_by_folder_nested_paths_group_separately_bare_names_group_alone():
    names = [
        "arhiv.zip/shablony/forma1.docx",
        "arhiv.zip/shablony/forma2.docx",
        "arhiv.zip/baza/otchet1.docx",
        "bare.docx",
    ]

    groups = _group_by_folder(names)

    assert groups["arhiv.zip/shablony"] == ["arhiv.zip/shablony/forma1.docx", "arhiv.zip/shablony/forma2.docx"]
    assert groups["arhiv.zip/baza"] == ["arhiv.zip/baza/otchet1.docx"]
    assert groups["bare.docx"] == ["bare.docx"]


def test_extract_source_chunks_stops_at_corpus_ceiling_and_reports_skipped(monkeypatch):
    """Change B: raising the per-file caps removed the only thing that kept
    the total corpus bounded by accident. A document that would push the
    running total over MAX_SOURCE_CHARS_TOTAL must not be loaded, but must be
    named — and a smaller document later in the list still gets its chance."""
    from conversations import conversation_manager

    big = "x" * (MAX_SOURCE_CHARS_TOTAL // 2 + 1)
    docs = [
        {"filename": "a.pdf", "content": big},
        {"filename": "b.pdf", "content": big},
        {"filename": "c.pdf", "content": "small enough to still fit"},
    ]
    monkeypatch.setattr(conversation_manager, "get_documents", lambda uid: docs)

    chunks, truncated, skipped = dg._extract_source_chunks(1)

    loaded = {c.doc_name for c in chunks}
    assert "a.pdf" in loaded, "the first document under the ceiling must load"
    assert "b.pdf" not in loaded, "the document that would exceed the ceiling must not load"
    assert "b.pdf" in skipped, "a document dropped by the ceiling must be named as skipped"
    assert "c.pdf" in loaded, "a smaller document after a skip must still be checked and fit"
    assert not truncated


def test_classify_documents_hint_names_template_and_knowledge_files():
    chunks = [
        Chunk(id=0, doc_name="шаблон_итт.docx", title="t", text="x", tokens=frozenset()),
        Chunk(id=1, doc_name="данные.xlsx", title="t", text="x", tokens=frozenset()),
    ]

    hint = _classify_documents_hint(chunks)

    assert hint != ""
    assert "шаблон_итт.docx" in hint
    assert "данные.xlsx" in hint


def test_classify_documents_hint_empty_without_template_like_names():
    chunks = [
        Chunk(id=0, doc_name="данные.xlsx", title="t", text="x", tokens=frozenset()),
        Chunk(id=1, doc_name="отчёт.docx", title="t", text="x", tokens=frozenset()),
    ]

    assert _classify_documents_hint(chunks) == ""


def test_section_context_matches_select_then_build_context_block():
    chunks = [
        Chunk(id=0, doc_name="шаблон.docx", title="Гидравлика", text="структура раздела про гидравлику", tokens=frozenset({"гидравлика"}), title_tokens=frozenset({"гидравлика"})),
        Chunk(id=1, doc_name="данные.xlsx", title="Прочее", text="конкретные цифры", tokens=frozenset({"цифры"}), title_tokens=frozenset({"прочее"})),
    ]
    section = Section(id=0, title="Гидравлика", brief="Опиши гидравлику", complexity="simple")

    block = _section_context(section, chunks, {"шаблон.docx"})

    expected = _build_context_block(
        _select_relevant_chunks(section.title, section.brief, chunks), MAX_CHUNK_CHARS_PER_SECTION, {"шаблон.docx"}
    )
    assert block == expected
    assert "Формат по шаблону:" in block


def test_document_previews_produces_labeled_blocks_capped_per_doc():
    chunks = [
        Chunk(id=0, doc_name="a.docx", title="t", text="A" * 2000, tokens=frozenset()),
        Chunk(id=1, doc_name="b.docx", title="t", text="B" * 2000, tokens=frozenset()),
    ]

    previews = _document_previews(chunks, budget_per_doc=1500)

    assert "--- a.docx (начало файла) ---" in previews
    assert "--- b.docx (начало файла) ---" in previews
    a_block = previews.split("--- b.docx")[0]
    assert a_block.count("A") == 1500
    b_block = previews.split("--- b.docx")[1]
    assert b_block.count("B") == 1500


def _previews_for(archives: int, per_archive: int) -> str:
    chunks = [
        Chunk(
            id=a * per_archive + i,
            doc_name=f"архив{a}.zip/f{i}.docx",
            title="t",
            text="X" * 2000,
            tokens=frozenset(),
        )
        for a in range(archives)
        for i in range(per_archive)
    ]
    return _document_previews(chunks)


def test_document_previews_size_stops_growing_with_file_count():
    """The previous per-file budget had a 300-char floor, so past ~200 files it
    grew linearly again — 346k chars at 1000 files, worse than the blowup it was
    meant to fix. Previews are per archive now, so the block must plateau."""
    small = _previews_for(archives=10, per_archive=5)      # 50 files
    large = _previews_for(archives=10, per_archive=500)    # 5000 files, same archives

    # 100x the files may only move the block by the width of the counts printed
    # in each archive header — nothing proportional to the file count.
    assert len(large) < len(small) * 1.2, (
        "preview size must depend on archive count, not file count: "
        f"{len(small)} chars for 50 files vs {len(large)} for 5000"
    )


def test_document_previews_ceiling_holds_at_any_scale():
    huge = _previews_for(archives=200, per_archive=25)  # 5000 files across 200 archives

    # A fixed ceiling, not one that grows with the input: preview bodies are
    # capped by PLANNER_PREVIEW_BUDGET and only PLANNER_PREVIEW_GROUPS archives
    # get one, so headers and name lists add bounded overhead on top.
    assert len(huge) < PLANNER_PREVIEW_BUDGET * 2


def test_document_previews_names_every_archive_even_without_a_preview():
    previews = _previews_for(archives=PLANNER_PREVIEW_GROUPS + 5, per_archive=2)

    for a in range(PLANNER_PREVIEW_GROUPS + 5):
        assert f"архив{a}.zip" in previews, f"архив{a}.zip missing from previews entirely"


# Task 15 — «было → стало»: find repeating requisites, parse a filled reply,
# apply it deterministically.

def test_parse_replacements_accepts_arrow_separator():
    assert dg._parse_replacements("АБВГ.123456.789 → АБВГ.999999.001") == {
        "АБВГ.123456.789": "АБВГ.999999.001"
    }


def test_parse_replacements_accepts_ascii_arrow_separator():
    assert dg._parse_replacements("Иванов И.И. -> Петров П.П.") == {"Иванов И.И.": "Петров П.П."}


def test_parse_replacements_accepts_coloneq_separator():
    # ::= is what app/api/documents.py's /edit-docx norm-control endpoint
    # already parses — same idea, same syntax.
    assert dg._parse_replacements("ООО «Ромашка» ::= ООО «Новый заказчик»") == {
        "ООО «Ромашка»": "ООО «Новый заказчик»"
    }


def test_parse_replacements_strips_label_prefix():
    assert dg._parse_replacements("Заказчик: ООО «Ромашка» → ООО «Новый заказчик»") == {
        "ООО «Ромашка»": "ООО «Новый заказчик»"
    }


def test_parse_replacements_empty_right_side_means_placeholder():
    assert dg._parse_replacements("Заказчик: ООО «Ромашка» → ") == {"ООО «Ромашка»": ""}


def test_parse_replacements_kak_est_drops_the_pair():
    assert dg._parse_replacements("Заказчик: ООО «Ромашка» → как есть") == {}


def test_replacement_table_kak_est_is_still_consent_to_start():
    """A table of «как есть» used to look identical to ordinary prose
    (_parse_replacements is {} for both), so the mode re-asked confirmation
    instead of generating. The table shape itself is consent."""
    mapping, is_table = dg._replacement_table(
        "Заказчик: ООО «Ромашка» → как есть\n"
        "Децимальный номер: АБВГ.123456.789 → как есть"
    )
    assert mapping == {}
    assert is_table is True


def test_replacement_table_prose_and_lone_number_are_not_consent():
    for prose in (
        "Просто обычное сообщение без разделителей.",
        "я думаю -> надо переделать раздел",
        "100000 -> 120000",
    ):
        mapping, is_table = dg._replacement_table(prose)
        assert mapping == {}
        assert is_table is False, f"prose counted as a replacement table: {prose!r}"


def test_marker_button_takes_empty_cells_from_the_confirmation_as_placeholders():
    confirmation = (
        "Разделов: 2, примерно 3 стр.\n"
        "Заказчик: ООО «Ромашка» → \n"
        "Децимальный номер: АБВГ.123456.789 → \n"
    )
    messages = [
        {"role": "user", "content": "сделай документ"},
        {"role": "assistant", "content": confirmation},
        {"role": "user", "content": f"Уточнения по задаче:\n1. Начинать генерацию? – {dg._CONFIRM_START}"},
    ]
    offered = dg._replacements_offered_in(messages)
    assert offered == {"ООО «Ромашка»": "", "АБВГ.123456.789": ""}
    swept = dg._apply_replacements("Заказчик ООО «Ромашка», номер АБВГ.123456.789.", offered)
    assert "ООО «Ромашка»" not in swept
    assert "АБВГ.123456.789" not in swept
    assert "[УКАЗАТЬ:" in swept


def test_replacements_offered_in_uses_only_the_last_assistant_message():
    messages = [
        {"role": "assistant", "content": "Заказчик: ООО «Ромашка» → \n"},
        {"role": "user", "content": "другой запрос"},
        {"role": "assistant", "content": "Разделов: 2, примерно 3 стр. Шаблон оформления не определён."},
    ]
    assert dg._replacements_offered_in(messages) == {}


def test_parse_replacements_no_separator_returns_empty_dict():
    # Not mistaken for a confirmation: Task 12's fail-safe gate relies on
    # this returning falsy for an ordinary message.
    assert dg._parse_replacements("Просто обычное сообщение без разделителей.") == {}


def test_parse_replacements_ignores_prose_that_merely_contains_an_arrow():
    """A non-empty mapping is taken as consent to start generating, so prose
    with a stray arrow must not produce one — otherwise an ordinary chat
    message costs thousands of model calls. The left side has to look like a
    requisite, the same shapes the table offered in the first place."""
    for prose in (
        "я думаю -> надо переделать раздел",
        "давай так: сначала план -> потом текст",
        "сроки сдвинулись ::= перенеси на май",
    ):
        assert dg._parse_replacements(prose) == {}, f"prose parsed as a replacement list: {prose!r}"


def test_parse_replacements_lone_number_line_is_not_consent():
    """A digits-only pair matches the document-number shape, but «100000 ->
    120000» is also how a budget or a deadline gets discussed in passing. On
    its own it is not the filled-in table, so it must not read as consent."""
    for prose in (
        "100000 -> 120000",
        "смета:\n100000 -> 120000",
        "бюджет 100000 -> 120000 и сроки",
    ):
        assert dg._parse_replacements(prose) == {}, f"lone number parsed as consent: {prose!r}"

    # The same line inside a real list still parses: one recognisable
    # neighbour is enough to tell the table from the aside.
    assert dg._parse_replacements(
        "100000 -> 120000\nЗаказчик: ООО «Ромашка» → ООО «Вектор»"
    ) == {"100000": "120000", "ООО «Ромашка»": "ООО «Вектор»"}

    # A lone pair with a recognisable shape is a table of one, not an aside.
    assert dg._parse_replacements("АБВГ.123456.789 -> ЖЗИК.987654.321") == {
        "АБВГ.123456.789": "ЖЗИК.987654.321"
    }


def test_requested_sections_reads_the_page_target():
    """Without this the planner sizes the document by feel — a request for
    5000 pages produced a few dozen sections and the user never learned why."""
    assert dg._requested_sections("напиши документацию на 5000 страниц") == 3334
    assert dg._requested_sections("документ на 5 тыс. страниц") == 3334
    assert dg._requested_sections("7к страниц") == 4667
    assert dg._requested_sections("сделай 300 разделов") == 300


def test_requested_sections_takes_the_upper_bound_of_a_range():
    # «5-7 тысяч страниц» — under-delivering on a stated range is the worse
    # error here: the run is confirmed before it starts, so too big is visible.
    assert dg._requested_sections("нужен документ 5-7 тысяч страниц") == 4667


def test_requested_sections_needs_a_unit_word():
    assert dg._requested_sections("пояснительная записка по ГОСТ 19.201-78") is None
    assert dg._requested_sections("сгенерируй документацию по этой базе") is None
    assert dg._requested_sections("страница 5 шаблона важна") is None


def test_requested_sections_clamped_to_max_sections():
    assert dg._requested_sections("нужно 100000 страниц") == dg.MAX_SECTIONS


def test_requested_documents_reads_file_count_not_pages():
    assert dg._requested_documents("сделай 1000 документов по шаблонам") == 1000
    assert dg._requested_documents("нужно 10 документов") == 10
    assert dg._requested_documents("2к документов") == 2000
    # «документацию» is one volume, not N files.
    assert dg._requested_documents("напиши документацию на 5000 страниц") is None
    assert dg._requested_documents("сделай ИТТ по шаблону") is None


def test_requested_documents_clamped_to_max_documents():
    assert dg._requested_documents("миллион документов") is None
    assert dg._requested_documents("100000 документов") == dg.MAX_DOCUMENTS


def test_assign_templates_matches_by_name_uniquely():
    titles = ["ИТТ на ПЛК", "ИТТ на насос"]
    templates = {"шаблоны.zip/ИТТ_ПЛК.docx", "шаблоны.zip/ИТТ_насос.docx"}
    assigned = dg._assign_templates(titles, templates)
    assert assigned["ИТТ на ПЛК"].endswith("ИТТ_ПЛК.docx")
    assert assigned["ИТТ на насос"].endswith("ИТТ_насос.docx")
    assert assigned["ИТТ на ПЛК"] != assigned["ИТТ на насос"]


def test_assign_templates_one_template_is_shared():
    assigned = dg._assign_templates(["А", "Б"], {"форма.docx"})
    assert assigned["А"] == assigned["Б"] == "форма.docx"


def test_assign_templates_matches_by_digest_when_names_differ():
    """forma1/forma2 say nothing; the start of each file does."""
    titles = ["ИТТ на ПЛК", "ИТТ на насос"]
    templates = {"шаблоны.zip/форма1.docx", "шаблоны.zip/форма2.docx"}
    digests = {
        "шаблоны.zip/форма1.docx": "Исходные требования на программируемый контроллер ПЛК Siemens",
        "шаблоны.zip/форма2.docx": "Исходные требования на центробежный насос насосной станции",
    }
    title_texts = {
        "ИТТ на ПЛК": "контроллер ПЛК ток пять ампер Siemens",
        "ИТТ на насос": "центробежный насос напор двенадцать метров",
    }
    assigned = dg._assign_templates(titles, templates, digests=digests, title_texts=title_texts)
    assert assigned["ИТТ на ПЛК"].endswith("форма1.docx")
    assert assigned["ИТТ на насос"].endswith("форма2.docx")


def test_scope_chunks_uses_digest_when_filenames_are_generic():
    chunks = [
        Chunk(id=0, doc_name="шаблоны.zip/форма1.docx", title="ИТТ",
              text="требования на контроллер ПЛК Siemens",
              tokens=frozenset({"требования", "контроллер", "плк", "siemens"}),
              title_tokens=frozenset({"итт"})),
        Chunk(id=1, doc_name="шаблоны.zip/форма2.docx", title="ИТТ",
              text="требования на центробежный насос станции",
              tokens=frozenset({"требования", "центробежный", "насос", "станции"}),
              title_tokens=frozenset({"итт"})),
        Chunk(id=2, doc_name="знания.zip/данные1.xlsx", title="данные",
              text="контроллер ПЛК Siemens ток пять ампер",
              tokens=frozenset({"контроллер", "плк", "siemens", "ток", "пять", "ампер"}),
              title_tokens=frozenset({"данные"})),
        Chunk(id=3, doc_name="знания.zip/данные2.xlsx", title="данные",
              text="центробежный насос напор двенадцать метров",
              tokens=frozenset({"центробежный", "насос", "напор", "двенадцать", "метров"}),
              title_tokens=frozenset({"данные"})),
    ]
    templates = {"шаблоны.zip/форма1.docx", "шаблоны.zip/форма2.docx"}
    digests = dg._file_digests(chunks)
    assigned = dg._assign_templates(
        ["ИТТ на ПЛК", "ИТТ на насос"], templates,
        digests=digests,
        title_texts={
            "ИТТ на ПЛК": digests["знания.zip/данные1.xlsx"],
            "ИТТ на насос": digests["знания.zip/данные2.xlsx"],
        },
    )
    section = Section(
        id=0, title="Требования", brief="ток", complexity="simple", document="ИТТ на ПЛК",
    )
    scoped, tmpl_set = dg._scope_chunks_for_section(
        section, chunks, templates, assigned, digests=digests,
    )
    names = {c.doc_name for c in scoped}
    assert assigned["ИТТ на ПЛК"].endswith("форма1.docx")
    assert "шаблоны.zip/форма1.docx" in names
    assert "шаблоны.zip/форма2.docx" not in names
    assert "знания.zip/данные1.xlsx" in names
    assert "знания.zip/данные2.xlsx" not in names
    assert tmpl_set == {"шаблоны.zip/форма1.docx"}


def test_file_digests_take_title_and_start_without_a_model():
    chunks = [
        Chunk(id=0, doc_name="a.docx", title="Контроллер ПЛК", text="A" * 2000, tokens=frozenset()),
        Chunk(id=1, doc_name="a.docx", title="хвост", text="B" * 2000, tokens=frozenset()),
        Chunk(id=2, doc_name="b.docx", title="Насос", text="напор", tokens=frozenset()),
    ]
    digests = dg._file_digests(chunks)
    assert "Контроллер ПЛК" in digests["a.docx"]
    assert "A" in digests["a.docx"]
    assert "B" not in digests["a.docx"], "digest must stay at the start of the file"
    assert len(digests["a.docx"]) <= dg.DIGEST_CHARS
    assert "Насос" in digests["b.docx"]


def test_seed_document_titles_from_matching_knowledge_count():
    knowledge = [f"знания.zip/узел_{i}.xlsx" for i in range(10)]
    templates = {f"шаблоны.zip/форма_{i}.docx" for i in range(10)}
    titles = dg._seed_document_titles(10, knowledge, templates)
    assert titles is not None
    assert len(titles) == 10
    assert "узел 0" in titles[0][0]
    assert titles[0][1].endswith("узел_0.xlsx")


def test_seed_document_titles_none_when_counts_do_not_match():
    knowledge = [f"знания.zip/узел_{i}.xlsx" for i in range(3)]
    templates = {f"шаблоны.zip/форма_{i}.docx" for i in range(2)}
    assert dg._seed_document_titles(10, knowledge, templates) is None


def test_scope_chunks_keeps_only_related_knowledge_and_own_template():
    chunks = [
        Chunk(id=0, doc_name="шаблоны.zip/ПЛК.docx", title="t",
              text="структура ПЛК", tokens=frozenset({"структура"}), title_tokens=frozenset()),
        Chunk(id=1, doc_name="шаблоны.zip/насос.docx", title="t",
              text="структура насоса", tokens=frozenset({"структура"}), title_tokens=frozenset()),
        Chunk(id=2, doc_name="знания.zip/ПЛК.xlsx", title="t",
              text="ток 5А", tokens=frozenset({"ток"}), title_tokens=frozenset()),
        Chunk(id=3, doc_name="знания.zip/насос.xlsx", title="t",
              text="напор 10 м", tokens=frozenset({"напор"}), title_tokens=frozenset()),
    ]
    templates = {"шаблоны.zip/ПЛК.docx", "шаблоны.zip/насос.docx"}
    assigned = dg._assign_templates(["ИТТ на ПЛК", "ИТТ на насос"], templates)
    section = Section(id=0, title="Требования", brief="ток", complexity="simple", document="ИТТ на ПЛК")
    scoped, tmpl_set = dg._scope_chunks_for_section(section, chunks, templates, assigned)
    names = {c.doc_name for c in scoped}
    assert "шаблоны.zip/ПЛК.docx" in names
    assert "шаблоны.zip/насос.docx" not in names
    assert "знания.zip/ПЛК.xlsx" in names
    assert "знания.zip/насос.xlsx" not in names
    assert tmpl_set == {"шаблоны.zip/ПЛК.docx"}


def test_section_context_does_not_mix_foreign_knowledge():
    chunks = [
        Chunk(id=0, doc_name="шаблоны.zip/ПЛК.docx", title="ПЛК",
              text="форма ПЛК", tokens=frozenset({"форма", "плк"}), title_tokens=frozenset({"плк"})),
        Chunk(id=1, doc_name="знания.zip/ПЛК.xlsx", title="ПЛК",
              text="ток пять ампер уникальный факт плк",
              tokens=frozenset({"ток", "пять", "ампер", "уникальный", "факт", "плк"}),
              title_tokens=frozenset({"плк"})),
        Chunk(id=2, doc_name="знания.zip/насос.xlsx", title="насос",
              text="напор двенадцать метров чужой факт насос",
              tokens=frozenset({"напор", "двенадцать", "метров", "чужой", "факт", "насос"}),
              title_tokens=frozenset({"насос"})),
    ]
    templates = {"шаблоны.zip/ПЛК.docx"}
    assigned = {"ИТТ на ПЛК": "шаблоны.zip/ПЛК.docx"}
    section = Section(id=0, title="Требования к ПЛК", brief="ток", complexity="simple", document="ИТТ на ПЛК")
    block = _section_context(section, chunks, templates, assigned_templates=assigned)
    assert "ток пять ампер" in block
    assert "напор двенадцать метров" not in block


def test_plan_seeds_one_file_per_matching_archive_member():
    """A 10-document order with 10 knowledge files must not wait on the planner
    to invent ten titles — that is how a 1000-file archive used to collapse."""
    import asyncio as _asyncio

    chunks = []
    for i in range(10):
        chunks.append(Chunk(
            id=i, doc_name=f"шаблоны.zip/узел_{i}.docx", title="t",
            text="форма", tokens=frozenset(), title_tokens=frozenset(),
        ))
        chunks.append(Chunk(
            id=100 + i, doc_name=f"знания.zip/узел_{i}.xlsx", title="t",
            text="цифры", tokens=frozenset(), title_tokens=frozenset(),
        ))

    called = []

    async def fake_plan(*args, **kwargs):
        called.append(1)
        return [dg.Section(id=0, title="не должно", brief="", complexity="simple")], set(), False

    orig = dg._plan_outline
    dg._plan_outline = fake_plan
    try:
        sections, templates, failed = _asyncio.run(
            dg._plan_document("сделай 10 документов по шаблонам", chunks, 1)
        )
    finally:
        dg._plan_outline = orig

    assert not failed
    assert not called, "matching archive members must seed the plan without the LLM"
    assert len({s.document for s in sections}) == 10
    assert len(templates) == 10


def test_template_formatting_is_not_overridden_when_template_is_used():
    """apply_paragraph_formatting centres headings and adds a 1.25cm first-line
    indent — GOST templates already define both, and forcing them is how the
    generated files stopped looking like the uploaded form."""
    from docx import Document as _Document
    from docx_generator import blank_copy_of_template, convert_markdown_to_docx

    source = _Document()
    source.sections[0].header.paragraphs[0].text = "HEADER_PLC"
    buf = io.BytesIO()
    source.save(buf)
    blank = blank_copy_of_template(buf.getvalue())
    with_template = convert_markdown_to_docx(
        "## Заголовок\n\nОбычный абзац текста документа.", base_template_bytes=blank,
    )
    out = _Document(io.BytesIO(with_template))
    assert "HEADER_PLC" in out.sections[0].header.paragraphs[0].text
    body = [p for p in out.paragraphs if p.text.strip() == "Обычный абзац текста документа."]
    assert body
    indent = body[0].paragraph_format.first_line_indent
    assert indent is None or indent == 0, f"template body got a forced first-line indent: {indent}"

    without = convert_markdown_to_docx("## Заголовок\n\nОбычный абзац текста документа.")
    plain = _Document(io.BytesIO(without))
    plain_body = [p for p in plain.paragraphs if p.text.strip() == "Обычный абзац текста документа."]
    assert plain_body
    # Default path still applies a first-line indent; the exact EMU value is
    # the 1.25cm we set plus whatever htmldocx already wrote.
    assert plain_body[0].paragraph_format.first_line_indent


def test_deepseek_is_what_keeps_the_run_inside_the_hundred_dollar_budget():
    """The mode's stated budget is 5-7 thousand pages under $100. With DeepSeek
    on the escalated sections the worst case fits with room to spare; without
    it, the same worst case does not — which is exactly why the confirmation
    screen warns when the key is missing instead of quoting a number and
    hoping. Both halves are asserted so neither can drift unnoticed.
    """
    outline = [
        dg.Section(id=i, title=f"Раздел {i}", brief="", complexity="complex")
        for i in range(dg._requested_sections("7000 страниц"))
    ]
    original = dg.DEEPSEEK_API_KEY
    try:
        dg.DEEPSEEK_API_KEY = "sk-present"
        with_deepseek = dg._estimate_cost_usd(outline, has_sources=True)
        dg.DEEPSEEK_API_KEY = ""
        without_deepseek = dg._estimate_cost_usd(outline, has_sources=True)
    finally:
        dg.DEEPSEEK_API_KEY = original

    assert with_deepseek < 100, f"budget blown even with DeepSeek: ${with_deepseek:.2f}"
    assert without_deepseek > with_deepseek * 2, "DeepSeek must be the cheap path, not a rounding difference"


def test_estimate_cost_grows_with_escalation_and_size():
    simple = [dg.Section(id=i, title="x", brief="", complexity="simple") for i in range(100)]
    complex_ = [dg.Section(id=i, title="x", brief="", complexity="complex") for i in range(100)]
    assert dg._estimate_cost_usd(complex_, True) > dg._estimate_cost_usd(simple, True)
    assert dg._estimate_cost_usd(simple, True) > dg._estimate_cost_usd(simple, False)
    assert dg._estimate_cost_usd(simple[:50], True) < dg._estimate_cost_usd(simple, True)


def test_chapter_count_never_exceeds_max_sections():
    # MAX_CHAPTERS x SECTIONS_PER_CHAPTER must stay equal to MAX_SECTIONS:
    # a drift here either caps the mode below its promise or blows past the
    # ceiling that keeps a run from lasting days.
    assert dg.MAX_CHAPTERS * dg.SECTIONS_PER_CHAPTER == dg.MAX_SECTIONS


def test_runaway_chapter_plan_does_not_fire_an_expansion_call_per_chapter():
    """A planner that answers with far more chapters than it was asked for
    used to cost one expansion call each, and the sections beyond the ceiling
    were then thrown away anyway — money spent on output nobody sees."""
    import asyncio as _asyncio

    expansions = []

    async def fake_plan(user_text, chunks, user_id, extra_instruction="", want_count=None, as_chapters=False):
        assert as_chapters and want_count
        chapters = [
            dg.Section(id=i, title=f"Глава {i}", brief="", complexity="simple")
            for i in range(want_count * 10)  # runaway: ten times what was asked
        ]
        return chapters, set(), False

    async def fake_expand(chapter, user_text, catalog, per_chapter, user_id):
        expansions.append(chapter.title)
        return [dg.Section(id=0, title=f"{chapter.title}.1", brief="", complexity="simple")]

    orig_plan, orig_expand = dg._plan_outline, dg._expand_chapter
    dg._plan_outline, dg._expand_chapter = fake_plan, fake_expand
    try:
        sections, _, failed = _asyncio.run(
            dg._plan_document("документ на 5000 страниц", [], 1)
        )
    finally:
        dg._plan_outline, dg._expand_chapter = orig_plan, orig_expand

    assert not failed
    expected_chapters = -(-dg._requested_sections("5000 страниц") // dg.SECTIONS_PER_CHAPTER)
    assert len(expansions) == expected_chapters, (
        f"expanded {len(expansions)} chapters, asked for {expected_chapters}"
    )


def test_short_chapter_plan_still_reaches_the_requested_size():
    """The planner usually returns fewer chapters than it was asked for. With
    a hard 25-sections-per-chapter ceiling that alone would turn a 3350-section
    order into 750 — so a short chapter list makes each chapter carry more."""
    import asyncio as _asyncio

    asked_per_chapter = []

    async def fake_plan(user_text, chunks, user_id, extra_instruction="", want_count=None, as_chapters=False):
        # A third of what was asked for — a realistic shortfall, not a failure.
        chapters = [
            dg.Section(id=i, title=f"Глава {i}", brief="", complexity="simple")
            for i in range(max(1, want_count // 3))
        ]
        return chapters, set(), False

    async def fake_expand(chapter, user_text, catalog, per_chapter, user_id):
        asked_per_chapter.append(per_chapter)
        return [
            dg.Section(id=0, title=f"{chapter.title}.{i}", brief="", complexity="simple")
            for i in range(per_chapter)
        ]

    orig_plan, orig_expand = dg._plan_outline, dg._expand_chapter
    dg._plan_outline, dg._expand_chapter = fake_plan, fake_expand
    try:
        sections, _, failed = _asyncio.run(dg._plan_document("документ на 5000 страниц", [], 1))
    finally:
        dg._plan_outline, dg._expand_chapter = orig_plan, orig_expand

    target = dg._requested_sections("5000 страниц")
    assert not failed
    assert max(asked_per_chapter) > dg.SECTIONS_PER_CHAPTER, "a short chapter list must widen each chapter"
    assert max(asked_per_chapter) <= dg.MAX_SECTIONS_PER_EXPANSION
    assert len(sections) >= target * 0.9, f"order shortfall: {len(sections)} of {target}"


def test_expand_chapter_keeps_the_document_name():
    """A two-level plan for «10 отдельных ИТТ на 5000 страниц» decides the
    file split on the chapter list. Expansion JSON has no `document` field,
    so without copying it from the chapter every file collapses into one."""
    import asyncio as _asyncio

    async def fake_chat(messages, model=None, user_id=None, use_tools=False, use_skills=False, **kwargs):
        return '{"sections": [{"title": "Общие сведения", "brief": "состав", "complexity": "simple"}]}', [], "", []

    orig = dg.get_chat_response
    dg.get_chat_response = fake_chat
    try:
        chapter = dg.Section(
            id=0, title="ИТТ на ПЛК", brief="требования к ПЛК",
            complexity="simple", document="ИТТ на ПЛК",
        )
        sections = _asyncio.run(dg._expand_chapter(chapter, "сделай 10 ИТТ на 5000 страниц", "", 3, 1))
    finally:
        dg.get_chat_response = orig

    assert sections
    assert all(s.document == "ИТТ на ПЛК" for s in sections)
    assert all(s.chapter == "ИТТ на ПЛК" for s in sections)


def test_progress_line_counts_pages_from_real_text_not_from_the_plan():
    """A run lasts hours and the user ordered pages, not sections. The pages
    shown have to come from what was actually written: a model writing shorter
    than planned must be visible straight away, not at the end."""
    # Half the sections done, but each one came out half the expected length.
    written = 50 * (dg.CHUNK_TARGET_CHARS // 2)
    line = dg._progress_line(done=50, total=100, chars_written=written)
    assert "раздел 50 из 100" in line
    assert "50%" in line
    assert f"готово ~{written // dg.CHARS_PER_PAGE} стр." in line
    # The target still reflects the plan, so the shortfall is visible as a gap.
    assert f"из ~{100 * dg.CHUNK_TARGET_CHARS // dg.CHARS_PER_PAGE}" in line


def test_progress_bar_fills_and_never_overflows():
    assert dg._progress_bar(0, 10) == "▱" * 10
    assert dg._progress_bar(10, 10) == "▰" * 10
    assert len(dg._progress_bar(3, 10)) == 10
    # Degenerate inputs must not raise or produce a ragged bar — this runs
    # inside the status update of a multi-hour job.
    assert len(dg._progress_bar(0, 0)) == 10
    assert len(dg._progress_bar(99, 10)) == 10


def test_estimate_cost_covers_the_model_writing_longer_than_planned():
    """Measured on the server: a 6-page order came back 11 pages. The estimate
    says «не больше», so it has to hold when the model overshoots — otherwise
    the budget promise breaks on a successful run, which is the worst time."""
    outline = [dg.Section(id=i, title="x", brief="", complexity="complex") for i in range(100)]
    estimate = dg._estimate_cost_usd(outline, has_sources=True)

    from app.billing.costs import model_cost_usd
    from model_context import estimate_tokens
    # What a run actually costs if every section comes out 1.8x the plan.
    real_input = estimate_tokens("x" * (dg.MAX_CHUNK_CHARS_PER_SECTION + 1200))
    real_output = estimate_tokens("x" * int(dg.CHUNK_TARGET_CHARS * 1.8))
    escalated = "deepseek-v4-pro" if dg.DEEPSEEK_API_KEY else dg.ESCALATED_WRITER_MODEL
    actual = 100 * model_cost_usd(escalated, real_input, real_output)
    assert estimate >= actual, f"estimate ${estimate:.2f} under-promises the real ${actual:.2f}"


def _sections(*pairs):
    return [
        dg.Section(id=i, title=title, brief="", complexity="simple", document=document)
        for i, (title, document) in enumerate(pairs)
    ]


def test_sections_without_a_document_stay_one_file():
    """The old behaviour, and still the common case: nothing asked for several
    documents, so nothing gets split."""
    outline = _sections(("Раздел 1", ""), ("Раздел 2", ""))
    groups = dg._group_by_document(outline, ["текст один", "текст два"])
    assert len(groups) == 1
    assert groups[0][0] == ""
    assert "текст один" in groups[0][1] and "текст два" in groups[0][1]


def test_each_document_becomes_its_own_file():
    outline = _sections(("Общие сведения", "ИТТ на ПЛК"), ("Требования", "ИТТ на ПЛК"),
                        ("Общие сведения", "ИТТ на шкаф"))
    groups = dg._group_by_document(outline, ["а", "б", "в"])
    assert [title for title, _ in groups] == ["ИТТ на ПЛК", "ИТТ на шкаф"]
    assert "а" in groups[0][1] and "б" in groups[0][1]
    assert "в" in groups[1][1] and "а" not in groups[1][1]


def test_interleaved_sections_do_not_split_a_document_in_two():
    """The planner is asked to keep a document's sections together, but it is a
    model. Grouping by position would turn ten documents into twenty files,
    half of them named «… (2)» — worse than useless to the customer."""
    outline = _sections(("Р1", "Документ А"), ("Р1", "Документ Б"), ("Р2", "Документ А"))
    groups = dg._group_by_document(outline, ["а1", "б1", "а2"])
    assert len(groups) == 2, f"interleaving split a document: {[t for t, _ in groups]}"
    assert "а1" in groups[0][1] and "а2" in groups[0][1]


def test_document_filename_is_safe_and_unique():
    used = set()
    assert dg._document_filename("ИТТ на ПЛК", used) == "ИТТ на ПЛК.docx"
    # Path separators in a model-supplied title must not escape into a path.
    assert dg._document_filename("узел/подузел: часть 1", used) == "узел подузел часть 1.docx"
    assert dg._document_filename("ИТТ на ПЛК", used) == "ИТТ на ПЛК (2).docx"
    assert dg._document_filename("", used) == "Документ.docx"


def test_many_documents_are_packed_into_one_archive():
    """Ten separate attachments are ten cards to download one by one; the user
    asked for an archive."""
    import zipfile
    files = [{"filename": f"Док {i}.docx", "bytes": b"PK-fake-%d" % i} for i in range(3)]
    packed = dg._zip_documents(files)
    assert packed["filename"].endswith(".zip")
    with zipfile.ZipFile(io.BytesIO(packed["bytes"])) as archive:
        assert archive.namelist() == ["Док 0.docx", "Док 1.docx", "Док 2.docx"]
        assert archive.read("Док 1.docx") == b"PK-fake-1"


def _docx_with_landscape_appendix() -> bytes:
    from docx import Document as _Document
    from docx.enum.section import WD_ORIENT

    doc = _Document()
    doc.add_paragraph("Основная часть")
    appendix = doc.add_section()
    appendix.orientation = WD_ORIENT.LANDSCAPE
    appendix.page_width, appendix.page_height = appendix.page_height, appendix.page_width
    doc.add_paragraph("Приложение")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_template_page_setup_comes_from_the_first_section():
    """A template whose last section is a landscape appendix used to make the
    whole generated document landscape — the final sectPr is the one python-docx
    hands over, and it belongs to the appendix, not to the body."""
    from docx import Document as _Document
    from docx.enum.section import WD_ORIENT
    from docx_generator import blank_copy_of_template

    blank = blank_copy_of_template(_docx_with_landscape_appendix())
    sections = _Document(io.BytesIO(blank)).sections
    assert len(sections) == 1
    assert sections[0].orientation == WD_ORIENT.PORTRAIT, "inherited the appendix, not the body"


def test_template_without_heading_style_still_renders_formatted_text():
    """The catastrophic case: htmldocx raises on a missing style, and
    convert_markdown_to_docx then writes RAW markdown into the document and
    returns it as a valid .docx — so the customer gets «## Заголовок» as body
    text while the summary claims the template was applied."""
    import re as _re
    import zipfile as _zipfile
    from docx import Document as _Document
    from docx_generator import (
        blank_copy_of_template, convert_markdown_to_docx, CONVERSION_FAILED_MARKER,
    )

    source = io.BytesIO()
    _Document().save(source)
    stripped = io.BytesIO()
    with _zipfile.ZipFile(io.BytesIO(source.getvalue())) as zin, \
            _zipfile.ZipFile(stripped, "w", _zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/styles.xml":
                data = _re.sub(
                    r'<w:style [^>]*w:styleId="Heading2".*?</w:style>', "",
                    data.decode("utf-8"), flags=_re.S,
                ).encode("utf-8")
            zout.writestr(item, data)

    blank = blank_copy_of_template(stripped.getvalue())
    result = convert_markdown_to_docx("## Заголовок\n\nТекст.", base_template_bytes=blank)
    paragraphs = [(p.style.name, p.text.strip()) for p in _Document(io.BytesIO(result)).paragraphs
                  if p.text.strip()]
    assert not any(CONVERSION_FAILED_MARKER in text for _, text in paragraphs), paragraphs
    assert ("Heading 2", "Заголовок") in paragraphs, paragraphs


def test_unrenderable_template_is_refused_rather_than_shipped():
    """If a template still cannot render, standard formatting is the right
    answer — ten documents of raw markdown is not."""
    assert dg._template_renders_cleanly(b"not a docx at all") is False


def test_knowledge_archive_is_not_mistaken_for_templates():
    """«База ин-ФОРМА-ции для наполнения.zip» — the customer's real archive
    name — matched the template stem «форма» as a substring, so the entire
    knowledge base was classified as formatting samples. Every section then
    wrote with zero facts, silently, while the template block looked fine."""
    knowledge = "База информации для наполнения.zip/узел1.docx"
    template = "Пример оформления.zip/ИТТ_ледокол.docx"
    assert not dg._TEMPLATE_NAME_RE.search(knowledge), "knowledge archive read as a template"
    assert dg._TEMPLATE_NAME_RE.search(template), "real template no longer recognised"

    chunks = [
        dg.Chunk(id=0, doc_name=template, title="", text="ИСХОДНЫЕ ТЕХНИЧЕСКИЕ ТРЕБОВАНИЯ"),
        dg.Chunk(id=1, doc_name=knowledge, title="", text="ПЛК мощность 2500 кВт"),
    ]
    assert dg._heuristic_template_names(chunks) == {template}


def test_template_stems_cover_plural_folder_names():
    # «Образцы» / «Формы» are at least as common as the singular forms on a
    # customer's disk, and neither contains the singular stem.
    for name in ("Образцы.zip/f.docx", "Формы.zip/f.docx", "ГОСТ формы.zip/f.docx",
                 "Шаблоны.zip/f.docx", "Бланки.zip/f.docx"):
        assert dg._TEMPLATE_NAME_RE.search(name), name


def test_weak_knowledge_match_does_not_drop_the_whole_base():
    """The cutoff was non-monotonic: a best score of 13 against a floor of 16
    emptied the knowledge set, so the section saw only its template — while a
    section with NO match at all kept every file. A hint about the topic made
    the outcome strictly worse."""
    def chunk(i, doc, text):
        return dg.Chunk(id=i, doc_name=doc, title="", text=text,
                        tokens=frozenset(dg._tokenize(text)), title_tokens=frozenset())

    template = "f1.docx"
    chunks = [
        chunk(0, template, "Общие положения требования раздел"),
        chunk(1, "d1.xlsx", "подача 120 кубометров напор 45 метров агрегат"),
        chunk(2, "d2.xlsx", "мощность 2500 киловатт частота 50 герц"),
        chunk(3, "d3.xlsx", "температура 80 градусов давление 16 бар"),
    ]
    digests = dg._file_digests(chunks)
    section = dg.Section(id=0, title="Сведения", brief="агрегат",
                         complexity="simple", document="Первый")
    scoped, _ = dg._scope_chunks_for_section(
        section, chunks, {template}, {"Первый": template}, digests,
    )
    names = {c.doc_name for c in scoped}
    assert "d1.xlsx" in names, f"the best-matching knowledge file was dropped: {sorted(names)}"


def test_plain_start_and_marker_start_are_different_orders():
    """Two intentions, two buttons. Pressing the primary «Да, начинай» to see a
    result must not silently replace twenty-five requisites with [УКАЗАТЬ: …]
    across the finished files; asking for the markers is a separate click."""
    assert dg._CONFIRM_START != dg._CONFIRM_START_PLACEHOLDERS
    # Neither label may contain the other, or the gate's substring checks
    # would read one click as the other.
    assert dg._CONFIRM_START not in dg._CONFIRM_START_PLACEHOLDERS
    assert dg._CONFIRM_START_PLACEHOLDERS not in dg._CONFIRM_START
    # Both must fit the clarify chip limit, or normalize_questions drops the
    # whole question and the confirmation card loses its start button.
    from clarify import MAX_OPTION_CHARS

    assert len(dg._CONFIRM_START_PLACEHOLDERS) <= MAX_OPTION_CHARS


def test_hand_formatted_template_still_sets_the_font():
    """Engineering examples are usually formatted by hand — select all, Times
    New Roman 14 — not through styles. That formatting lives on the runs and
    disappears with them when the body is cleared, so the generated document
    came out in the default font while the example was set by ГОСТ. The
    dominant direct format is promoted into Normal before the body goes."""
    from docx import Document as _Document
    from docx.shared import Pt
    from docx_generator import blank_copy_of_template, convert_markdown_to_docx

    doc = _Document()
    for text in ("Первый абзац примера.", "Второй абзац примера."):
        paragraph = doc.add_paragraph(text)
        paragraph.runs[0].font.name = "Times New Roman"
        paragraph.runs[0].font.size = Pt(14)
    source = io.BytesIO()
    doc.save(source)

    blank = blank_copy_of_template(source.getvalue())
    normal = _Document(io.BytesIO(blank)).styles["Normal"]
    assert normal.font.name == "Times New Roman", normal.font.name
    assert normal.font.size == Pt(14), normal.font.size

    result = _Document(io.BytesIO(convert_markdown_to_docx("Текст.", base_template_bytes=blank)))
    assert result.styles["Normal"].font.name == "Times New Roman"


def test_template_own_style_wins_over_measured_formatting():
    """A template that defines Normal properly must keep it: the measurement
    is a fallback for hand formatting, not a second opinion."""
    from docx import Document as _Document
    from docx.shared import Pt
    from docx_generator import blank_copy_of_template

    doc = _Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)
    paragraph = doc.add_paragraph("Абзац, набранный другим шрифтом вручную.")
    paragraph.runs[0].font.name = "Times New Roman"
    paragraph.runs[0].font.size = Pt(20)
    source = io.BytesIO()
    doc.save(source)

    normal = _Document(io.BytesIO(blank_copy_of_template(source.getvalue()))).styles["Normal"]
    assert normal.font.name == "Arial", "measurement overrode the template's own style"
    assert normal.font.size == Pt(11)


def _template_without_style(style_id: str) -> bytes:
    """A .docx whose styles.xml has no definition for that style — exactly what
    a customer template looks like when it never used bullet lists."""
    import re as _re
    import zipfile as _zipfile
    from docx import Document as _Document

    source = io.BytesIO()
    _Document().save(source)
    out = io.BytesIO()
    with _zipfile.ZipFile(io.BytesIO(source.getvalue())) as zin, \
            _zipfile.ZipFile(out, "w", _zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/styles.xml":
                data = _re.sub(
                    r'<w:style [^>]*w:styleId="%s".*?</w:style>' % style_id, "",
                    data.decode("utf-8"), flags=_re.S,
                ).encode("utf-8")
            zout.writestr(item, data)
    return out.getvalue()


def test_template_without_list_styles_still_renders_lists():
    """The customer's finished documents contained raw markdown («- Заказчик –
    ООО …») because htmldocx builds <ul> with 'List Bullet' and <ol> with
    'List Number'. Neither was in the backfill list, so a template that never
    used lists crashed the conversion on the document's first bullet."""
    from docx import Document as _Document
    from docx_generator import (
        blank_copy_of_template, convert_markdown_to_docx, CONVERSION_FAILED_MARKER,
    )

    for style_id in ("ListBullet", "ListNumber"):
        blank = blank_copy_of_template(_template_without_style(style_id))
        result = convert_markdown_to_docx(
            "## Раздел\n\n- Заказчик – ООО «РЭО»\n- Разработчик – ООО «БИТ»\n\n1. первый\n2. второй",
            base_template_bytes=blank,
        )
        paragraphs = [(p.style.name, p.text.strip()) for p in _Document(io.BytesIO(result)).paragraphs
                      if p.text.strip()]
        texts = [text for _, text in paragraphs]
        assert not any(CONVERSION_FAILED_MARKER in text for text in texts), (style_id, paragraphs)
        assert not any(text.startswith("- ") for text in texts), f"raw markdown survived: {paragraphs}"
        assert any("Заказчик" in text for text in texts), paragraphs


def test_preflight_reads_the_document_not_the_zip_bytes():
    """The guard searched the .docx BYTES for the failure marker — but a .docx
    is a zip, so the compressed text never matched and the guard answered «the
    template is fine» every single time. That is why a broken template reached
    the customer."""
    import docgen_router
    from docx import Document as _Document
    from docx_generator import CONVERSION_FAILED_MARKER

    real_convert = docgen_router.convert_markdown_to_docx if hasattr(
        docgen_router, "convert_markdown_to_docx") else None

    def broken_convert(markdown_text, base_template_bytes=None):
        doc = _Document()
        doc.add_paragraph(CONVERSION_FAILED_MARKER)
        doc.add_paragraph(markdown_text)
        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    import docx_generator
    original = docx_generator.convert_markdown_to_docx
    docx_generator.convert_markdown_to_docx = broken_convert
    try:
        source = io.BytesIO()
        _Document().save(source)
        assert dg._template_renders_cleanly(source.getvalue()) is False
    finally:
        docx_generator.convert_markdown_to_docx = original


def test_every_heading_is_black_with_and_without_a_template():
    """A business document has no blue headings. python-docx defines Heading N
    in accent blue, so both paths must force black: the run level wins over
    the style, and injected styles must not carry the colour in either."""
    from docx import Document as _Document
    from docx.shared import RGBColor
    from docx_generator import blank_copy_of_template, convert_markdown_to_docx

    markdown_text = "## ИСХОДНЫЕ ТЕХНИЧЕСКИЕ ТРЕБОВАНИЯ\n\nТекст.\n\n### Подраздел\n\n- пункт"
    source = io.BytesIO()
    _Document().save(source)
    blank = blank_copy_of_template(_template_without_style("Heading2"))

    for label, template in (("без шаблона", None), ("по шаблону", blank)):
        document = _Document(io.BytesIO(convert_markdown_to_docx(markdown_text, base_template_bytes=template)))
        for paragraph in document.paragraphs:
            for run in paragraph.runs:
                if not run.text.strip():
                    continue
                assert run.font.color is not None and run.font.color.rgb == RGBColor(0, 0, 0), (
                    f"{label}: {paragraph.style.name} run is {run.font.color.rgb}: {run.text[:40]!r}"
                )

    # And the style injected into a customer template brings no colour of its own.
    injected = _Document(io.BytesIO(blank)).styles["Heading 2"]
    assert injected.font.color is None or injected.font.color.type is None, (
        "the backfilled heading style carried python-docx's blue into the template"
    )


def test_apply_replacements_longest_first_prevents_partial_overlap():
    mapping = {
        "АБВГ.123456.789": "СТАЛО.000000.001",
        "123456": "999999",
    }
    text = "Документ АБВГ.123456.789, отдельно упомянут номер 123456."

    result = dg._apply_replacements(text, mapping)

    assert "СТАЛО.000000.001" in result
    assert "999999" in result
    assert "123456" not in result, f"old value survived a partial overlap: {result!r}"


def test_apply_replacements_empty_value_becomes_placeholder():
    mapping = {"ООО «Ромашка»": ""}
    text = "Заказчик: ООО «Ромашка»."

    result = dg._apply_replacements(text, mapping)

    assert "ООО «Ромашка»" not in result
    assert "[УКАЗАТЬ: Организация]" in result


def test_find_replacement_candidates_needs_two_documents():
    chunks = [
        Chunk(id=0, doc_name="a.docx", title="t", text="Изделие АБВГ.123456.789 испытано.", tokens=frozenset()),
        Chunk(id=1, doc_name="b.docx", title="t", text="См. АБВГ.123456.789 в приложении.", tokens=frozenset()),
        Chunk(id=2, doc_name="a.docx", title="t", text="Разовый номер АБВГ.999999.999 только здесь.", tokens=frozenset()),
    ]

    candidates = dg._find_replacement_candidates(chunks)
    values = [v for v, _ in candidates]

    assert "АБВГ.123456.789" in values, "a designation in two documents must be proposed"
    assert "АБВГ.999999.999" not in values, "a designation in a single document must not be proposed"


def test_find_replacement_candidates_never_proposes_gost_reference():
    chunks = [
        Chunk(id=0, doc_name="a.docx", title="t", text="Материал по ГОСТ 123456-99 применяется.", tokens=frozenset()),
        Chunk(id=1, doc_name="b.docx", title="t", text="Материал по ГОСТ 123456-99 применяется здесь же.", tokens=frozenset()),
    ]

    candidates = dg._find_replacement_candidates(chunks)
    values = [v for v, _ in candidates]

    assert not any("123456" in v for v in values), f"a ГОСТ reference leaked into candidates: {values!r}"
