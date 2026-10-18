import hashlib
import io
import json
from types import SimpleNamespace
from uuid import uuid4
import pytest
from fastapi.testclient import TestClient
from app.main import app
from app.auth import get_current_user
from app.core.repository import repo
from app.db.base import Base
from app.db.engine import SyncSessionLocal, sync_engine
from app.db.models import User, Conversation, ConversationMember, Message, Document
from app.services import project_memory as memory, project_results as results, project_assets as assets


@pytest.fixture
def account():
    Base.metadata.create_all(sync_engine)
    email=uuid4().hex+'@example.com';uid=repo._bot_id(email);cid=uuid4().hex
    with SyncSessionLocal() as session:
        session.add(User(email=email,password_hash='test',role='user',bot_user_id=uid))
        session.add(Conversation(id=cid,user_id=uid,title='Проект'));session.commit()
    pid=memory.create(uid,'Плата')['id'];memory.attach(uid,cid,pid)
    app.dependency_overrides[get_current_user]=lambda:email
    yield email,uid,cid,pid
    app.dependency_overrides.clear()


def table():
    return {'type':'table','title':'Сравнение','columns':[{'id':'part','label':'Компонент'},{'id':'price','label':'Цена'}],
            'rows':[{'id':'r1','cells':{'part':'LM2596','price':'200'}},{'id':'r2','cells':{'part':'AMS1117','price':0}}]}


def test_targeted_patch_preserves_other_rows_and_conflicts(account):
    _,uid,_,pid=account
    saved=results.save(uid,pid,table());eid=saved['id']
    results.patch(uid,pid,eid,1,[{'row_id':'r1','column_id':'price','value':'180'}])
    current=results.read(uid,pid,eid)
    assert current['spec']['rows'][1]['cells']=={'part':'AMS1117','price':'0'}
    assert current['spec']['rows'][0]['id']=='r1' and current['revision']==2
    with pytest.raises(ValueError,match='уже изменена'):results.patch(uid,pid,eid,1,[{'row_id':'r1','column_id':'price','value':'1'}])
    with pytest.raises(ValueError):results.patch(uid,pid,eid,None,[{'row_id':'r1','column_id':'price','value':'1'}])
    assert results.read(uid,pid,eid,1)['spec']['rows'][0]['cells']['price']=='200'
    with pytest.raises(ValueError):results.read(uid+1,pid,eid)


def test_exports_are_real_and_spreadsheet_values_cannot_be_formulas(account):
    from openpyxl import load_workbook
    from docx import Document as Word
    from pptx import Presentation
    _,uid,_,pid=account
    spec=table();spec['rows'][0]['cells']['part']='=HYPERLINK("https://example.com")'
    eid=results.save(uid,pid,spec)['id'];data,mime=results.export(uid,pid,eid,'xlsx')
    workbook=load_workbook(io.BytesIO(data));assert workbook.active['A2'].data_type=='s' and workbook.active['B3'].value=='0'
    eid=results.save(uid,pid,{'type':'report','title':'Отчёт','text':'Исходная формулировка\nИтог.'})['id']
    results.patch(uid,pid,eid,1,[{'old':'Итог.','value':'Проверенный итог.'}])
    data,_=results.export(uid,pid,eid,'docx');assert any(p.text=='Проверенный итог.' for p in Word(io.BytesIO(data)).paragraphs)
    eid=results.save(uid,pid,{'type':'slides','title':'Доклад','slides':[{'id':'s1','title':'Задача','text':'Описание'},{'id':'s2','title':'Итог','text':'Вывод'}]})['id']
    results.patch(uid,pid,eid,1,[{'slide_id':'s2','field':'text','value':'Новый вывод'}])
    data,_=results.export(uid,pid,eid,'pptx');deck=Presentation(io.BytesIO(data));assert len(deck.slides)==2
    assert any(getattr(shape,'text','')=='Новый вывод' for shape in deck.slides[1].shapes)


def test_api_versions_and_project_scope(account):
    _,uid,cid,pid=account;client=TestClient(app)
    saved=client.post(f'/projects/{pid}/results',json={'spec':table()}).json();eid=saved['id']
    assert client.put(f'/projects/{pid}/results/{eid}',json={'spec':table()}).status_code==400
    assert client.get(f'/projects/{pid}/results/{eid}/export/xlsx').status_code==200
    assert memory.workspace(uid,pid)['chats'][0]['id']==cid
    with SyncSessionLocal() as session:session.add(ConversationMember(conversation_id=cid,user_id=uid+1));session.commit()
    assert memory.workspace(uid,pid)['chats']==[]
    with pytest.raises(ValueError):memory.workspace(uid+1,pid)


def test_purchase_plan_is_bounded_and_unconfirmed(account):
    from app.services.run_protocol import RunRecord,current_run
    _,uid,_,pid=account
    with pytest.raises(ValueError):results.purchase_spec('Закупка',[{'part':'R1'}]*21)
    with pytest.raises(ValueError):results.purchase_spec('Закупка',[{'quantity':20}])
    spec=results.purchase_spec('Закупка',[{'part':'LM2596','quantity':10,'requirements':'SOIC-8'}])
    assert spec['rows'][0]['state']=='pending'
    spec['rows'][0]['state']='checked';spec['rows'][0]['cells']['source']='https://example.com/item'
    token=current_run.set(RunRecord())
    try:eid=results.save(uid,pid,spec,model=True)['id']
    finally:current_run.reset(token)
    assert results.read(uid,pid,eid)['spec']['rows'][0]['state']=='pending'
    token=current_run.set(RunRecord())
    try:assert results.save(uid,pid,spec,model=True)['id']==eid
    finally:current_run.reset(token)


def test_empty_project_is_known_and_oversized_text_is_not_silently_lost(account):
    _,uid,cid,pid=account
    assert pid in memory.context(uid,cid,'подготовь таблицу')
    with pytest.raises(ValueError):results.save(uid,pid,{'type':'report','title':'Отчёт','text':'x'*40001})


def test_paged_tool_read_cannot_silently_replace_complete_result(account):
    _,uid,_,pid=account;spec=table();spec['rows']=[{'id':str(i),'cells':{'part':'A'*500,'price':'20'}} for i in range(10)]
    eid=results.save(uid,pid,spec)['id'];page=results.tool_read(uid,pid,eid)
    assert len(page['spec']['rows'])==4 and page['next_offset']==4 and page['shortened_cells']
    assert len(results.read(uid,pid,eid)['spec']['rows'])==10


def test_purchase_offer_checks_quotes_and_preserves_previous_observations(account):
    from app.services.run_protocol import RunRecord,current_run
    _,uid,_,pid=account
    spec=results.purchase_spec('Компоненты',[{'part':'LM2596','quantity':10},{'part':'AMS1117','quantity':20}])
    eid=results.save(uid,pid,spec)['id']
    record=RunRecord();record.observe_page('URL: https://example.com/lm2596\n\nLM2596. Цена: 200 рублей. В наличии: 12 шт. '+ 'Характеристики. '*10)
    token=current_run.set(record)
    try:
        with pytest.raises(ValueError):results.purchase_offer(uid,pid,eid,1,'row-1','https://example.com/lm2596',{'price':'Цена: 1 рубль','available':'В наличии: 12 шт.'})
        results.purchase_offer(uid,pid,eid,1,'row-1','https://example.com/lm2596',{'price':'Цена: 200 рублей.','available':'В наличии: 12 шт.'})
    finally:current_run.reset(token)
    old=results.read(uid,pid,eid)['spec']['rows'][0]
    assert old['state']=='checked' and old['cells']['checked']==record.sources[0]['read_at']
    record2=RunRecord();record2.observe_page('URL: https://example.com/ams1117\n\nAMS1117. Цена: 20 рублей. В наличии: 100 шт. '+ 'Характеристики. '*10)
    token=current_run.set(record2)
    try:results.purchase_offer(uid,pid,eid,2,'row-2','https://example.com/ams1117',{'price':'Цена: 20 рублей.','available':'В наличии: 100 шт.'})
    finally:current_run.reset(token)
    current=results.read(uid,pid,eid)['spec']['rows'];assert current[0]==old and current[1]['state']=='checked'


def test_final_table_is_captured_without_model_and_with_links(account):
    _,uid,cid,pid=account
    result=results.capture_answer(uid,cid,'Сравни две рации','| Товар | Цена |\n| --- | --- |\n| [Рация А](https://example.com/a) | 2000 ₽ |\n| Рация Б | 1800 ₽ |')
    spec=results.read(uid,pid,result['id'])['spec']
    assert len(spec['rows'])==2 and 'https://example.com/a' in spec['rows'][0]['cells']['c0']
    assert results.capture_answer(uid,cid,'Привет','Привет') is None


def test_originals_are_private_deduplicated_and_survive_source_removal(account,tmp_path,monkeypatch):
    email,uid,cid,pid=account
    monkeypatch.setattr(assets,'get_settings',lambda:SimpleNamespace(UPLOAD_DIR=tmp_path/'uploads',DATA_DIR=tmp_path/'data'))
    key=hashlib.sha256(email.encode()).hexdigest()[:16];source=tmp_path/'uploads'/'chat'/key/'original.zip';source.parent.mkdir(parents=True);source.write_bytes(b'original archive')
    with SyncSessionLocal() as session:
        session.add(Message(conversation_id=cid,role='user',content='Материалы',attachment=json.dumps({'name':'1.zip','url':f'/uploads/chat/{key}/original.zip'})))
        docs=[Document(conversation_id=cid,filename=f'1.zip/{i}.pdf',content='Текст документа') for i in range(2)]
        session.add_all(docs);session.commit();ids=[d.id for d in docs]
    entries=[memory.publish_document(uid,pid,cid,did)['id'] for did in ids]
    source.unlink()
    path,name=assets.original(uid,pid,entries[0]);assert path.read_bytes()==b'original archive' and name=='1.zip'
    assert not path.is_relative_to(tmp_path/'uploads')
    assert len(list(assets.folder(uid,pid).glob('*.zip')))==1
    with pytest.raises(ValueError):assets.original(uid+1,pid,entries[0])
    memory.remove(uid,pid,entries[0]);assert assets.original(uid,pid,entries[1])[0].exists()
    memory.remove(uid,pid,entries[1]);assert not path.exists()
