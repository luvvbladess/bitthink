import asyncio
import hashlib
import io
import json
import zipfile
from types import SimpleNamespace
from urllib.parse import parse_qs, urlsplit

import pytest
from fastapi import HTTPException
from jose import jwt
from lxml import etree as xml
from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference

from app.config import get_settings
from app.services.office_viewer import spreadsheet_source, spreadsheet_viewer


def workbook_bytes():
    w = Workbook(); s = w.active; s.title = 'Отчёт'
    s.append(['Месяц', 'Выручка, ₽', 'Расходы, ₽'])
    s.append(['Август', 480000, 350000]); s.append(['Сентябрь', 540000, 380000])
    s['D2'] = '=B2-C2'
    chart = LineChart(); chart.title = 'Выручка и расходы'
    chart.add_data(Reference(s, min_col=2, max_col=3, min_row=1, max_row=3), titles_from_data=True)
    chart.set_categories(Reference(s, min_col=1, min_row=2, max_row=3))
    s.add_chart(chart, 'A6')
    b = io.BytesIO(); w.save(b); return b.getvalue()


def test_viewing_copy_repairs_chart_axes_and_caches_without_changing_cells():
    data = workbook_bytes(); viewed = spreadsheet_source(data, 'Отчёт.xlsx')
    with zipfile.ZipFile(io.BytesIO(data)) as original, zipfile.ZipFile(io.BytesIO(viewed)) as result:
        assert original.namelist() == result.namelist()
        for name in original.namelist():
            if not name.startswith('xl/charts/chart'):
                assert original.read(name) == result.read(name), name
        ns = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
        root = xml.fromstring(result.read('xl/charts/chart1.xml'))
        assert root.find('.//c:catAx/c:axPos', ns).get('val') == 'b'
        assert root.find('.//c:cat/c:strRef/c:strCache/c:pt/c:v', ns).text == 'Август'
        assert root.find('.//c:val/c:numRef/c:numCache/c:pt/c:v', ns).text == '480000'
    parsed = load_workbook(io.BytesIO(viewed))
    assert parsed.active['D2'].value == '=B2-C2'
    assert parsed.active['B2'].value == 480000


def test_native_viewer_is_read_only_scoped_to_content_and_does_not_prompt_for_name(monkeypatch):
    settings = get_settings(); monkeypatch.setattr(settings, 'ONLYOFFICE_JWT_SECRET', 'viewer-test-secret')
    data = workbook_bytes(); viewer = spreadsheet_viewer(data, 'Отчёт.xlsx', 'my-chat', True)
    c = viewer['config']; token = c.pop('token')
    assert jwt.decode(token, 'viewer-test-secret', algorithms=['HS256']) == c
    assert c['documentType'] == 'cell' and c['type'] == 'mobile'
    assert c['editorConfig']['mode'] == 'view' and c['editorConfig']['user']['name']
    assert not c['document']['permissions']['edit'] and 'callbackUrl' not in c['editorConfig']
    assert 'plugins' not in c['editorConfig']
    query = parse_qs(urlsplit(c['document']['url']).query)
    assert query['v'] == [hashlib.sha256(data).hexdigest()]
    from app.api.editor import check_file_link
    assert check_file_link(query['t'][0]) == ('my-chat', 'Отчёт.xlsx')
    assert spreadsheet_viewer(data+b'changed', 'Отчёт.xlsx', 'my-chat')['config']['document']['key'] != c['document']['key']
    assert spreadsheet_viewer(data, 'Отчёт.xlsx', 'other-chat')['config']['document']['key'] != c['document']['key']
    assert spreadsheet_viewer(data, 'Отчёт.pdf', 'my-chat') is None


def test_native_opt_in_bypasses_print_conversion_and_keeps_legacy_clients(monkeypatch):
    from app.api import editor, conversations
    from app.services import document_preview
    settings = get_settings(); monkeypatch.setattr(settings, 'ONLYOFFICE_JWT_SECRET', 'viewer-test-secret')
    async def allowed(*args): return 1, object()
    monkeypatch.setattr(editor, '_conversation_for', allowed)
    monkeypatch.setattr(editor, 'file_bytes', lambda *args: workbook_bytes())
    calls = []
    monkeypatch.setattr(document_preview, 'preview', lambda *args: calls.append(args) or {'pages': 2})
    response = asyncio.run(conversations.preview_conversation_file('mine', 'Report.xlsx', user_id='owner', native=True))
    assert json.loads(response.body)['viewer']['config']['documentType'] == 'cell' and not calls
    response = asyncio.run(conversations.preview_conversation_file('mine', 'Report.xlsx', user_id='owner'))
    assert json.loads(response.body)['pages'] == 2 and len(calls) == 1


def test_viewer_source_checks_signature_conversation_and_content_version(monkeypatch):
    from app.api import editor, conversations
    data = workbook_bytes(); accesses = []
    monkeypatch.setattr(editor, 'file_bytes', lambda *args: accesses.append(args) or data)
    monkeypatch.setattr(editor, 'check_file_link', lambda token: ('mine', 'Report.xlsx'))
    with pytest.raises(HTTPException) as error:
        asyncio.run(conversations.spreadsheet_viewer_source('other', 'token', hashlib.sha256(data).hexdigest()))
    assert error.value.status_code == 403 and not accesses
    monkeypatch.setattr(conversations.repo, 'conversations', SimpleNamespace(owner=lambda cid: 1, view=lambda uid, cid: object()))
    with pytest.raises(HTTPException) as error:
        asyncio.run(conversations.spreadsheet_viewer_source('mine', 'token', 'stale'))
    assert error.value.status_code == 409
    response = asyncio.run(conversations.spreadsheet_viewer_source('mine', 'token', hashlib.sha256(data).hexdigest()))
    assert response.headers['cache-control'] == 'private, no-store'
    assert load_workbook(io.BytesIO(response.body)).active['B2'].value == 480000


def test_missing_configuration_unsupported_and_invalid_workbooks(monkeypatch):
    settings = get_settings(); monkeypatch.setattr(settings, 'ONLYOFFICE_JWT_SECRET', '')
    assert spreadsheet_viewer(b'x', 'Report.xlsx', 'mine') is None
    with pytest.raises(ValueError, match='прочитать'):
        spreadsheet_source(b'not a zip', 'Report.xlsx')
    assert spreadsheet_source(b'legacy', 'Report.xls') == b'legacy'
