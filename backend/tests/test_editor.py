"""Редактор документов: открытие, сохранения от OnlyOffice и ИИ-правка фрагмента."""

import io

import pytest
from fastapi.testclient import TestClient
from jose import jwt

from app.config import get_settings
from app.bot_core.db_conversations import DatabaseConversationManager


def _docx(text: str) -> bytes:
    from docx import Document

    document = Document()
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def editor(tmp_path, monkeypatch):
    settings = get_settings()
    monkeypatch.setattr(settings, "ONLYOFFICE_JWT_SECRET", "ds-secret")
    monkeypatch.setattr(settings, "DATA_DIR", tmp_path / "data")
    monkeypatch.setattr(settings, "UPLOAD_DIR", tmp_path / "uploads")
    monkeypatch.setattr(settings, "ONLYOFFICE_PUBLIC_URL", "https://site.example/onlyoffice")
    monkeypatch.setattr(settings, "ONLYOFFICE_INTERNAL_URL", "http://onlyoffice")
    from app.api import editor as module

    return module


def _room_with_contract(settings, owner: int, guest: int):
    manager = DatabaseConversationManager()
    conv = manager.create_conversation(owner, title="Договор")
    stored = settings.UPLOAD_DIR / "generated" / "abcdef0123456789" / ("0" * 32 + ".docx")
    stored.parent.mkdir(parents=True, exist_ok=True)
    stored.write_bytes(_docx("Срок оплаты 30 дней."))
    manager.add_message(
        owner, "assistant", "Готово", conv_id=conv.id,
        attachment={"name": "Договор.docx", "status": "done", "type": "generated",
                    "url": f"/uploads/generated/abcdef0123456789/{'0' * 32}.docx"},
    )
    token = manager.create_share(owner, conv.id)
    manager.join_share(guest, token)
    return manager, manager.conversation_view(owner, conv.id)


def test_members_share_one_document_and_saves_version_it(editor):
    settings = get_settings()
    manager, conv = _room_with_contract(settings, 881_001, 881_002)
    guest_view = manager.conversation_view(881_002, conv.id)

    first = editor._open_record(conv.id, "Договор.docx", 881_001, lambda: editor.source_bytes(conv, "Договор.docx"))
    second = editor._open_record(guest_view.id, "Договор.docx", 881_002, lambda: editor.source_bytes(guest_view, "Договор.docx"))
    assert first == second  # тот же key – совместная правка вживую
    doc_id, version = first
    assert version == 1 and editor.version_path(doc_id, 1).is_file()

    # Принудительное сохранение не рвёт живую сессию: версия та же.
    assert editor.store_saved_version(doc_id, _docx("Срок оплаты 10 дней."), bump=False) == 1
    # Все закрыли документ – следующее открытие получит новый key.
    assert editor.store_saved_version(doc_id, _docx("Срок оплаты 10 рабочих дней."), bump=True) == 2
    assert editor._open_record(conv.id, "Договор.docx", 881_001, lambda: None) == (doc_id, 2)
    # Чат видит исправленный текст.
    docs = manager.get_documents(881_001, conv_id=conv.id)
    assert any("10 рабочих дней" in item["content"] for item in docs)

    manager.delete_conversation(881_001, conv.id)


def test_edit_becomes_the_chat_file_and_the_assistant_knows(editor):
    settings = get_settings()
    manager, conv = _room_with_contract(settings, 881_031, 881_032)
    doc_id, _ = editor._open_record(conv.id, "Договор.docx", 881_031, lambda: editor.source_bytes(conv, "Договор.docx"))
    edited = _docx("Срок оплаты 10 рабочих дней.")
    # Автосохранение во время правки (без повышения версии) уже обновляет файл беседы.
    editor.store_saved_version(doc_id, edited, bump=False, edited_by=881_032)

    stored = settings.UPLOAD_DIR / "generated" / "abcdef0123456789" / ("0" * 32 + ".docx")
    assert stored.read_bytes() == edited  # карточка в чате отдаёт правку
    assert editor.version_path(doc_id, 1).is_file()  # исходник остался v1

    files = editor.conversation_files(manager.conversation_view(881_032, conv.id))
    item = next(f for f in files if f["name"] == "Договор.docx")
    assert item["uploaded_by"] == "Bit-Think" and item["edited_at"] and item["editable"]
    conv_id, name = editor.check_file_link(item["download_url"].split("t=", 1)[1])
    assert (conv_id, name) == (conv.id, "Договор.docx")
    assert editor.file_bytes(conv, "Договор.docx") == edited

    manager.add_message(881_031, "user", "какой срок оплаты в договоре?", conv_id=conv.id, author_user_id=881_031)
    api = manager.get_messages_for_api(881_031, "системный", requesting_user_id=881_031, conv_id=conv.id)
    blob = "\n".join(str(m.get("content") or "") for m in api)
    assert "Изменён в редакторе" in blob and "10 рабочих дней" in blob
    assert "Работай с их текущей версией" in blob
    manager.delete_conversation(881_031, conv.id)


def test_phone_edit_lands_as_tracked_change_and_waits_for_desktop_editors(editor):
    from app.auth import create_access_token
    from app.core.repository import repo

    settings = get_settings()
    email = "phone-owner@example.ru"
    owner = repo._bot_id(email)
    manager, conv = _room_with_contract(settings, owner, 881_042)
    from app.main import app
    doc_id, _ = editor._open_record(conv.id, "Договор.docx", owner, lambda: editor.source_bytes(conv, "Договор.docx"))
    client = TestClient(app)
    auth = {"Authorization": f"Bearer {create_access_token({'sub': email})}"}

    listed = client.get(f"/editor/{doc_id}/paragraphs", headers=auth).json()
    item = listed["paragraphs"][0]
    assert item["text"] == "Срок оплаты 30 дней." and listed["busy"] == []

    # Someone has it open in OnlyOffice on a computer: the phone waits.
    editor._EDITING[doc_id] = {"881042"}
    body = {"index": item["index"], "base_text": item["text"], "text": "Срок оплаты 10 рабочих дней."}
    assert client.post(f"/editor/{doc_id}/paragraphs", json=body, headers=auth).status_code == 423
    # Their own desktop session counts too: it would save over the phone edit.
    editor._EDITING[doc_id] = {str(owner)}
    assert client.get(f"/editor/{doc_id}/paragraphs", headers=auth).json()["busy"] == ["вы"]
    assert client.post(f"/editor/{doc_id}/paragraphs", json=body, headers=auth).status_code == 423
    editor._EDITING.pop(doc_id)

    saved = client.post(f"/editor/{doc_id}/paragraphs", json=body, headers=auth)
    assert saved.status_code == 200, saved.text
    assert saved.json()["version"] == 2
    # Same stale text again: the paragraph already changed.
    assert client.post(f"/editor/{doc_id}/paragraphs", json=body, headers=auth).status_code == 409

    stored = settings.UPLOAD_DIR / "generated" / "abcdef0123456789" / ("0" * 32 + ".docx")
    from docx import Document

    xml = Document(io.BytesIO(stored.read_bytes())).element.xml
    assert "w:del " in xml and "w:ins " in xml  # the chat file carries the tracked change
    files = editor.conversation_files(manager.conversation_view(owner, conv.id))
    assert files[0]["version"] == 2 and files[0]["edited_at"]

    stranger = {"Authorization": f"Bearer {create_access_token({'sub': 'stranger@example.ru'})}"}
    assert client.get(f"/editor/{doc_id}/paragraphs", headers=stranger).status_code == 404
    manager.delete_conversation(owner, conv.id)


def test_callback_requires_document_server_signature(editor):
    settings = get_settings()
    manager, conv = _room_with_contract(settings, 881_011, 881_012)
    doc_id, _ = editor._open_record(conv.id, "Договор.docx", 881_011, lambda: editor.source_bytes(conv, "Договор.docx"))
    from app.main import app

    client = TestClient(app)
    link = editor.sign_link(doc_id, "callback")
    body = {"status": 1, "key": f"{doc_id}-1"}
    assert client.post(f"/editor/{doc_id}/callback?t={link}", json=body).status_code == 403
    signed = {"token": jwt.encode(body, "ds-secret", algorithm="HS256")}
    assert client.post(f"/editor/{doc_id}/callback?t={link}", json=signed).json() == {"error": 0}
    # Чужая ссылка (от файла) не подходит к колбэку.
    wrong = editor.sign_link(doc_id, "file")
    assert client.post(f"/editor/{doc_id}/callback?t={wrong}", json=signed).status_code == 403
    manager.delete_conversation(881_011, conv.id)


def test_saved_file_is_fetched_only_from_internal_document_server(editor):
    assert (
        editor.internal_download_url("https://site.example/onlyoffice/cache/files/data/key/output.docx?md5=x")
        == "http://onlyoffice/cache/files/data/key/output.docx?md5=x"
    )
    assert editor.internal_download_url("http://169.254.169.254/latest/meta-data") == "http://onlyoffice/latest/meta-data"


def test_ai_edit_needs_document_token_and_room_access(editor, monkeypatch):
    settings = get_settings()
    manager, conv = _room_with_contract(settings, 881_021, 881_022)
    doc_id, _ = editor._open_record(conv.id, "Договор.docx", 881_021, lambda: editor.source_bytes(conv, "Договор.docx"))
    seen = {}

    async def fake_chat(messages, **kwargs):
        seen["prompt"] = messages[-1]["content"]
        return "«Срок оплаты – 10 рабочих дней.»", [], "", []

    import openai_client

    monkeypatch.setattr(openai_client, "get_chat_response", fake_chat)
    from app.main import app

    client = TestClient(app)
    payload = {"selection": "Срок оплаты 30 дней.", "instruction": "сделай 10 рабочих дней"}
    guest_token = editor.sign_link(doc_id, "ai", user=881_022)
    ok = client.post(f"/editor/{doc_id}/ai", json=payload, headers={"Authorization": f"Bearer {guest_token}"})
    assert ok.status_code == 200, ok.text
    assert ok.json() == {"text": "Срок оплаты – 10 рабочих дней."}
    assert "Срок оплаты 30 дней." in seen["prompt"] and "сделай 10 рабочих дней" in seen["prompt"]

    stranger = editor.sign_link(doc_id, "ai", user=881_099)
    assert client.post(f"/editor/{doc_id}/ai", json=payload, headers={"Authorization": f"Bearer {stranger}"}).status_code == 403
    file_link = editor.sign_link(doc_id, "file", user=881_022)
    assert client.post(f"/editor/{doc_id}/ai", json=payload, headers={"Authorization": f"Bearer {file_link}"}).status_code == 403
    manager.delete_conversation(881_021, conv.id)


def test_library_lists_files_from_every_chat_including_shared(editor):
    from app.auth import create_access_token
    from app.core.repository import repo

    settings = get_settings()
    email = "library-guest@example.ru"
    guest = repo._bot_id(email)
    manager, room = _room_with_contract(settings, 881_051, guest)
    from app.main import app  # after the uploads dir exists: the app mounts it
    own = manager.create_conversation(guest, title="Мой чат")
    manager.add_message(guest, "assistant", "Готово", conv_id=own.id, attachment={
        "name": "Смета.xlsx", "status": "done", "type": "generated",
        "url": "/uploads/generated/abcdef0123456789/x.xlsx"})

    try:
        listed = TestClient(app).get(
            "/conversations/library", headers={"Authorization": f"Bearer {create_access_token({'sub': email})}"}
        ).json()
        where = {(item["name"], item["conversation_id"]): item for item in listed}
        assert ("Договор.docx", room.id) in where  # the shared room counts too
        sheet = where[("Смета.xlsx", own.id)]
        assert sheet["conversation_title"] == "Мой чат" and sheet["download_url"]
    finally:
        manager.delete_conversation(881_051, room.id)
        manager.delete_conversation(guest, own.id)
