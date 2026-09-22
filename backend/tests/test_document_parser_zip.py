import asyncio
import io
import zipfile

from app.config import get_settings  # noqa: F401 — adds bot_core to sys.path
from document_parser import extract_zip_archive


def _make_zip(files: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, data in files.items():
            zf.writestr(name, data)
    return buf.getvalue()


def test_extract_zip_archive_reads_supported_files_and_skips_junk():
    zip_bytes = _make_zip({
        "spec.txt": "Техническое задание".encode("utf-8"),
        "notes.txt": "Заметки инженера".encode("utf-8"),
        "image.bin": b"\x00\x01\x02",
        "__MACOSX/._spec.txt": b"junk",
        "folder/": b"",
    })

    results = asyncio.run(extract_zip_archive(zip_bytes, "archive.zip", user_id=None))

    names = [name for name, _ in results]
    assert names == ["archive.zip/spec.txt", "archive.zip/notes.txt"]
    assert "Техническое задание" in dict(results)["archive.zip/spec.txt"]


def test_extract_zip_archive_normalizes_windows_backslash_paths():
    """Some Windows zip tools store 'shablony\\forma.txt' instead of '/'.
    Folder grouping in docgen splits on '/', so a backslash would hide the
    template folder from the planner.

    ZipInfo() itself rewrites os.sep to '/' on Windows, so the filename is
    assigned after construction to keep the backslash the way a foreign
    zip tool would have stored it.
    """
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for raw_name, body in (
            ("shablony\\forma.txt", "Шаблон формы".encode("utf-8")),
            ("baza\\otchet.txt", "База знаний".encode("utf-8")),
        ):
            info = zipfile.ZipInfo("placeholder.txt")
            info.filename = raw_name
            zf.writestr(info, body)
    zip_bytes = buf.getvalue()

    results = asyncio.run(extract_zip_archive(zip_bytes, "arhiv.zip", user_id=None))
    names = [name for name, _ in results]

    assert "arhiv.zip/shablony/forma.txt" in names
    assert "arhiv.zip/baza/otchet.txt" in names
    assert not any("\\" in name for name in names)


def test_extract_zip_archive_empty_zip_returns_empty_list():
    zip_bytes = _make_zip({})

    results = asyncio.run(extract_zip_archive(zip_bytes, "empty.zip", user_id=None))

    assert results == []

def test_doc_text_is_read_without_a_converter():
    """Four of the five formatting examples a customer sent were legacy .doc,
    and extraction used to return None — the files were dropped whole, so the
    mode saw one example out of five. Word 97+ keeps its text as UTF-16LE
    inside the OLE container, which needs no dependency at all."""
    import asyncio

    from app.bot_core import document_parser as parser

    body = "ИСХОДНЫЕ ТЕХНИЧЕСКИЕ ТРЕБОВАНИЯ (ИТТ) на поставку системы водотушения"
    # Префикс обязан быть чётной длины: UTF-16 читается парами байт,
    # и лишний байт сдвинул бы весь текст.
    blob = bytes([0, 1, 2, 4]) + body.encode("utf-16-le") + bytes([255, 254])
    original = parser.convert_doc_to_docx
    parser.convert_doc_to_docx = lambda data, name="source.doc": None
    try:
        text = asyncio.run(parser.extract_text_from_doc(blob))
    finally:
        parser.convert_doc_to_docx = original
    assert "ИСХОДНЫЕ ТЕХНИЧЕСКИЕ ТРЕБОВАНИЯ" in text


def test_doc_becomes_a_usable_template_only_when_converted(tmp_path):
    """Only a .docx can be opened as a formatting base, so a .doc example is
    stored converted. Without LibreOffice nothing is stored — the mode then
    falls back to standard formatting instead of a broken template."""
    from docx import Document

    from app.config import get_settings
    from app.bot_core import document_parser as parser

    settings = get_settings()
    original_dir, settings.UPLOAD_DIR = settings.UPLOAD_DIR, tmp_path
    original_convert = parser.convert_doc_to_docx
    try:
        parser.convert_doc_to_docx = lambda data, name="source.doc": None
        parser.store_source_docx(1, "пример.doc", b"binary .doc payload")
        assert parser.load_source_docx(1, "пример.doc") is None

        buffer = io.BytesIO()
        Document().save(buffer)
        parser.convert_doc_to_docx = lambda data, name="source.doc": buffer.getvalue()
        parser.store_source_docx(1, "пример.doc", b"binary .doc payload")
        stored = parser.load_source_docx(1, "пример.doc")
        assert stored is not None and stored[:2] == b"PK"
    finally:
        parser.convert_doc_to_docx = original_convert
        settings.UPLOAD_DIR = original_dir


def test_converter_is_optional():
    """The image may not carry LibreOffice; the call must answer None rather
    than raise, because uploads run through this path."""
    from app.bot_core import document_parser as parser

    original = parser._soffice_binary
    parser._soffice_binary = lambda: None
    try:
        assert parser.convert_doc_to_docx(b"whatever") is None
    finally:
        parser._soffice_binary = original


def _xlsx_bytes(rows, sheet="Данные", extra_sheets=None) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    for row in rows:
        worksheet.append(list(row))
    for name, extra_rows in extra_sheets or ():
        other = workbook.create_sheet(name)
        for row in extra_rows:
            other.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pdf_bytes(text: str) -> bytes:
    import fitz

    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    data = document.tobytes()
    document.close()
    return data


def test_extract_text_from_file_reads_excel_and_pdf():
    from app.bot_core import document_parser as parser

    xlsx = _xlsx_bytes(
        [
            ("Узел", "Мощность, кВт", "Среда"),
            ("Пиролизный реактор", 2500, "азот"),
        ],
        extra_sheets=(("Нормы", [("Давление", "16 бар")]),),
    )
    excel_text = asyncio.run(parser.extract_text_from_file(xlsx, "база.XLSX"))
    assert "Пиролизный реактор" in excel_text
    assert "2500" in excel_text
    assert "16 бар" in excel_text
    assert "Нормы" in excel_text

    xlsm_text = asyncio.run(parser.extract_text_from_file(xlsx, "таблица.xlsm"))
    assert "Пиролизный реактор" in xlsm_text

    nameless = asyncio.run(parser.extract_text_from_file(xlsx, "база знаний"))
    assert "Пиролизный реактор" in nameless

    pdf_text = asyncio.run(parser.extract_text_from_file(_pdf_bytes("ITT pump specification 120 m3h"), "scan.PDF"))
    assert "ITT pump specification 120 m3h" in pdf_text


def test_extract_zip_archive_reads_pdf_and_excel():
    from app.bot_core import document_parser as parser

    zip_bytes = _make_zip({
        "шаблоны/ITT.pdf": _pdf_bytes("ITT template for the pump unit"),
        "знания/параметры.xlsx": _xlsx_bytes([
            ("Параметр", "Значение"),
            ("Подача", "120 м3/ч"),
        ]),
        "знания/NOTES.TXT": "Заметка инженера".encode("utf-8"),
    })

    results = asyncio.run(parser.extract_zip_archive(zip_bytes, "пакет.zip", user_id=None))
    bodies = dict(results)

    assert any(name.endswith("ITT.pdf") for name in bodies)
    assert any(name.endswith("параметры.xlsx") for name in bodies)
    assert "ITT template for the pump unit" in "\n".join(bodies.values())
    assert "120 м3/ч" in "\n".join(bodies.values())
